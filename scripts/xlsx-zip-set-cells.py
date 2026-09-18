#!/usr/bin/env python3
"""xlsx-zip-set-cells.py — xlsx / xlsm の値セルだけを zip 直編集で書き換える (Excel も openpyxl も使わない。 他の zip member は byte 同一、 VBA・drawings・form control は無傷。 office-automation.md#xlsx-cell-value-zip-surgery の実装)

なぜ: openpyxl の save は drawings / form control / 一部の書式を落とし (#openpyxl-destroys-drawings)、
Excel の AppleScript は遅いマシンで cold-start 90-120 秒・マクロ付き様式で crash する (#xlsm-macro-export-trap)。
値を入れるだけなら worksheet XML の <c> 要素を差し替えるのが一番壊れない。

使い方:
  xlsx-zip-set-cells.py BOOK --sheet NAME --set C12=文字列 [--set D12=#45672] (--out OUT | --in-place)
  xlsx-zip-set-cells.py BOOK --spec cells.json (--out OUT | --in-place)
      cells.json = {"sheet": "NAME", "cells": {"C12": "文字列", "D12": 45672, "E12": null}}
  xlsx-zip-set-cells.py --selftest

値の書き方 (--set):  文字列 → inline string (t="inlineStr"、 sharedStrings を触らない)
                     #<数> → 数値 <v> (日付書式済みの cell には date serial を入れる)
                     ##... → 先頭の # を 1 つ落とした文字列 / 空 (C12=) → 値を消す (style は残す)
挙動:
  - 既存 cell の s= (style) は保持、 t= は付け直す。 cell が無ければ同じ行に列順を保って挿入し、 style は
    同じ列の最寄りの行の cell から借りる。 行が無ければ行順を保って挿入。 <dimension> と行の spans を広げる
  - 数式 cell (<f>) は既定で拒否 (--allow-formula で上書き。 共有数式の親は常に拒否)。 数式を消したら
    calcChain.xml を外す (Excel が開くとき作り直す = 残すと stale で「修復」 になる)
  - 書いた後に読み戻して全 cell の値を照合 + 他 member の CRC 一致 + check-xlsx-integrity.py を回す。
    どれか落ちたら出力を消して exit 1
  - 既定は --out 必須 (上書きは --in-place を明示)。 git 管理下なら元に戻せるので --in-place でよい
exit: 0 = 書いた / 1 = 検査落ち・拒否 / 2 = 引数・file の誤り
"""
from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from html import unescape
from pathlib import Path

HERE = Path(__file__).resolve().parent
INTEGRITY = HERE / "check-xlsx-integrity.py"
ROW_RE = re.compile(r"<row\b[^>]*?/>|<row\b[^>]*?>.*?</row>", re.S)
CELL_RE = re.compile(r"<c\b[^>]*?/>|<c\b[^>]*?>.*?</c>", re.S)
REF_RE = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")


class Refused(Exception):
    pass


def col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def col_letters(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _xml_text(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def sheet_part(z: zipfile.ZipFile, name: str):
    """sheet 名 → (worksheet part path, sheetId)。 workbook.xml と workbook.xml.rels から引く。"""
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    names = []
    for m in re.finditer(r"<sheet\b[^>]*/?>", wb):
        tag = m.group(0)
        nm = unescape(re.search(r'\bname="([^"]*)"', tag).group(1))
        names.append(nm)
        if nm != name:
            continue
        rid = re.search(r'\br:id="([^"]+)"', tag) or re.search(r'\b\w+:id="([^"]+)"', tag)
        sid = re.search(r'\bsheetId="(\d+)"', tag).group(1)
        rel = re.search(r'<Relationship\b[^>]*\bId="%s"[^>]*/?>' % re.escape(rid.group(1)), rels).group(0)
        target = re.search(r'\bTarget="([^"]+)"', rel).group(1)
        path = target.lstrip("/") if target.startswith("/") else "xl/" + target
        return os.path.normpath(path).replace(os.sep, "/"), sid
    raise SystemExit(f"xlsx-zip-set-cells: sheet「{name}」 が無い (在るのは {names})")


def _attr(tag, key):
    m = re.search(r'\b%s="([^"]*)"' % key, tag)
    return m.group(1) if m else None


def _new_cell(ref: str, style, value) -> str:
    s = f' s="{style}"' if style is not None else ""
    if value is None:
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        v = str(int(value)) if float(value).is_integer() else repr(float(value))
        return f'<c r="{ref}"{s}><v>{v}</v></c>'
    text = str(value)
    sp = ' xml:space="preserve"' if text != text.strip() or "\n" in text else ""
    return f'<c r="{ref}"{s} t="inlineStr"><is><t{sp}>{_xml_text(text)}</t></is></c>'


def _style_from_column(sheet_xml: str, col: str, row: int):
    best = None
    for m in re.finditer(r'<c\b[^>]*\br="%s(\d+)"[^>]*>' % col, sheet_xml):
        st = _attr(m.group(0), "s")
        if st is None:
            continue
        d = abs(int(m.group(1)) - row)
        if best is None or d < best[0]:
            best = (d, st)
    return best[1] if best else None


def set_cell(xml: str, ref: str, value, allow_formula=False):
    """1 cell を差し替えた xml と、 消した数式があれば True を返す。"""
    col, row = REF_RE.match(ref).groups()
    row, ci = int(row), col_index(col)
    sd = re.search(r"<sheetData\s*/>|<sheetData>(.*?)</sheetData>", xml, re.S)
    if sd.group(0).endswith("/>"):
        xml = xml[:sd.start()] + "<sheetData></sheetData>" + xml[sd.end():]
        sd = re.search(r"<sheetData>(.*?)</sheetData>", xml, re.S)
    base = sd.start(1)
    body = sd.group(1)
    rows = [(int(_attr(m.group(0)[:m.group(0).index(">") + 1], "r")), m) for m in ROW_RE.finditer(body)]
    hit = [m for r, m in rows if r == row]
    style = _style_from_column(xml, col, row)
    if not hit:  # 行が無い → 行順に挿入
        after = [m for r, m in rows if r > row]
        pos = base + (after[0].start() if after else len(body))
        new_row = f'<row r="{row}">{_new_cell(ref, style, value)}</row>'
        return xml[:pos] + new_row + xml[pos:], False
    m = hit[0]
    rtext = m.group(0)
    open_end = rtext.index(">") + 1
    rowtag = rtext[:open_end]
    if rowtag.endswith("/>"):
        rowtag, inner, close = rowtag[:-2] + ">", "", "</row>"
    else:
        inner, close = rtext[open_end:-len("</row>")], "</row>"
    spans = _attr(rowtag, "spans")
    if spans and ":" in spans:
        a, b = (int(x) for x in spans.split(":"))
        if not a <= ci <= b:
            rowtag = rowtag.replace(f'spans="{spans}"', f'spans="{min(a, ci)}:{max(b, ci)}"')
    cells = list(CELL_RE.finditer(inner))
    dropped_formula = False
    for c in cells:
        ctag = c.group(0)
        if _attr(ctag[:ctag.index(">") + 1], "r") != ref:
            continue
        if "<f" in ctag:
            f = re.search(r"<f\b[^>]*>", ctag)
            if f and _attr(f.group(0), "ref") and _attr(f.group(0), "t") == "shared":
                raise Refused(f"{ref}: 共有数式の親 cell (上書きすると子の数式が壊れる)")
            if not allow_formula:
                raise Refused(f"{ref}: 数式 cell (上書きするなら --allow-formula)")
            dropped_formula = True
        st = _attr(ctag[:ctag.index(">") + 1], "s")
        inner = inner[:c.start()] + _new_cell(ref, st, value) + inner[c.end():]
        break
    else:  # cell が無い → 列順に挿入
        later = [c for c in cells if col_index(REF_RE.match(_attr(c.group(0), "r")).group(1)) > ci]
        pos = later[0].start() if later else len(inner)
        inner = inner[:pos] + _new_cell(ref, style, value) + inner[pos:]
    new_row = rowtag + inner + close
    s0 = base + m.start()
    return xml[:s0] + new_row + xml[base + m.end():], dropped_formula


def _widen_dimension(xml: str, refs) -> str:
    m = re.search(r'<dimension ref="([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?"\s*/>', xml)
    if not m:
        return xml
    c0, r0 = col_index(m.group(1)), int(m.group(2))
    c1, r1 = (col_index(m.group(3)), int(m.group(4))) if m.group(3) else (c0, r0)
    for ref in refs:
        c, r = REF_RE.match(ref).groups()
        c, r = col_index(c), int(r)
        c0, r0, c1, r1 = min(c0, c), min(r0, r), max(c1, c), max(r1, r)
    return xml[:m.start()] + f'<dimension ref="{col_letters(c0)}{r0}:{col_letters(c1)}{r1}"/>' + xml[m.end():]


def read_cell(xml: str, ref: str):
    for c in CELL_RE.finditer(xml):
        t = c.group(0)
        if _attr(t[:t.index(">") + 1], "r") != ref:
            continue
        if 't="inlineStr"' in t:
            return unescape("".join(re.findall(r"<t\b[^>]*>(.*?)</t>", t, re.S)))
        v = re.search(r"<v>(.*?)</v>", t)
        return float(v.group(1)) if v else None
    return None


def set_cells(book, sheet: str, cells: dict, out, allow_formula=False, integrity=True) -> list:
    book, out = str(book), str(out)
    with zipfile.ZipFile(book) as zin:
        part, _sid = sheet_part(zin, sheet)
        xml = zin.read(part).decode("utf-8")
        dropped = False
        for ref, val in cells.items():
            if not REF_RE.match(ref):
                raise SystemExit(f"xlsx-zip-set-cells: cell 参照が不正: {ref!r}")
            xml, d = set_cell(xml, ref, val, allow_formula)
            dropped |= d
        xml = _widen_dimension(xml, cells)
        drop = set()
        extra = {}
        if dropped and "xl/calcChain.xml" in zin.namelist():
            drop.add("xl/calcChain.xml")
            rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            extra["xl/_rels/workbook.xml.rels"] = re.sub(r"<Relationship\b[^>]*calcChain[^>]*/>", "", rels).encode("utf-8")
            ct = zin.read("[Content_Types].xml").decode("utf-8")
            extra["[Content_Types].xml"] = re.sub(r'<Override\b[^>]*PartName="/xl/calcChain.xml"[^>]*/>', "", ct).encode("utf-8")
        fd, tmp = tempfile.mkstemp(suffix=Path(out).suffix, dir=str(Path(out).resolve().parent))
        os.close(fd)
        with zipfile.ZipFile(tmp, "w") as zout:
            for info in zin.infolist():
                if info.filename in drop:
                    continue
                if info.filename == part:
                    data = xml.encode("utf-8")
                else:
                    data = extra.get(info.filename, zin.read(info.filename))
                zout.writestr(info, data, compress_type=info.compress_type)
        # 検査: 値の読み戻し + 他 member の CRC 一致 + integrity gate
        problems = []
        with zipfile.ZipFile(tmp) as zchk:
            got = zchk.read(part).decode("utf-8")
            for ref, val in cells.items():
                want = None if val is None else (float(val) if isinstance(val, (int, float)) and not isinstance(val, bool) else str(val))
                if read_cell(got, ref) != want:
                    problems.append(f"{ref}: 読み戻し {read_cell(got, ref)!r} ≠ {want!r}")
            before = {i.filename: i.CRC for i in zin.infolist() if i.filename not in drop | set(extra) | {part}}
            after = {i.filename: i.CRC for i in zchk.infolist()}
            for name, crc in before.items():
                if after.get(name) != crc:
                    problems.append(f"{name}: 触っていない member が変わった")
    if integrity and INTEGRITY.exists() and not problems:
        r = subprocess.run([sys.executable, str(INTEGRITY), tmp], capture_output=True, text=True)
        if r.returncode != 0:
            problems.append("check-xlsx-integrity: " + (r.stdout + r.stderr).strip()[-600:])
    if problems:
        os.remove(tmp)
        raise Refused("; ".join(problems))
    os.replace(tmp, out)
    notes = [f"{ref} = {val!r}" for ref, val in cells.items()]
    if drop:
        notes.append("数式を消したので calcChain.xml を外した (Excel が作り直す)")
    return notes


def _parse_set(items):
    cells = {}
    for it in items:
        if "=" not in it:
            raise SystemExit(f"--set は REF=値: {it!r}")
        ref, val = it.split("=", 1)
        if val == "":
            cells[ref] = None
        elif val.startswith("##"):
            cells[ref] = val[1:]
        elif val.startswith("#"):
            cells[ref] = float(val[1:]) if "." in val else int(val[1:])
        else:
            cells[ref] = val
    return cells


def _selftest():
    import openpyxl
    d = tempfile.mkdtemp()
    src = os.path.join(d, "in.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表 & 一覧"
    ws["A1"] = "見出し"
    ws["C3"] = 45672
    ws["C3"].number_format = "0"
    ws["C4"] = 1
    ws["C4"].number_format = "0.00"
    ws["B5"] = "=1+1"
    ws["A7"] = "行7"
    ws["F7"] = "右"
    wb.save(src)
    out = os.path.join(d, "out.xlsx")
    n_ok = 0
    set_cells(src, "表 & 一覧", {"C3": "2025/01/15", "D7": "新", "A10": "新しい行", "E3": 12.5, "A1": " 前後に空白 "}, out)
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2["表 & 一覧"]
    assert ws2["C3"].value == "2025/01/15" and ws2["C3"].number_format == "0", (ws2["C3"].value, ws2["C3"].number_format)
    assert ws2["D7"].value == "新" and ws2["A10"].value == "新しい行" and ws2["E3"].value == 12.5
    assert ws2["A1"].value == " 前後に空白 " and ws2["F7"].value == "右" and ws2["A7"].value == "行7"
    n_ok += 1
    with zipfile.ZipFile(out) as z:
        x = z.read(sheet_part(z, "表 & 一覧")[0]).decode()
    row7 = re.search(r'<row r="7".*?</row>', x, re.S).group(0)
    assert row7.index('r="A7"') < row7.index('r="D7"') < row7.index('r="F7"'), row7  # 列順に挿入
    assert x.index('<row r="7"') < x.index('<row r="10"'), "行順に挿入"
    assert re.search(r'<dimension ref="A1:F10"/>', x), re.search(r"<dimension[^>]*>", x).group(0)
    n_ok += 1
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        part = sheet_part(a, "表 & 一覧")[0]
        for i in a.infolist():
            if i.filename != part:
                assert a.read(i.filename) == b.read(i.filename), i.filename  # 他 member は byte 同一
    n_ok += 1
    try:
        set_cells(src, "表 & 一覧", {"B5": "x"}, os.path.join(d, "f.xlsx"))
        raise AssertionError("数式 cell を黙って上書きした")
    except Refused:
        n_ok += 1
    set_cells(src, "表 & 一覧", {"B5": "x", "C4": None}, os.path.join(d, "f.xlsx"), allow_formula=True)
    ws3 = openpyxl.load_workbook(os.path.join(d, "f.xlsx"))["表 & 一覧"]
    assert ws3["B5"].value == "x" and ws3["C4"].value is None and ws3["C4"].number_format == "0.00"
    n_ok += 1
    assert _parse_set(["A1=#3", "A2=#2.5", "A3=##x", "A4=", "A5=文"]) == {"A1": 3, "A2": 2.5, "A3": "#x", "A4": None, "A5": "文"}
    n_ok += 1
    print(f"xlsx-zip-set-cells selftest: {n_ok}/6 PASS")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("book", nargs="?")
    ap.add_argument("--sheet")
    ap.add_argument("--set", action="append", default=[], metavar="REF=値")
    ap.add_argument("--spec", help='JSON {"sheet": ..., "cells": {...}}')
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--out")
    g.add_argument("--in-place", action="store_true")
    ap.add_argument("--allow-formula", action="store_true")
    ap.add_argument("--no-integrity", action="store_true", help="check-xlsx-integrity.py を回さない (selftest 等)")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        _selftest()
        return 0
    if not a.book or not os.path.exists(a.book):
        ap.error("BOOK が無い")
    if not (a.out or a.in_place):
        ap.error("--out か --in-place を指定 (上書きは明示)")
    sheet, cells = a.sheet, _parse_set(a.set)
    if a.spec:
        spec = json.load(open(a.spec, encoding="utf-8"))
        sheet = spec.get("sheet", sheet)
        cells.update(spec["cells"])
    if not sheet or not cells:
        ap.error("--sheet と書く cell (--set / --spec) が要る")
    try:
        notes = set_cells(a.book, sheet, cells, a.out or a.book, a.allow_formula, not a.no_integrity)
    except Refused as e:
        print(f"xlsx-zip-set-cells: ✗ {e}", file=sys.stderr)
        return 1
    for n in notes:
        print("  ✓", n)
    print(f"  → {a.out or a.book} (他 member は byte 同一、 integrity {'skip' if a.no_integrity else 'PASS'})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
