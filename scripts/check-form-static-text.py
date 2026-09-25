#!/usr/bin/env python3
"""check-form-static-text.py — 様式の雛形が紙に出すもの (図形の字・cell の見出し・字の無い図形 = 画像) が出力 PDF に在るかを、 雛形と素刷りを参照に見る。

WHY (office-automation.md#openpyxl-destroys-drawings、 form-case-pipeline.md#fidelity):
  様式の見出しの一部は cell でなく**図形** (DrawingML の textbox) に刷られている。 openpyxl で保存した workbook
  (記入した案件の xlsx でも、 PDF にするための使い捨ての temp でも) はその図形を落とし、 PDF から見出しが消える。
  既存の検査はどれも**書いたもの** (記入値の照合・字の切れ・記入要領の残置・label の上書き) を見ており、
  **雛形が元から紙に出すもの**を出力で確かめる段が無かった。 目視も「在るもの」 に目が行き、 無いものには
  比べる参照が無いと気づけない (実測: 見出しが消えたままの提出物が続き、 gate は全部通っていた)。
  この script はその参照を雛形 (と、 雛形を道具を通さず app で刷った素刷り) から機械で作り、 出力 PDF と突き合わせる。
  参照を「前の出力」 にしない (前の欠けを正しいと認定する = docs/convention-design-principles.md#parity-against-own-output)。

WHAT (xlsx の雛形):
  - 雛形 xlsx の各 sheet の図形 (twoCellAnchor / oneCellAnchor / absoluteAnchor、 mc:AlternateContent の中も) から
    字のある textbox を集める (hidden の図形・--drop で名指しした図形・枠を切る設定で枠が 1 字より狭い図形 =
    雛形を刷っても字が出ないもの、 は除く)。 段落ごとに照合する = 🔴 で exit 1
  - --filled 案件.xlsx (記入済み、 または体裁を当てた temp): 印刷範囲の中の cell の見出し (label) のうち、 **案件で
    書き換えていない cell** (雛形と同じ字。 記入した cell は除く = 記入欄は雛形との差から導く、 人が番地を書かない) の字も
    照合する。 案件で隠した・潰した行 (hidden / 高さ 3.5pt 以下) の label は数えない。 欠け = 「label」 として数え、
    --strict-labels でなければ exit code に入れない (誤検出の型を実物で測る段階 = 呼び元が warn で出す)。
    記入値 (雛形と違う cell の値、 数式は計算済みの値) も消し込み、 残った字 = 「増えた字」 (extra) を情報として出す
  - --blank 素刷り.pdf (雛形を道具を通さず app で刷った PDF): 対象の頁を label で選び、 **画像の数** (Excel は form control の
    checkbox を PDF に画像として描く = 実測) と線・矩形の数を比べる。 画像が減っていれば「image」 として数える
    (--strict-images でなければ exit code に入れない)。 線の減少は情報 (行を潰す体裁で減るので止めない)
  - 対象 = --target の sheet (と印刷範囲)。 範囲を書かなければ雛形の印刷範囲、 それも無ければ sheet 全体。
    範囲の外に anchor がある図形 (記入例・作成上の注意など横に置かれた注記) は数えない
  - 対象ごとに PDF の頁を 1 つ選ぶ = 雛形のその範囲の cell の字 (label) が一番多く出ている頁
  - 照合は NFKC + 空白除去 + 私用領域 (Wingdings 等) の記号を除き、 CJK 部首の互換字 (「⻑」 U+2ED1 等 = PDF の text 層が
    返す字、 NFKC で戻らない) は 1 字の揺れとして受ける。 長い字から順に 1 回ずつ消し込むので、
    様式番号「甲-外部資金様式③-1」 が在っても区分の枠「外部資金」 が無ければ無いと言う。 枠で字が切れた
    (「⑥-2」 が「⑥-」) も無いと言う

WHAT (docx の雛形): 雛形.docx を渡すと、 textbox (w:txbxContent)・header・footer の段落の字を「図形の字」 として (exit 1)、
  本文と表の段落の字を「label」 として照合する (--filled 案件.docx があれば、 案件で書き換えた段落は除く)。 頁は選ばず
  文書全体の text で消し込む。

射程外 (= この検査が黙っても在るとは限らない):
  - 図形の位置のずれ・重なり・字が潰れる二重刷り = 在るかしか見ない (位置は段階 2)
  - raster (画像だけ) の PDF = text 層が無いので ⚪
  - 素刷りを渡さなければ字の無い図形 (選択の丸・線・checkbox の箱) は数だけ出す

使い方:
  check-form-static-text.py 雛形.xlsx 出力.pdf                                  # 全 sheet、 雛形の印刷範囲
  check-form-static-text.py 雛形.xlsx 出力.pdf --target '出張願!A1:AH60' --target '報告書'
  check-form-static-text.py 雛形.xlsx 出力.pdf --target 請求書 --drop '請求書!楕円 2'
  check-form-static-text.py 雛形.xlsx 出力.pdf --filled 案件.xlsx --blank 素刷り.pdf   # label と画像の数も
  check-form-static-text.py 雛形.docx 出力.pdf --filled 案件.docx
  check-form-static-text.py ... --json      # 機械向け
  check-form-static-text.py ... --strict    # 対象の頁が見つからない (⚪) も exit 1
  check-form-static-text.py --selftest
exit: 0 = 全部在る (⚪ は --strict でなければ 0) / 1 = 無い図形の字がある (--strict-labels / --strict-images で label・画像も) /
      2 = 使い方・読めない file
"""
from __future__ import annotations

import argparse
import json
import posixpath
import re
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}
REL_DRAWING = NS["r"] + "/drawing"
ANCHORS = ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor")
COLLAPSED_PT = 3.5          # これ以下の行高は「潰した行」 (formcase の日程表の正規化 = 3pt)
_PUA = re.compile(r"[-]")
_RADICAL = "⺀-⿟"  # CJK 部首補助・康熙部首 (PDF の text 層が返す互換字。 NFKC で戻らないものがある)


def norm(s) -> str:
    return _PUA.sub("", re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(s or ""))))


def _find(want: str, text: str):
    """want が text に在れば (start, end)。 完全一致 → 無ければ CJK の字を部首の互換字でも受ける正規表現で。"""
    i = text.find(want)
    if i >= 0:
        return i, i + len(want)
    if not re.search(r"[㐀-鿿]", want):
        return None
    pat = "".join(f"(?:{re.escape(c)}|[{_RADICAL}])" if "㐀" <= c <= "鿿" else re.escape(c) for c in want)
    m = re.search(pat, text)
    return (m.start(), m.end()) if m else None


def _consume(want: str, text: str, loose: bool = False):
    """want を text から 1 回消す。 (消した後の text, 見つかったか)。
    loose = 長い字 (12 字以上) は 6 字ずつの断片が全部在れば「在る」 とする (消さない) = 折り返した cell の 2 行目の
    間に隣の cell の字が挟まる PDF の text 順 (実測) を、 無いと言わないため。"""
    sp = _find(want, text)
    if sp is not None:
        return text[:sp[0]] + "\0" * (sp[1] - sp[0]) + text[sp[1]:], True
    if loose and len(want) >= 12:
        grams = [want[i:i + 6] for i in range(0, len(want) - 5, 3)]
        if grams and all(_find(g, text) is not None for g in grams):
            return text, True
    return text, False


# ---------------------------------------------------------------------------
# 雛形の図形 (xlsx)
# ---------------------------------------------------------------------------
def _resolve(base, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base), target))


def _rels(z, part):
    name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if name not in z.namelist():
        return []
    return [(r.get("Id"), r.get("Type"), r.get("Target")) for r in ET.fromstring(z.read(name))]


def sheet_parts(z) -> dict:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    by_id = {rid: _resolve("xl/workbook.xml", t) for rid, _ty, t in _rels(z, "xl/workbook.xml")}
    out = {}
    for s in wb.iter(f"{{{NS['m']}}}sheet"):
        rid = s.get(f"{{{NS['r']}}}id")
        if rid in by_id:
            out[s.get("name")] = by_id[rid]
    return out


def _anchor_elems(root):
    """wsDr 直下の anchor (mc:AlternateContent の中は mc:Choice の側)。"""
    for ch in root:
        tag = ch.tag.split("}")[-1]
        if tag in ANCHORS:
            yield ch
        elif tag == "AlternateContent":
            branch = ch.find("mc:Choice", NS)
            if branch is None:
                branch = ch.find("mc:Fallback", NS)
            for sub in (branch if branch is not None else []):
                if sub.tag.split("}")[-1] in ANCHORS:
                    yield sub


def _cell_of(anchor, which):
    el = anchor.find(f"xdr:{which}", NS)
    if el is None:
        return None
    return int(el.find("xdr:row", NS).text) + 1, int(el.find("xdr:col", NS).text) + 1


def _offs(anchor, which):
    el = anchor.find(f"xdr:{which}", NS)
    if el is None:
        return None
    return tuple(int(el.find(f"xdr:{k}", NS).text) for k in ("col", "colOff", "row", "rowOff"))


EMU_PT = 12700


def _col_pt(ws, ci):
    """列 ci (1 始まり) の幅 (pt、 近似)。 Excel の列幅は既定 font の数字幅 (≒ 7px) 単位 + 余白 5px。"""
    from openpyxl.utils import get_column_letter

    cd = ws.column_dimensions.get(get_column_letter(ci))
    if cd is not None and cd.hidden:
        return 0.0
    w = cd.width if cd is not None and cd.width else (ws.sheet_format.defaultColWidth or 8.43)
    return (float(w) * 7 + 5) * 0.75


def _row_pt(ws, ri):
    rd = ws.row_dimensions.get(ri)
    if rd is not None and rd.hidden:
        return 0.0
    if rd is not None and rd.height is not None:
        return float(rd.height)
    return float(ws.sheet_format.defaultRowHeight or 15)


def _box_pt(ws, anchor):
    """図形の外枠 (幅, 高さ) の pt。 one/absolute は ext から、 twoCellAnchor は列幅・行高から。 求まらなければ None。"""
    ext = anchor.find("xdr:ext", NS)
    if ext is not None:
        return int(ext.get("cx")) / EMU_PT, int(ext.get("cy")) / EMU_PT
    f, t = _offs(anchor, "from"), _offs(anchor, "to")
    if f is None or t is None or ws is None:
        return None
    w = sum(_col_pt(ws, c + 1) for c in range(f[0], t[0])) - f[1] / EMU_PT + t[1] / EMU_PT
    h = sum(_row_pt(ws, r + 1) for r in range(f[2], t[2])) - f[3] / EMU_PT + t[3] / EMU_PT
    return max(0.0, w), max(0.0, h)


def _invisible(sp, box) -> bool:
    """字が紙に出ない図形 = 枠の外を切る (clip) 設定で、 枠の内側が字 1 つ分より狭い (実測: 様式の旧版の見出しが
    幅 1 字未満の列に残り、 雛形を Excel で刷っても字が出ない)。 縦書きも幅と高さの両方に 1 字要る。"""
    if box is None:
        return False
    bp = sp.find("xdr:txBody/a:bodyPr", NS)
    if bp is None:
        return False
    sizes = [int(r.get("sz")) / 100 for r in sp.iter(f"{{{NS['a']}}}rPr") if r.get("sz")]
    em = max(sizes) if sizes else 11.0
    ins = {k: int(bp.get(k, d)) / EMU_PT for k, d in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720))}
    narrow_w = bp.get("horzOverflow") == "clip" and box[0] - ins["lIns"] - ins["rIns"] < 0.6 * em
    narrow_h = bp.get("vertOverflow") == "clip" and box[1] - ins["tIns"] - ins["bIns"] < 0.6 * em
    return narrow_w or narrow_h


def shapes(template, book=None) -> dict:
    """sheet 名 → [{name, hidden, invisible, paras, from, to}] (字の無い図形は paras=[])。 book = 雛形の openpyxl
    workbook (列幅・行高から図形の大きさを出す = 字が出ない狭い図形を除くため。 無ければ大きさは見ない)。"""
    out = {}
    with zipfile.ZipFile(template) as z:
        for sheet, part in sheet_parts(z).items():
            ws = book[sheet] if book is not None and sheet in book.sheetnames else None
            items = []
            for _rid, typ, tgt in _rels(z, part):
                if typ != REL_DRAWING:
                    continue
                root = ET.fromstring(z.read(_resolve(part, tgt)))
                for a in _anchor_elems(root):
                    frm, to = _cell_of(a, "from"), _cell_of(a, "to")
                    box = _box_pt(ws, a)
                    for sp in a.iter(f"{{{NS['xdr']}}}sp"):
                        nv = sp.find("xdr:nvSpPr/xdr:cNvPr", NS)
                        name = nv.get("name") if nv is not None else "?"
                        hidden = nv is not None and nv.get("hidden") in ("1", "true")
                        paras = []
                        body = sp.find("xdr:txBody", NS)
                        if body is not None:
                            for p in body.findall("a:p", NS):
                                t = norm("".join(x.text or "" for x in p.iter(f"{{{NS['a']}}}t")))
                                if t:
                                    paras.append(t)
                        items.append({"name": name, "hidden": hidden, "invisible": _invisible(sp, box),
                                      "paras": paras, "from": frm, "to": to})
            out[sheet] = items
    return out


# ---------------------------------------------------------------------------
# 範囲・cell の label
# ---------------------------------------------------------------------------
def _bounds(rng: str):
    from openpyxl.utils.cell import range_boundaries

    rng = rng.split("!")[-1].replace("$", "").strip()
    c0, r0, c1, r1 = range_boundaries(rng)
    return c0, r0, c1, r1


def _find_sheet(names, want):
    for n in names:
        if n.strip() == want.strip():
            return n
    return None


def _book(path, data_only=False):
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=data_only)


def _template_book(template):
    return _book(template, data_only=False)


def _ranges_of(ws, rng):
    if rng:
        return [rng]
    pa = ws.print_area
    if pa:
        parts = pa if isinstance(pa, (list, tuple)) else str(pa).split(",")
        return [p.split("!")[-1].replace("$", "") for p in parts if p.strip()]
    return [ws.dimensions]


def _labels(ws, rng) -> list:
    c0, r0, c1, r1 = _bounds(rng)
    out = []
    for row in ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            v = c.value
            if isinstance(v, str) and not v.startswith("="):
                t = norm(v)
                if len(t) >= 2:
                    out.append(t)
    return out


def _row_collapsed(ws, r) -> bool:
    rd = ws.row_dimensions.get(r)
    return rd is not None and (bool(rd.hidden) or (rd.height is not None and float(rd.height) <= COLLAPSED_PT))


def _col_hidden(ws, ci) -> bool:
    from openpyxl.utils import get_column_letter

    cd = ws.column_dimensions.get(get_column_letter(ci))
    return cd is not None and bool(cd.hidden)


def _untouched_labels(wt, wf, rng) -> list:
    """案件で書き換えていない cell の見出し = [(coord, 字)]。 wt = 雛形の sheet、 wf = 案件 (or 体裁を当てた temp) の同名 sheet。
    記入した cell (値が違う)・数式・潰した/隠した行・隠した列・2 字未満は除く。"""
    c0, r0, c1, r1 = _bounds(rng)
    out = []
    for row in wt.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            v = c.value
            if not isinstance(v, str) or v.startswith("="):
                continue
            t = norm(v)
            if len(t) < 2:
                continue
            fv = wf[c.coordinate].value
            if not isinstance(fv, str) or norm(fv) != t:
                continue                                    # 記入した cell (= 雛形との差) は label でない
            if _row_collapsed(wf, c.row) or _col_hidden(wf, c.column):
                continue
            # 改行を含む label (「出発␊␊到着」) は PDF の text で行の間に隣の cell の字が挟まる = 行ごとに照合する (実測)
            pieces = [norm(x) for x in v.split("\n")] if "\n" in v else [t]
            for pc in pieces:
                if len(pc) >= 2:
                    out.append((c.coordinate, pc))
    return out


def _filled_values(wt, wfv, wf, rng) -> list:
    """消し込んでよい字 = 記入値 (雛形と値が違う cell) と導出値 (数式 cell の計算済みの値。 雛形と同じ数式でも、 app が
    開いた時に再計算するので雛形の cache と違う = `=TODAY()`)。 wfv = 案件を data_only で読んだ sheet。"""
    c0, r0, c1, r1 = _bounds(rng)
    out = []
    for row in wf.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            tv = wt[c.coordinate].value
            fv = c.value
            if fv is None:
                continue
            is_formula = isinstance(fv, str) and str(fv).startswith("=")
            if fv == tv and not is_formula:
                continue
            val = wfv[c.coordinate].value if is_formula else fv
            if val is None:
                continue
            t = norm(val if isinstance(val, str) else (val.strftime("%Y-%m-%d") if hasattr(val, "strftime") else val))
            if len(t) >= 1:
                out.append(t)
    return out


def _single_chars(wt, rng) -> set:
    """1 字の cell (「円」「日」「印」「□」 等) = 増えた字の判定から除く字。"""
    c0, r0, c1, r1 = _bounds(rng)
    out = set()
    for row in wt.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                t = norm(c.value)
                if len(t) == 1:
                    out.add(t)
    return out


_DIGITY = re.compile(r"[\d/:.,()~\-年月日時分曜()]")


def _extra_fragments(text2: str, singles: set) -> list:
    """消し込みの残り = 増えた字。 1 字の cell の字と、 数字・日付の記号が 6 割以上の断片 (表示書式で形の変わる日時) は除く。"""
    left = []
    for x in re.split(r"\0+", text2):
        x = "".join(ch for ch in x if ch not in singles)
        if len(x) < 2 or not re.search(r"[^\d\W_]", x):
            continue
        if len(_DIGITY.findall(x)) >= 0.6 * len(x):
            continue
        left.append(x)
    return left


def _inside(cell, rng) -> bool:
    if cell is None:
        return False
    c0, r0, c1, r1 = _bounds(rng)
    r, c = cell
    return r0 <= r <= r1 and c0 <= c <= c1


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def page_texts(pdf) -> list:
    import fitz

    with fitz.open(pdf) as d:
        return [norm(p.get_text()) for p in d]


def page_visuals(pdf) -> list:
    """頁ごとの {images, drawings} (Excel は checkbox の箱を画像として描く = 画像の数で字の無い図形の消失を見る)。"""
    import fitz

    out = []
    with fitz.open(pdf) as d:
        for p in d:
            # get_images = 頁が参照する画像 (xref) の数。 Excel は checkbox の箱を 1 個ずつ別の画像で描く (実測: 9 個 = 9)。
            # get_image_info (描かれた回数) は点線の pattern の tile まで数えて体裁で増減する (実測: 日程表 11 → 0) = 使わない
            out.append({"images": len(p.get_images()), "drawings": len(p.get_drawings())})
    return out


def page_layout(pdf) -> list:
    """段階 2 (D4、 2026-09-25) の材料 = 頁ごとの {images: [(xref, bbox)], hlines: [(y, x0, x1)], words: [(字, bbox)]}。
    画像 = 描かれた位置つき (pattern の tile = 同じ xref が何度も描かれるものは除く)。 線 = 線と細い矩形 (罫線) の水平なもの。"""
    import fitz

    out = []
    with fitz.open(pdf) as d:
        for p in d:
            info = p.get_image_info(xrefs=True)
            per = {}
            for it in info:
                per.setdefault(it.get("xref"), []).append(it["bbox"])
            images = [(x, b) for x, bs in per.items() if x and len(bs) <= 4 for b in bs]
            hl = {}
            for dr in p.get_drawings():
                for it in dr["items"]:
                    if it[0] == "l":
                        (x0, y0), (x1, y1) = (it[1].x, it[1].y), (it[2].x, it[2].y)
                        if abs(y0 - y1) <= 0.6 and abs(x1 - x0) >= 5:
                            hl.setdefault(round((y0 + y1) / 2, 1), []).append((min(x0, x1), max(x0, x1)))
                    elif it[0] == "re":
                        r = it[1]
                        if r.height <= 1.5 and r.width >= 5:
                            hl.setdefault(round((r.y0 + r.y1) / 2, 1), []).append((r.x0, r.x1))
            hlines = []
            for y, segs in hl.items():
                segs.sort()
                cur = list(segs[0])
                for a, b in segs[1:]:
                    if a <= cur[1] + 3:
                        cur[1] = max(cur[1], b)
                    else:
                        hlines.append((y, cur[0], cur[1]))
                        cur = [a, b]
                hlines.append((y, cur[0], cur[1]))
            words = [(norm(w[4]), fitz.Rect(w[:4])) for w in p.get_text("words")]
            out.append({"images": images, "hlines": hlines, "words": words, "height": p.rect.height, "width": p.rect.width})
    return out


def _label_points(layout: dict, labels) -> dict:
    """頁の中で 1 回だけ出る label → その中心 (x, y)。 位置の対応づけの基準点。"""
    pts = {}
    for t in set(labels):
        if len(t) < 2:
            continue
        hits = [r for w, r in layout["words"] if w == t]
        if len(hits) == 1:
            r = hits[0]
            pts[t] = ((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
    return pts


def layout_diff(blank: dict, out: dict, labels) -> dict:
    """素刷り (blank) と出力 (out) を label の位置の組で対応づけ、 画像 (checkbox の箱・図) と水平の罫線の位置を比べる。
    x は 1 次 (紙 1 枚に収める縮尺)、 y は label の組の間の区分線形 (行を伸ばした分だけ下がる)。
    返り値 = {pairs, images_missing: [bbox], images_moved: [(bbox, Δx, Δy)], images_added: n, hlines_missing: [(y, x0, x1)],
    double: [字]} (pairs < 4 なら対応づけできず {"pairs": n} だけ)。"""
    bp, op = _label_points(blank, labels), _label_points(out, labels)
    common = sorted(set(bp) & set(op), key=lambda t: bp[t][1])
    if len(common) < 4:
        return {"pairs": len(common)}
    xs = [(bp[t][0], op[t][0]) for t in common]
    n = len(xs)
    mx, my = sum(a for a, _ in xs) / n, sum(b for _, b in xs) / n
    var = sum((a - mx) ** 2 for a, _ in xs)
    ax = (sum((a - mx) * (b - my) for a, b in xs) / var) if var > 1e-6 else 1.0
    bx = my - ax * mx
    ys = sorted({(round(bp[t][1], 1), op[t][1]) for t in common})

    def fy(y):
        if y <= ys[0][0]:
            (y0, o0), (y1, o1) = ys[0], ys[1]
        elif y >= ys[-1][0]:
            (y0, o0), (y1, o1) = ys[-2], ys[-1]
        else:
            k = max(i for i in range(len(ys) - 1) if ys[i][0] <= y)
            (y0, o0), (y1, o1) = ys[k], ys[k + 1]
        return o0 + (o1 - o0) * (y - y0) / (y1 - y0) if y1 != y0 else o0

    tol = max(4.0, 0.006 * out["height"])
    used = set()
    missing, moved = [], []
    for _x, b in blank["images"]:
        cx, cy = ax * (b[0] + b[2]) / 2 + bx, fy((b[1] + b[3]) / 2)
        best = None
        for i, (_ox, ob) in enumerate(out["images"]):
            if i in used:
                continue
            d = ((ob[0] + ob[2]) / 2 - cx, (ob[1] + ob[3]) / 2 - cy)
            dist = (d[0] ** 2 + d[1] ** 2) ** 0.5
            if best is None or dist < best[0]:
                best = (dist, i, d)
        if best is not None and best[0] <= tol:
            used.add(best[1])
        elif best is not None and best[0] <= 8 * tol:
            used.add(best[1])
            moved.append(([round(v, 1) for v in b], round(best[2][0], 1), round(best[2][1], 1)))
        else:
            missing.append([round(v, 1) for v in b])
    added = len(out["images"]) - len(used)
    # 罫線は label の組に挟まれた帯 (隣の組まで BAND pt 以内) だけ比べる = 帯の外 (label の無い表 = 日程表の block を
    # 値で潰す様式) は行の高さが値で決まり、 素刷りから写像できない = 圏外 (数だけ出す)
    BAND = 250.0
    hmiss, checked, skipped = [], 0, 0
    for y, x0, x1 in blank["hlines"]:
        if not (ys[0][0] - tol <= y <= ys[-1][0] + tol):
            skipped += 1
            continue
        k = max((i for i in range(len(ys) - 1) if ys[i][0] <= y), default=0)
        if ys[min(k + 1, len(ys) - 1)][0] - ys[k][0] > BAND:
            skipped += 1
            continue
        checked += 1
        ey, ex0, ex1 = fy(y), ax * x0 + bx, ax * x1 + bx
        ok = any(abs(oy - ey) <= tol and min(ox1, ex1) - max(ox0, ex0) >= 0.5 * (ex1 - ex0) for oy, ox0, ox1 in out["hlines"])
        if not ok:
            hmiss.append((round(y, 1), round(x0, 1), round(x1, 1)))
    double = []
    for t in common:
        hits = [r for w, r in out["words"] if w == t]
        if len(hits) >= 2:
            double.append(t)
    return {"pairs": len(common), "images_missing": missing, "images_moved": moved, "images_added": max(0, added),
            "hlines_missing": hmiss, "hlines_checked": checked, "hlines_skipped": skipped, "double": double}


def _score(labels, text) -> int:
    return sum(1 for t in set(labels) if t in text)


def _assign_pages(jobs, texts) -> None:
    """頁の割り当て = label の出た数が多い順に、 1 頁 1 対象 (同じ sheet の別範囲も別の頁)。 jobs[i]["page"] に書く。"""
    pairs = sorted(((_score(j["labels"], t), -i, k) for i, j in enumerate(jobs) for k, t in enumerate(texts)),
                   reverse=True)
    taken_j, taken_p = set(), set()
    for sc, negi, k in pairs:
        i = -negi
        need = max(2, int(0.3 * len(set(jobs[i]["labels"]))))
        if i in taken_j or k in taken_p or sc < need:
            continue
        jobs[i]["page"], jobs[i]["score"] = k + 1, sc
        taken_j.add(i)
        taken_p.add(k)


# ---------------------------------------------------------------------------
# 照合 (xlsx)
# ---------------------------------------------------------------------------
def check(template, pdf, targets=None, drop=None, filled=None, blank=None) -> dict:
    """targets = [(sheet, range or None)] (None = 雛形の全 sheet)。 drop = {(sheet, 図形名)}。
    filled = 案件 (or 体裁を当てた temp) の xlsx = cell の label と記入値を見る。 blank = 素刷りの PDF = 画像・線の数を比べる。"""
    drop = {(s.strip(), n) for s, n in (drop or set())}
    wb = _template_book(template)
    shp = shapes(template, wb)
    texts = page_texts(pdf)
    wf = _book(filled) if filled else None
    wfv = _book(filled, data_only=True) if filled else None
    if targets is None:
        targets = [(n, None) for n in wb.sheetnames]
    jobs = []
    for sheet, rng in targets:
        name = _find_sheet(wb.sheetnames, sheet)
        if name is None:
            raise ValueError(f"雛形に sheet {sheet!r} が無い")
        ws = wb[name]
        fname = _find_sheet(wf.sheetnames, sheet) if wf is not None else None
        for r in _ranges_of(ws, rng):
            labels = _labels(ws, r)
            items = [s for s in shp.get(name, []) if _inside(s["from"], r)]
            outside = sum(1 for s in shp.get(name, []) if not _inside(s["from"], r))
            job = {"sheet": name, "range": r, "labels": labels, "items": items, "outside": outside,
                   "untouched": [], "values": [], "singles": _single_chars(ws, r)}
            if fname is not None:
                job["untouched"] = _untouched_labels(ws, wf[fname], r)
                job["values"] = _filled_values(ws, wfv[fname], wf[fname], r)
            jobs.append(job)
    _assign_pages(jobs, texts)
    bl_texts = bl_vis = None
    if blank:
        bl_texts, bl_vis = page_texts(blank), page_visuals(blank)
        bjobs = [{"labels": j["labels"]} for j in jobs]
        _assign_pages(bjobs, bl_texts)
        for j, bj in zip(jobs, bjobs):
            j["blank_page"] = bj.get("page")
    vis = page_visuals(pdf) if blank else None
    bl_lay = out_lay = None
    report = []
    for j in jobs:
        res = {"sheet": j["sheet"], "range": j["range"], "page": j.get("page"), "checked": 0, "missing": [],
               "textless": 0, "dropped": 0, "hidden": 0, "invisible": 0, "outside": j["outside"],
               "labels_checked": 0, "missing_labels": [], "extra": None, "blank": None}
        want = []
        for s in j["items"]:
            if (j["sheet"].strip(), s["name"]) in drop:
                res["dropped"] += 1
            elif s["hidden"]:
                res["hidden"] += 1
            elif s["invisible"] and s["paras"]:
                res["invisible"] += 1
            elif not s["paras"]:
                res["textless"] += 1
            else:
                want += [(p, s["name"]) for p in s["paras"]]
        res["checked"] = len(want)
        res["labels_checked"] = len(j["untouched"])
        if res["page"] is not None:
            text = texts[res["page"] - 1]
            # cell の label が図形の字を含むなら、 その label の分を先に消す (= label で図形の代わりにしない)
            for lab in sorted(set(j["labels"]), key=len, reverse=True):
                if any(p in lab for p, _n in want):
                    text, _ok = _consume(lab, text)
            # 素刷り (雛形を道具を通さず app で刷った PDF) にも無い字 = 雛形自身の欠陥 (枠が字幅に足りず末尾が消える等)。
            # 出力と素刷りの差ではないので「消えた」 に数えない (⚪ で出す)。 素刷りが無ければ従来どおり全部「消えた」
            btext = bl_texts[j["blank_page"] - 1] if blank and j.get("blank_page") else None
            res["template_defect"] = []
            for p, n in sorted(want, key=lambda x: len(x[0]), reverse=True):
                text, ok = _consume(p, text)
                if not ok:
                    if btext is not None and not _consume(p, btext)[1]:
                        res["template_defect"].append({"name": n, "text": p})
                    else:
                        res["missing"].append({"name": n, "text": p})
            if j["untouched"] or j["values"]:
                text2 = texts[res["page"] - 1]
                # label (書き換えていない cell) → 図形の字 → 記入値・導出値 の順に長い方から消し込み、 残りが「増えた字」
                # (label を先に = 図形の字を含む label 「甲-外部資金様式の説明」 を図形の字で穴あきにしない)
                for coord, lab in sorted(j["untouched"], key=lambda x: len(x[1]), reverse=True):
                    text2, ok = _consume(lab, text2, loose=True)
                    if not ok:
                        res["missing_labels"].append({"cell": coord, "text": lab})
                for p, _n in sorted(want, key=lambda x: len(x[0]), reverse=True):
                    text2, _ok = _consume(p, text2)
                for v in sorted(set(j["values"]), key=len, reverse=True):
                    while True:                                # 数式で複製された値は複数回出る
                        text2, ok = _consume(v, text2)
                        if not ok:
                            break
                res["extra"] = _extra_fragments(text2, j["singles"])
        if blank:
            bp = j.get("blank_page")
            if bp is not None and res["page"] is not None:
                b, o = bl_vis[bp - 1], vis[res["page"] - 1]
                res["blank"] = {"page": bp, "images": [b["images"], o["images"]], "drawings": [b["drawings"], o["drawings"]]}
                # 段階 2 (D4): 位置の写像 = 動いた画像・素刷りの罫線の欠け・二重刷り (warn)
                if bl_lay is None:
                    bl_lay, out_lay = page_layout(blank), page_layout(pdf)
                labels_pos = [t for _c, t in j["untouched"]] if j["untouched"] else j["labels"]
                res["blank"]["layout"] = layout_diff(bl_lay[bp - 1], out_lay[res["page"] - 1], labels_pos)
            else:
                res["blank"] = {"page": bp, "images": None, "drawings": None}
        report.append(res)
    return {"template": str(template), "pdf": str(pdf), "kind": "xlsx", "targets": report,
            "missing_total": sum(len(r["missing"]) for r in report),
            "missing_labels_total": sum(len(r["missing_labels"]) for r in report),
            "missing_images_total": sum(max(0, r["blank"]["images"][0] - r["blank"]["images"][1]) for r in report
                                        if r.get("blank") and r["blank"].get("images")),
            "unmatched": sum(1 for r in report if r["page"] is None and r["checked"])}


# ---------------------------------------------------------------------------
# 照合 (docx)
# ---------------------------------------------------------------------------
def _docx_strings(path) -> dict:
    """{"boxes": [textbox / header / footer の段落の字], "body": [本文・表の段落の字]} (NFKC、 2 字以上)。"""
    W = f"{{{NS['w']}}}"
    boxes, body = [], []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for n in sorted(names):
            if not re.match(r"^word/(document|header\d*|footer\d*)\.xml$", n):
                continue
            root = ET.fromstring(z.read(n))
            in_doc = n == "word/document.xml"
            inbox = set()
            for tb in root.iter(W + "txbxContent"):
                for p in tb.iter(W + "p"):
                    inbox.add(id(p))
                    t = norm("".join(x.text or "" for x in p.iter(W + "t")))
                    if len(t) >= 2:
                        boxes.append(t)
            for p in root.iter(W + "p"):
                if id(p) in inbox or p.find(".//" + W + "txbxContent") is not None:
                    continue                            # textbox の中の段落と、 textbox を包む段落は本文に数えない
                t = norm("".join(x.text or "" for x in p.iter(W + "t")))
                if len(t) < 2:
                    continue
                (boxes if not in_doc else body).append(t)
    return {"boxes": boxes, "body": body}


def check_docx(template, pdf, filled=None) -> dict:
    tpl = _docx_strings(template)
    text = "".join(page_texts(pdf))
    body = tpl["body"]
    if filled:
        got = set(_docx_strings(filled)["body"])
        body = [t for t in body if t in got]          # 案件で書き換えた段落 (= 記入欄) は除く
    missing, missing_labels = [], []
    for t in sorted(tpl["boxes"], key=len, reverse=True):
        text, ok = _consume(t, text)
        if not ok:
            missing.append({"name": "textbox/header/footer", "text": t})
    for t in sorted(body, key=len, reverse=True):
        text, ok = _consume(t, text)
        if not ok:
            missing_labels.append({"cell": "-", "text": t})
    rep = {"sheet": "document", "range": "-", "page": 1 if text is not None else None, "checked": len(tpl["boxes"]),
           "missing": missing, "textless": 0, "dropped": 0, "hidden": 0, "invisible": 0, "outside": 0,
           "labels_checked": len(body), "missing_labels": missing_labels, "extra": None, "blank": None}
    return {"template": str(template), "pdf": str(pdf), "kind": "docx", "targets": [rep],
            "missing_total": len(missing), "missing_labels_total": len(missing_labels), "missing_images_total": 0,
            "unmatched": 0}


# ---------------------------------------------------------------------------
# 表示
# ---------------------------------------------------------------------------
def render(rep) -> list:
    lines = []
    for r in rep["targets"]:
        where = f"{r['sheet'].strip()}!{r['range']}" if rep.get("kind") != "docx" else "docx"
        extra = []
        if r["textless"]:
            extra.append(f"字の無い図形 {r['textless']}" + (" (射程外)" if not r.get("blank") else ""))
        if r["dropped"]:
            extra.append(f"drop {r['dropped']}")
        if r["invisible"]:
            extra.append(f"枠が 1 字より狭く字の出ない図形 {r['invisible']}")
        tail = f" [{' / '.join(extra)}]" if extra else ""
        if not r["checked"] and not r["labels_checked"]:
            lines.append(f"   ・ {where}: 字のある図形なし{tail}")
        elif r["page"] is None:
            lines.append(f"   ⚪ {where}: 雛形のこの範囲の頁が PDF に見つからない = 図形の字 {r['checked']} 段落を照合できない{tail}")
        elif r["missing"]:
            ms = ", ".join(f"{m['name']}「{m['text'][:24]}」" for m in r["missing"])
            lines.append(f"   🔴 {where} (PDF {r['page']} 頁): 雛形の図形の字が無い {len(r['missing'])}/{r['checked']} — {ms}{tail}")
        elif r["checked"]:
            lines.append(f"   ✅ {where} (PDF {r['page']} 頁): 図形の字 {r['checked']} 段落 全部在る{tail}")
        if r["page"] is not None and r["labels_checked"]:
            if r["missing_labels"]:
                ms = ", ".join(f"{m['cell']}「{m['text'][:16]}」" for m in r["missing_labels"][:8])
                more = f" … 他 {len(r['missing_labels']) - 8}" if len(r["missing_labels"]) > 8 else ""
                lines.append(f"   ⚠️ {where}: 雛形の見出し (書き換えていない cell) が無い {len(r['missing_labels'])}/{r['labels_checked']} — {ms}{more}")
            else:
                lines.append(f"   ✅ {where}: 雛形の見出し {r['labels_checked']} cell 全部在る")
            if r.get("extra"):
                lines.append(f"   ⚪ {where}: 雛形にも記入値にも無い字 {len(r['extra'])} 片: "
                             + ", ".join(f"「{x[:12]}」" for x in r["extra"][:6]) + (" …" if len(r["extra"]) > 6 else ""))
        if r.get("blank"):
            b = r["blank"]
            if b.get("images") is None:
                lines.append(f"   ⚪ {where}: 素刷りの頁が見つからない = 画像・線の数を比べられない")
            else:
                bi, oi = b["images"]
                bd, od = b["drawings"]
                if oi < bi:
                    lines.append(f"   🔴 {where}: 素刷りより画像が少ない {bi} → {oi} (checkbox の箱・図が紙に無い)")
                else:
                    lines.append(f"   ✅ {where}: 画像 {oi} = 素刷り {bi}" + (f"、 線・矩形 {od} (素刷り {bd})" if od != bd else f"、 線・矩形 {od} 同じ"))
                lay = b.get("layout") or {}
                if lay.get("pairs", 0) < 4:
                    lines.append(f"   ⚪ {where}: 位置の写像は label の組が {lay.get('pairs', 0)} で足りない")
                else:
                    probs = ([f"画像が無い {len(lay['images_missing'])}"] if lay.get("images_missing") else []) \
                        + ([f"動いた画像 {len(lay['images_moved'])}"] if lay.get("images_moved") else []) \
                        + ([f"罫線が無い {len(lay['hlines_missing'])} 本 (最下 y={max(h[0] for h in lay['hlines_missing'])})"] if lay.get("hlines_missing") else []) \
                        + (["二重刷り " + ", ".join(lay["double"][:3])] if lay.get("double") else [])
                    cov = f"罫線 {lay.get('hlines_checked', 0)}/{lay.get('hlines_checked', 0) + lay.get('hlines_skipped', 0)} 本を照合"
                    lines.append(f"   {'⚠️' if probs else '✅'} {where}: 位置の写像 (label {lay['pairs']} 組、 {cov})" + (": " + " / ".join(probs) if probs else " 画像は素刷りどおり"))
    return lines


def _parse_targets(vals):
    out = []
    for v in vals or []:
        if "!" in v:
            s, r = v.rsplit("!", 1)
            out.append((s, r))
        else:
            out.append((v, None))
    return out or None


def _parse_drop(vals):
    out = set()
    for v in vals or []:
        if "!" not in v:
            raise SystemExit(f"--drop は 'sheet!図形名' の形: {v!r}")
        s, n = v.split("!", 1)
        out.add((s, n))
    return out


def exit_code(rep, strict=False, strict_labels=False, strict_images=False) -> int:
    if rep["missing_total"] or (strict and rep["unmatched"]):
        return 1
    if strict_labels and rep["missing_labels_total"]:
        return 1
    if strict_images and rep["missing_images_total"]:
        return 1
    return 0


# ---------------------------------------------------------------------------
# selftest (合成の雛形 + 合成の PDF)
# ---------------------------------------------------------------------------
def _inject_drawing(xlsx, sheet_xml_part, anchors_xml):
    """openpyxl で作った xlsx の 1 sheet に drawing を差す (selftest 用)。"""
    import shutil
    import tempfile

    dr = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          f'<xdr:wsDr xmlns:xdr="{NS["xdr"]}" xmlns:a="{NS["a"]}">{anchors_xml}</xdr:wsDr>')
    tmp = tempfile.mktemp(suffix=".xlsx")
    with zipfile.ZipFile(xlsx) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            b = zi.read(i.filename)
            if i.filename == "[Content_Types].xml":
                b = b.replace(b"</Types>", b'<Override PartName="/xl/drawings/drawing1.xml" ContentType='
                              b'"application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
            if i.filename == sheet_xml_part:
                s = b.decode()
                s = s.replace("</worksheet>", f'<drawing xmlns:r="{NS["r"]}" r:id="rIdD1"/></worksheet>')
                b = s.encode()
            zo.writestr(i, b)
        rels = posixpath.join(posixpath.dirname(sheet_xml_part), "_rels", posixpath.basename(sheet_xml_part) + ".rels")
        zo.writestr(rels, f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{NS["pr"]}">'
                          f'<Relationship Id="rIdD1" Type="{REL_DRAWING}" Target="../drawings/drawing1.xml"/></Relationships>')
        zo.writestr("xl/drawings/drawing1.xml", dr)
    shutil.move(tmp, xlsx)


def _anchor(name, row, col, text, kind="oneCellAnchor", hidden=False):
    frm = f"<xdr:from><xdr:col>{col - 1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row - 1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
    size = '<xdr:ext cx="1000000" cy="200000"/>' if kind == "oneCellAnchor" else (
        f"<xdr:to><xdr:col>{col + 2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>")
    paras = "".join(f"<a:p><a:r><a:t>{t}</a:t></a:r></a:p>" for t in (text if isinstance(text, list) else [text]))
    body = f"<xdr:txBody><a:bodyPr/>{paras}</xdr:txBody>" if text else ""
    h = ' hidden="1"' if hidden else ""
    return (f"<xdr:{kind}>{frm}{size}<xdr:sp><xdr:nvSpPr><xdr:cNvPr id=\"2\" name=\"{name}\"{h}/><xdr:cNvSpPr/></xdr:nvSpPr>"
            f"<xdr:spPr/>{body}</xdr:sp><xdr:clientData/></xdr:{kind}>")


def _mk_docx(path, boxes=("枠の字",), body=("見出しA", "見出しB"), footer=("頁の字",)):
    W = NS["w"]

    def p(t):
        return f'<w:p><w:r><w:t>{t}</w:t></w:r></w:p>'

    doc = (f'<w:document xmlns:w="{W}"><w:body>' + "".join(p(t) for t in body)
           + "".join(f'<w:p><w:r><w:pict><w:txbxContent>{p(t)}</w:txbxContent></w:pict></w:r></w:p>' for t in boxes)
           + "</w:body></w:document>")
    ftr = f'<w:ftr xmlns:w="{W}">' + "".join(p(t) for t in footer) + "</w:ftr>"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", doc)
        z.writestr("word/footer1.xml", ftr)


def selftest() -> int:
    import os
    import tempfile

    import fitz
    import openpyxl

    d = tempfile.mkdtemp(prefix="static-text-selftest-")
    tpl = os.path.join(d, "tpl.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "出張願"
    for coord, v in {"B3": "出張者", "B4": "所属", "B5": "氏名", "B6": "用務", "B7": "期間", "B8": "宿泊先",
                     "D3": "甲-外部資金様式の説明", "B9": "支援課長代理", "B10": "潰す行の見出し", "D5": "記入例の字"}.items():
        ws[coord] = v
    ws.print_area = "A1:J20"
    ws2 = wb.create_sheet("報告書")
    for coord, v in {"B3": "報告者", "B4": "出張先", "B5": "成果", "B6": "日程"}.items():
        ws2[coord] = v
    wb.save(tpl)
    _inject_drawing(tpl, "xl/worksheets/sheet1.xml",
                    _anchor("Shape 1", 1, 1, "外部資金")
                    + _anchor("Shape 2", 1, 8, "甲-外部資金様式③-1", kind="twoCellAnchor")
                    + _anchor("Shape 3", 2, 30, "作成上の注意 横の注記")            # 印刷範囲の外
                    + _anchor("Shape 4", 10, 2, "")                                 # 字の無い丸
                    + _anchor("Shape 5", 11, 2, "隠れた字", hidden=True)
                    + _anchor("Shape 6", 12, 2, "回覧先の課")
                    + _anchor("Shape 7", 13, 2, ["複数", "段落の字"]))
    # 案件 = D5 の記入例の字を氏名に書き換え、 B10 の行を潰した (3pt)
    filled = os.path.join(d, "filled.xlsx")
    wbf = openpyxl.load_workbook(tpl)
    wbf["出張願"]["D5"] = "山田花子"
    wbf["出張願"].row_dimensions[10].height = 3
    wbf.save(filled)

    def pdf(name, pages, images=0):
        p = os.path.join(d, name)
        doc = fitz.open()
        for lines in pages:
            pg = doc.new_page()
            y = 60
            for t in lines:
                pg.insert_text((50, y), t, fontname="japan", fontsize=10)
                y += 16
            for k in range(images):
                pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8 + k, 8 + k), 0)
                pix.clear_with(0)
                pg.insert_image(fitz.Rect(400, 60 + 12 * k, 408 + k, 68 + 12 * k), pixmap=pix)
        doc.save(p)
        return p

    labels1 = ["出張者", "所属", "氏名", "用務", "期間", "宿泊先", "甲-外部資金様式の説明", "支援課⻑代理"]   # ⻑ = 部首の互換字
    labels2 = ["報告者", "出張先", "成果", "日程"]
    shapes_ok = ["外部資金", "甲-外部資金様式③-1", "回覧先の課", "複数", "段落の字"]
    cases = [
        ("全部在る", [labels1 + shapes_ok, labels2], {}, 0, 0),
        ("区分の枠だけ無い (様式番号と label の中の『外部資金』 で代わりにしない)",
         [labels1 + ["甲-外部資金様式③-1", "回覧先の課", "複数", "段落の字"], labels2], {}, 1, 1),
        ("様式番号の末尾が切れた (③-1 → ③-)",
         [labels1 + ["外部資金", "甲-外部資金様式③-", "回覧先の課", "複数", "段落の字"], labels2], {}, 1, 1),
        ("図形が全部落ちた", [labels1, labels2], {}, 1, 5),
        ("drop した図形は求めない", [labels1 + ["外部資金", "甲-外部資金様式③-1", "複数", "段落の字"], labels2],
         {"drop": {("出張願", "Shape 6")}}, 0, 0),
        ("頁の順が違っても label で頁を選ぶ", [labels2, labels1 + shapes_ok], {}, 0, 0),
        ("頁の text 層が無い (raster) = ⚪ で exit 0", [[], []], {}, 0, 0),
        ("頁の text 層が無い + strict = exit 1", [[], []], {"strict": True}, 1, 0),
    ]
    fails = 0
    for label, pages, opt, want_rc, want_missing in cases:
        p = pdf("out.pdf", pages)
        rep = check(tpl, p, targets=[("出張願", None), ("報告書", None)], drop=opt.get("drop"))
        rc = exit_code(rep, strict=opt.get("strict", False))
        ok = rc == want_rc and rep["missing_total"] == want_missing
        fails += not ok
        print(f"{'PASS' if ok else 'FAIL'} {label}: rc={rc} missing={rep['missing_total']}")
        if not ok:
            print("\n".join(render(rep)))
    # 範囲の外・隠れた図形・字の無い図形は数えない
    rep = check(tpl, pdf("out.pdf", [labels1 + shapes_ok, labels2]), targets=[("出張願", None)])
    t = rep["targets"][0]
    ok = t["outside"] == 1 and t["hidden"] == 1 and t["textless"] == 1 and t["checked"] == 5
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} 範囲外 1 / 隠れ 1 / 字なし 1 / 照合 5 段落: {t['outside']}/{t['hidden']}/{t['textless']}/{t['checked']}")
    # --filled: 書き換えていない label の欠け (「所属」 が無い) / 記入した cell の雛形の字 (記入例の字) は求めない /
    #           潰した行の label は求めない / 部首の互換字 (⻑) は受ける / 記入値と label 以外の字 = 増えた字
    out = pdf("out.pdf", [[x for x in labels1 if x != "所属"] + shapes_ok + ["山田花子", "余計な字"], labels2])
    rep = check(tpl, out, targets=[("出張願", None), ("報告書", None)], filled=filled)
    t = rep["targets"][0]
    miss = [m["text"] for m in t["missing_labels"]]
    ok = (miss == ["所属"] and t["labels_checked"] == 8 and rep["missing_labels_total"] == 1
          and exit_code(rep) == 0 and exit_code(rep, strict_labels=True) == 1
          and t["extra"] == ["余計な字"])
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} --filled: 欠けた見出し = {miss} (8 cell 照合、 記入例の字・潰した行は除く、 ⻑ を受ける)、 増えた字 = {t['extra']}")
    # --blank: 素刷りより画像が少ない (checkbox の箱の消失)
    blank = pdf("blank.pdf", [labels1 + shapes_ok, labels2], images=2)
    out = pdf("out2.pdf", [labels1 + shapes_ok + ["山田花子"], labels2], images=1)
    rep = check(tpl, out, targets=[("出張願", None), ("報告書", None)], filled=filled, blank=blank)
    b = rep["targets"][0]["blank"]
    # 合成 PDF は全頁に画像を置く = 2 対象 × (2 → 1) で合計 2
    ok = b and b["images"] == [2, 1] and rep["missing_images_total"] == 2 and exit_code(rep) == 0 and exit_code(rep, strict_images=True) == 1
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} --blank: 画像 2 → 1 を拾う (rc は strict-images の時だけ 1): {b}")
    ok = any("素刷りより画像が少ない" in x for x in render(rep))
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} render: 画像の減少の行")
    # docx: textbox / footer の字は図形の字 (exit 1)、 本文の見出しは label、 案件で書き換えた段落は除く
    dt, df = os.path.join(d, "t.docx"), os.path.join(d, "f.docx")
    _mk_docx(dt)
    _mk_docx(df, body=("見出しA", "記入した値"))
    out = pdf("d.pdf", [["見出しA", "記入した値", "頁の字"]])
    rep = check_docx(dt, out, filled=df)
    ok = (rep["missing_total"] == 1 and rep["targets"][0]["missing"][0]["text"] == "枠の字"
          and rep["missing_labels_total"] == 0 and exit_code(rep) == 1)
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} docx: textbox の字の欠けは 🔴、 書き換えた段落は求めない: missing={rep['missing_total']} labels={rep['missing_labels_total']}")
    out = pdf("d2.pdf", [["枠の字", "記入した値", "頁の字"]])
    rep = check_docx(dt, out, filled=df)
    ok = rep["missing_total"] == 0 and rep["missing_labels_total"] == 1 and rep["targets"][0]["missing_labels"][0]["text"] == "見出しA"
    fails += not ok
    print(f"{'PASS' if ok else 'FAIL'} docx: 本文の見出しの欠けは label: {rep['missing_labels_total']}")
    print("ALL PASS" if not fails else f"{fails} FAIL")
    return 1 if fails else 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("template", nargs="?")
    ap.add_argument("pdf", nargs="?")
    ap.add_argument("--target", action="append", help="'sheet' か 'sheet!A1:AH60' (繰り返し可)")
    ap.add_argument("--drop", action="append", help="'sheet!図形名' = 刷らないと決めた図形 (繰り返し可)")
    ap.add_argument("--filled", help="案件の xlsx / docx (記入済み、 または体裁を当てた temp) = cell の見出しと記入値も見る")
    ap.add_argument("--blank", help="雛形の素刷り PDF (道具を通さず app で刷ったもの) = 画像・線の数を比べる")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--strict", action="store_true", help="対象の頁が見つからない (⚪) も exit 1")
    ap.add_argument("--strict-labels", action="store_true", help="雛形の見出し (cell) の欠けも exit 1")
    ap.add_argument("--strict-images", action="store_true", help="素刷りより画像が少ないのも exit 1")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    if not a.template or not a.pdf:
        ap.print_usage(sys.stderr)
        return 2
    try:
        if str(a.template).lower().endswith(".docx"):
            rep = check_docx(a.template, a.pdf, filled=a.filled)
        else:
            rep = check(a.template, a.pdf, _parse_targets(a.target), _parse_drop(a.drop), filled=a.filled, blank=a.blank)
    except (ValueError, KeyError, zipfile.BadZipFile, FileNotFoundError) as e:
        print(f"check-form-static-text: {e}", file=sys.stderr)
        return 2
    if a.json:
        print(json.dumps(rep, ensure_ascii=False, indent=1))
    else:
        print(f"雛形の図形の字 ({rep['template'].split('/')[-1]} → {rep['pdf'].split('/')[-1]}):")
        print("\n".join(render(rep)))
    return exit_code(rep, strict=a.strict, strict_labels=a.strict_labels, strict_images=a.strict_images)


if __name__ == "__main__":
    sys.exit(main())
