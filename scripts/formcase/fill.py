"""案件の fill stub (fill_<doc>.py) の実行系。 値を Excel で書き → 読み戻し → gate を回して結果を出す。

EDITS の行 = (sheet, cell, kind, value[, group])。 ``--group G`` でその group の行だけ書き、 gate もその group だけ裁く
(= 後で出す group の欄が空のままでも止めない)。 value=None の行は未記入として一覧に出し、 書かない。
gate の FAIL は「まだ直すところがある」 の報告 (exit 1)。 PDF を作るのは build だけで、 build は gate が通らないと作らない。
"""
from __future__ import annotations

import datetime as _dt
import sys
import warnings
from pathlib import Path

from . import excel as X
from . import gates as GT
from . import guard as G
from . import manifest as M


def value_from(workbook, sheet: str, cell: str):
    """住所・口座のような個人情報を stub に書かずに EDITS へ入れる: 既にその値を持つ workbook (本人が前に出して受理された様式・
    案件の置き場の xlsx) の cell を実行時に読む。 値は新しい workbook の欄にだけ入り、 stub には出所の path と cell だけが残る
    (form-case-pipeline.md #pii-runtime-source)。 読めない・空なら止める (= 空欄のまま紙にしない)。"""
    import openpyxl

    path = Path(workbook).expanduser()
    if not path.is_file():
        raise SystemExit(f"🔴 値の出所の workbook が無い: {path}")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"🔴 値の出所に sheet が無い: {path.name} / {sheet}")
    v = wb[sheet][cell].value
    wb.close()
    if v in (None, ""):
        raise SystemExit(f"🔴 値の出所の cell が空: {path.name} / {sheet}!{cell}")
    return v


def value_from_text(path, pattern: str, group: int | str = 1) -> str:
    """value_from の text 版: 本人の回答を転記した記録 (暗号化した markdown の台帳など、 workbook でない置き場) から
    正規表現で 1 つ取る。 stub には出所の path と pattern だけが残る (form-case-pipeline.md #pii-runtime-source)。
    file が無い・pattern が当たらない・当たった値が空なら止める (= 空欄のまま紙にしない / 書式が変わったら気づく)。
    複数に当たる pattern も止める (= どれを取ったか分からない値を紙に出さない)。"""
    import re

    p = Path(path).expanduser()
    if not p.is_file():
        raise SystemExit(f"🔴 値の出所の file が無い: {p}")
    hits = list(re.finditer(pattern, p.read_text(encoding="utf-8"), re.M))
    if len(hits) != 1:
        raise SystemExit(f"🔴 値の出所 {p.name} で pattern の当たりが {len(hits)} 件 (1 件だけにする): {pattern}")
    v = (hits[0].group(group) or "").strip()
    if not v:
        raise SystemExit(f"🔴 値の出所 {p.name} で当たった値が空: {pattern}")
    return v


def run_docx(stub_file, doc_id, fields, choices, texts) -> int:
    """Word 様式の記入 (form-case-pipeline.md #docx)。 FIELDS を雛形から作り直した docx に打ち、 CHOICES / TEXTS を overlay.yaml に書き、
    読み戻し → gate。 value=None は未記入 (書かない)。 凍結 group がある document には書かない (先に reopen)。"""
    import yaml

    from . import docx_form as DF
    from . import specs as S

    dry = "--dry-run" in sys.argv
    case = G.find_manifest_dir(Path(stub_file).resolve().parent)
    if case is None:
        print("🔴 submission.yaml が見つからない (formcase.py new で作った dir で実行する)")
        return 2
    m = M.load(case)
    doc = m.doc(doc_id)
    spec = S.get(doc.get("form"))
    if not DF.is_docx(spec):
        print(f"🔴 form {doc.get('form')!r} は docx 様式でない")
        return 2
    frozen = [g for g in (doc.get("groups") or {}) if m.group_is_frozen(doc_id, g)]
    if frozen:
        print(f"🔴 凍結 group {frozen} の document = 書かない → 先に formcase.py reopen (新しい issue)")
        return 2
    bad_ids = ([f"FIELDS.{k}" for k in fields if k not in {f['id'] for f in DF.fields(spec)}]
               + [f"CHOICES.{k}" for k in choices if k not in {c['id'] for c in DF.choices(spec)}]
               + [f"TEXTS.{k}" for k in texts if k not in {x['id'] for x in DF.texts(spec)}])
    if bad_ids:
        print(f"🔴 spec に無い id: {bad_ids} (spec が変わったら stub を作り直すか id を直す)")
        return 2
    optional = {x["id"] for x in DF.fields(spec) + DF.choices(spec) + DF.texts(spec) if x.get("optional")}
    todo = [k for d in (fields, choices, texts) for k, v in d.items() if v is None and k not in optional]
    if todo:
        print(f"🟡 未記入 {len(todo)} 欄 (value=None、 書かない): " + ", ".join(todo[:12]) + (" …" if len(todo) > 12 else ""))
    if dry:
        return 0
    wb = m.workbook(doc_id)
    tpl = S.template_path(spec)
    try:
        if DF.word_available():
            DF.write_word(spec, tpl, fields, wb)          # D6 (2026-09-25): Word で書く = 欄の run の書式・図・field が雛形のまま
        else:
            print("⚪ Word が無い (か FORMCASE_DOCX_WRITER=python) = python-docx で書く (空欄の run の書式が既定に落ちうる、 下の ⚠️)")
            DF.write(spec, tpl, fields, wb)
    except DF.DocxFormError as e:
        print(f"🔴 {e}")
        return 2
    for line in DF.run_format_lines(spec, tpl, wb, fields):   # 欄の run の書式を雛形と照合 (D6)
        print("   " + line)
    ov = {"choices": {k: v for k, v in choices.items() if v is not None},
          "texts": {k: v for k, v in texts.items() if v is not None}}
    DF.overlay_path(wb).write_text(
        "# formcase の Word 様式: PDF に重ねる値 (fill_<doc>.py が書く生成物、 手で直さない)\n"
        + yaml.safe_dump(ov, allow_unicode=True, sort_keys=False, width=1000), encoding="utf-8")
    bad = DF.readback(spec, tpl, wb, fields)
    for fid, want, got in bad:
        print(f"🔴 読み戻し {fid}: 期待 {want!r} / 実際 {got!r}")
    ok = GT.run_scoped(m, doc_id, list(doc.get("groups") or {}), raise_on_fail=False)
    good = ok and not bad and not todo
    print("✅ 記入と gate が通った → formcase.py build で PDF" if good else "🔴 まだ直すところがある (上の 未記入 / FAIL / 読み戻し)")
    return 0 if good else 1


def _readback(wb_path, edits) -> list:
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(wb_path)
    bad = []
    for sheet, cell, kind, val in edits:
        if kind == "general":          # 表示書式だけ (値は次の formula 行が確かめる)
            continue
        if sheet not in wb.sheetnames:
            bad.append((sheet, cell, val, "<sheet 無し>"))
            continue
        if ":" in cell:                # 範囲 (merge の全域 clear)
            vals = [c.value for row in wb[sheet][cell] for c in row]
            if kind != "clear" or any(v not in (None, "") for v in vals):
                bad.append((sheet, cell, val, [v for v in vals if v not in (None, "")][:3]))
            continue
        if kind == "fontsize":
            got = wb[sheet][cell].font.sz
            if got is None or abs(float(got) - float(val)) > 0.01:
                bad.append((sheet, cell, f"font {val}", got))
            continue
        got = wb[sheet][cell].value
        if kind == "formula":
            ok = isinstance(got, str) and got.replace(" ", "") == str(val).replace(" ", "")
        elif kind == "clear":
            ok = got in (None, "")
        elif kind == "date":
            d = got.date() if isinstance(got, _dt.datetime) else got
            want = val.date() if isinstance(val, _dt.datetime) else val
            ok = isinstance(got, (_dt.date, _dt.datetime)) and d == want
        elif kind == "number":
            ok = isinstance(got, (int, float)) and float(got) == float(val)
        else:
            ok = str(got) == str(val)
        if not ok:
            bad.append((sheet, cell, val, got))
    return bad


def _writes(row) -> bool:
    """書く行か。 value=None は未記入 (書かない) だが、 clear / general は値を持たない kind なので書く。"""
    return row[3] is not None or row[2] in ("clear", "general")


def _arg(name):
    a = sys.argv
    return a[a.index(name) + 1] if name in a and a.index(name) + 1 < len(a) else None


def run(stub_file, doc_id, edits) -> int:
    dry = "--dry-run" in sys.argv
    only = _arg("--group")
    case = G.find_manifest_dir(Path(stub_file).resolve().parent)
    if case is None:
        print("🔴 submission.yaml が見つからない (formcase.py new で作った dir で実行する)")
        return 2
    m = M.load(case)
    groups = list((m.doc(doc_id).get("groups") or {}).keys())
    frozen = [g for g in groups if m.group_is_frozen(doc_id, g)]
    rows = [(e[0], e[1], e[2], e[3], e[4] if len(e) > 4 else None) for e in edits]
    if only:
        if only not in groups:
            print(f"🔴 group {only!r} は manifest に無い ({groups})")
            return 2
        rows = [r for r in rows if r[4] in (only, None)]
    if only in frozen or (not only and any(r[4] in frozen for r in rows if _writes(r))):
        print(f"🔴 凍結 group {frozen} の欄を書こうとしている → 先に formcase.py reopen (新しい issue)。 何も書かない")
        return 2
    todo = [(s, c) for s, c, k, v, _g in rows if not _writes((s, c, k, v))]
    if todo:
        print(f"🟡 未記入 {len(todo)} 欄 (value=None、 書かない): " + ", ".join(f"{s}!{c}" for s, c in todo[:12])
              + (" …" if len(todo) > 12 else ""))
    if dry:
        return 0
    edits4 = [(s, c, k, v) for s, c, k, v, _g in rows if _writes((s, c, k, v))]
    wb = m.workbook(doc_id)
    if edits4:
        try:
            X.write_cells(wb, edits4)
        except X.ExcelError as e:          # osascript の標準エラーは逐語で e に入っている (excel.run_osascript)
            print(f"🔴 {e}")
            return 2
    bad = _readback(wb, edits4)
    for s, c, want, got in bad:
        print(f"🔴 読み戻し {s}!{c}: 期待 {want!r} / 実際 {got!r}")
    targets = [only] if only else [g for g in groups if g not in frozen]
    ok = GT.run_scoped(m, doc_id, targets, raise_on_fail=False)
    print("✅ 記入と gate が通った → formcase.py build で PDF" if ok and not bad and not todo else
          "🔴 まだ直すところがある (上の 未記入 / FAIL / 読み戻し)")
    return 0 if ok and not bad and not todo else 1
