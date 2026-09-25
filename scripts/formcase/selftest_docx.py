"""Word 様式 (docx_form) と 刷る頁 (page_roles) の test。 fixture は全部合成 (Office 不要)。

実 spec に対する同じ検査 (雛形が在るか・欄の場所が在るか・頁の役割に矛盾が無いか) は instance の selftest が持つ
(= どの様式かは engine の知るところでない)。
"""
from __future__ import annotations

import warnings  # noqa: F401
from pathlib import Path

from . import manifest as M


def run_docx_tests(tmp: Path, expect) -> None:
    """Word 様式 (docx_form.py、 form-case-pipeline.md #docx) = 合成の雛形 docx と spec で、 記入・gate・digest・凍結の判定 (Office 不要)。"""
    import docx
    import yaml

    from . import docx_form as DF
    from . import fill as FI
    from . import scaffold as SC
    from . import specs as SP

    root = tmp / "docx"
    root.mkdir()
    tpl = root / "tpl.docx"
    d = docx.Document()
    d.add_paragraph("見出し")
    d.add_paragraph("20　　年　　月　　日")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 0).text = "氏名", "目的"
    t.cell(1, 1).text = "（注記）"
    d.save(tpl)
    spec = {"meta": {"id": "fx-docx", "kind": "docx", "pages": 1}, "groups": {"g": {"pages": [1]}},
            "docx": {"fields": [{"id": "date", "at": {"para": 1}, "label": "日付"},
                                {"id": "name", "at": {"table": 0, "row": 0, "col": 1}, "format": "{value}　印"},
                                {"id": "why", "at": {"table": 0, "row": 1, "col": 1}, "mode": "append"},
                                {"id": "memo", "at": {"table": 0, "row": 1, "col": 0}, "optional": True}],
                     "choices": [{"id": "role", "kind": "word", "options": ["教授", "講師"]},
                                 {"id": "a", "kind": "paren", "anchor": "A", "optional": True},
                                 {"id": "b", "kind": "paren", "anchor": "B", "optional": True}],
                     "choice_groups": [{"label": "AB", "ids": ["a", "b"], "min": 1}],
                     "texts": [{"id": "tel", "label_text": "電話"}]}}

    def ov(choices, texts):
        DF.overlay_path(out).write_text(yaml.safe_dump({"choices": choices, "texts": texts}, allow_unicode=True),
                                        encoding="utf-8")

    expect("docx: is_docx は meta.kind で判定", DF.is_docx(spec) and not DF.is_docx({"meta": {"id": "xl"}}))
    out = root / "case.docx"
    vals = {"date": "2026年9月18日", "name": "甲野太郎", "why": "研究発表"}
    DF.write(spec, tpl, vals, out)
    DF.write(spec, tpl, vals, out)          # 2 回目も雛形から = append が二重にならない
    dd = docx.Document(out)
    cells = [p.text for p in dd.tables[0].cell(1, 1).paragraphs]
    expect("docx: 記入 (text / format / append)、 2 回書いても append は 1 段落",
           dd.paragraphs[1].text == "2026年9月18日" and dd.tables[0].cell(0, 1).text == "甲野太郎　印"
           and cells == ["（注記）", "研究発表"], cells)
    expect("docx: 読み戻しが一致", DF.readback(spec, tpl, out, vals) == [], DF.readback(spec, tpl, out, vals))
    expect("docx: 図・field・記号の無い雛形なら消失の行は出ない", DF.object_loss(tpl, out) == [], DF.object_loss(tpl, out))
    # D6 (Word で書く): 段落番号の写像 = 本文 → 表の cell の段落 → 行末の記号 の順 (実物 2 雛形で 319/321 一致、 残り 2 = 改行・改頁の表記)
    dt = docx.Document(tpl)
    wi = DF.WordIndex(dt)
    expect("D6: WordIndex = 本文 [1, 2]、 表 (0,0)=3 (0,1)=4 行末=5 (1,0)=6 (1,1)=7",
           wi.body == [1, 2] and wi.of(dt, {"table": 0, "row": 0, "col": 1}) == [4]
           and wi.of(dt, {"table": 0, "row": 1, "col": 1}) == [7], (wi.body, wi.cells))
    edits, expect_txt = DF.word_edits(spec, tpl, vals)
    expect("D6: word_edits = text は番号と値、 append は insert_after、 照合の字は雛形の段落の字",
           (2, "text", "2026年9月18日") in edits and (4, "text", "甲野太郎　印") in edits
           and any(e[0] == 7 and e[1] == "insert_after" and e[2][0] == "研究発表" for e in edits)
           and expect_txt.get(7) == "（注記）" and expect_txt.get(2) == "20　　年　　月　　日", (edits, expect_txt))
    # run の書式の照合: 雛形の空欄の段落記号に sz=20 が付いていると、 python-docx の記入 (既定の run) は ⚠️
    from docx.oxml.ns import qn
    from lxml import etree

    tpl2 = root / "tpl-sz.docx"
    d2 = docx.Document(tpl)
    p2 = d2.tables[0].cell(0, 1).paragraphs[0]
    ppr = p2._p.get_or_add_pPr()
    rpr = etree.SubElement(ppr, qn("w:rPr"))
    etree.SubElement(rpr, qn("w:sz")).set(qn("w:val"), "20")
    d2.save(tpl2)
    out2b = root / "case-sz.docx"
    DF.write(spec, tpl2, vals, out2b)
    fl = DF.run_format_lines(spec, tpl2, out2b, vals)
    expect("D6: run_format_lines = 雛形の sz 20 が記入後の run に無い → ⚠️ (name)",
           len(fl) == 1 and "name" in fl[0] and "sz '20'" in fl[0], fl)
    expect("D6: 書式の指定が無い雛形 (tpl) は ⚠️ 0", DF.run_format_lines(spec, tpl, out, vals) == [])
    # 欄の段落の run に記号 (w:sym = Wingdings の □ 等) があると、 run の書き換えで黙って消える (2026-09-24 RCA)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tpl2, out2 = root / "tpl-sym.docx", root / "case-sym.docx"
    d2 = docx.Document(tpl)
    sym = OxmlElement("w:sym")
    sym.set(qn("w:font"), "Wingdings")
    sym.set(qn("w:char"), "F0A8")
    d2.paragraphs[1].add_run("")._r.append(sym)
    d2.save(tpl2)
    DF.write(spec, tpl2, vals, out2)
    loss = DF.object_loss(tpl2, out2)
    expect("docx: 欄の run の中の記号が書き換えで消えたら ⚠️ の行", len(loss) == 1 and "記号" in loss[0], loss)
    ok, lines = DF.gate(spec, tpl, out)
    expect("docx gate: 必須の選択肢・choice_group・短い値が未記入なら FAIL", not ok
           and any("role" in x and "未選択" in x for x in lines) and any("AB" in x for x in lines)
           and any("tel" in x and "未記入" in x for x in lines), lines)
    ov({"role": "教授", "b": True}, {"tel": "0"})
    ok, lines = DF.gate(spec, tpl, out)
    expect("docx gate: 全部そろえば通る (任意の欄の空は 🧊)", ok and any("memo" in x and "🧊" in x for x in lines), lines)
    ov({"role": "総長", "b": True, "zz": 1}, {"tel": "0"})
    ok, lines = DF.gate(spec, tpl, out)
    expect("docx gate: 選択肢に無い値 / spec に無い id は FAIL",
           not ok and any("総長" in x for x in lines) and any("zz" in x for x in lines), lines)
    DF.write(spec, tpl, {"date": "2026年9月18日", "why": "研究発表"}, out)
    ok, lines = DF.gate(spec, tpl, out)
    expect("docx gate: 必須の欄が雛形のまま = FAIL", not ok and any("name" in x and "未記入" in x for x in lines), lines)
    DF.write(spec, tpl, vals, out)
    dg = DF.docx_digest(out)
    DF.write(spec, tpl, vals, root / "same.docx")
    expect("docx digest: 同じ値で作り直すと同じ", DF.docx_digest(root / "same.docx") == dg)
    DF.write(spec, tpl, dict(vals, name="乙野次郎"), root / "other.docx")
    expect("docx digest: 値を変えると values が変わる", DF.docx_digest(root / "other.docx")["values"] != dg["values"])
    d3 = docx.Document(out)
    d3.paragraphs[1].runs[0].font.size = docx.shared.Pt(20)
    d3.save(root / "fmt.docx")
    g3 = DF.docx_digest(root / "fmt.docx")
    expect("docx digest: 字の大きさだけ変えると format だけが変わる", g3["values"] == dg["values"] and g3["format"] != dg["format"])
    ov({"role": "教授", "b": True}, {"tel": "0"})
    fr = {"docx_digest": dg, "overlay_digest": DF.overlay_digest(out)}
    expect("凍結 (docx): 変化なし = 差なし", M.frozen_sheet_changes(fr, out) == ([], []))
    expect("凍結 (docx): 書式だけの変化 = format 側", M.frozen_sheet_changes({"docx_digest": dg}, root / "fmt.docx") == ([], ["docx"]))
    ov({"role": "講師", "b": True}, {"tel": "0"})
    expect("凍結 (docx): PDF に重ねる値 (overlay.yaml) の変化も値の変化", M.frozen_sheet_changes(fr, out) == (["overlay.yaml"], []))
    expect("manifest: docx digest の版と has_source_digest", M.frozen_digest_version(fr) == "docx" and M.has_source_digest(fr))
    stub = SC.docx_stub_text(dict(spec, _path=root / "fx.yaml"), "case", "~/x", "doc", "fill_doc.py")
    expect("docx stub: FIELDS / CHOICES / TEXTS を並べて run_docx を呼ぶ",
           all(k in stub for k in ("FIELDS = {", "'why': None", "CHOICES = {", "'role': None", "TEXTS = {", "run_docx(")))
    expect("docx: fill.run_docx がある", callable(getattr(FI, "run_docx", None)))
    expect("docx: 欄の場所が雛形に無ければ DocxFormError (= 雛形の改訂を黙って通さない)",
           _raises(DF.DocxFormError, DF.write,
                   {"meta": {"id": "fx-docx", "kind": "docx", "pages": 1}, "groups": {"g": {"pages": [1]}},
                    "docx": {"fields": [{"id": "x", "at": {"table": 9, "row": 0, "col": 0}}]}},
                   tpl, {"x": "v"}, root / "bad.docx"))


def _raises(exc, fn, *a, **k) -> bool:
    try:
        fn(*a, **k)
    except exc:
        return True
    except Exception:  # noqa: BLE001
        return False
    return False


def run_page_role_tests(tmp: Path, expect) -> None:
    """刷る頁 = spec の page_roles の検査・Word の頁の目印・出力への宣言 (Office 不要)。 fixture は合成。"""
    import fitz

    from . import recipes as RC
    from . import specs as SP

    base = {"meta": {"id": "fx", "kind": "docx", "pages": 3}, "groups": {"g": {"pages": [1, 2]}},
            "page_roles": {1: {"role": "submit", "anchor": "申請書"}, 2: {"role": "submit", "anchor": "計画書"},
                           3: {"role": "instructions", "anchor": "について"}}}
    expect("page_roles: 様式 2 頁 + 説明書き 1 頁 = 矛盾なし", not SP.page_role_problems(base), SP.page_role_problems(base))
    for label, patch, word in (
            ("Word 様式で page_roles が全頁に無い", {"page_roles": {1: {"role": "submit", "anchor": "a"}}}, "全部"),
            ("group に説明書きの頁", {"groups": {"g": {"pages": [1, 2, 3]}}}, "窓口に出さない"),
            ("submit の頁がどの group にも無い", {"groups": {"g": {"pages": [1]}}}, "刷られない"),
            ("役割の語が無い語", {"page_roles": {**base["page_roles"], 3: {"role": "note"}}}, "どれでもない"),
            ("submit に anchor が無い", {"page_roles": {**base["page_roles"], 2: {"role": "submit"}}}, "anchor")):
        probs = SP.page_role_problems(dict(base, **patch))
        expect(f"page_roles の矛盾を検出: {label}", any(word in x for x in probs), probs)
    expect("Excel 様式 (kind なし) は page_roles が無くてよい",
           not SP.page_role_problems({"meta": {"id": "x"}, "groups": {"g": {"pages": [1]}}}))

    root = tmp / "pages"
    root.mkdir()
    pdf = root / "word.pdf"
    d = fitz.open()
    for t in ("国際 交流 申請書", "研究 計画書", "交流費について"):
        p = d.new_page()
        p.insert_text((72, 72), t, fontname="japan", fontsize=16)
        p.draw_line((60, 200), (500, 200))
    d.save(pdf)
    expect("頁の目印: 空白を無視して各頁に在る", RC.check_page_anchors(pdf, base) == [], RC.check_page_anchors(pdf, base))
    moved = dict(base, page_roles={**base["page_roles"], 2: {"role": "submit", "anchor": "について"}})
    expect("頁の目印: 雛形の頁がずれたら頁と目印を返す", RC.check_page_anchors(pdf, moved) == [(2, "について")])

    out = root / "g.pdf"
    RC._extract(pdf, [1, 2], out)
    lines = RC.declare_pages(out, base, "g", [1, 2])
    rec = RC.PP.read_record(fitz.open(out))
    expect("出力への宣言: 全頁 submit・外した頁の記録",
           [x["role"] for x in rec["pages"]] == ["submit", "submit"] and [x["from"] for x in rec["dropped"]] == [3]
           and not lines, (rec, lines))
    ras = root / "g_raster.pdf"
    RC._raster(out, ras, dpi=36)
    expect("raster の出力も宣言を引き継ぐ (刷る直前の gate が読む)", RC.PP.read_record(fitz.open(ras)) is not None)
    # Excel 様式 (page_roles なし) の group に記載例の頁が入ると止まる / 「…について」 だけなら行に出して通す
    ex = root / "ex.pdf"
    d = fitz.open()
    for t in ("旅費請求書", "記入例"):
        p = d.new_page()
        p.insert_text((72, 72), t, fontname="japan", fontsize=16)
        p.draw_line((60, 200), (500, 200))
    d.save(ex)
    try:
        RC.declare_pages(ex, {"meta": {"id": "x"}}, "g", [1, 2])
        expect("宣言の無い group に記載例の頁 = BuildError", False)
    except RC.BuildError as e:
        expect("宣言の無い group に記載例の頁 = BuildError", "記載例" in str(e), e)
    RC._extract(pdf, [1, 3], root / "weak.pdf")
    wl = RC.declare_pages(root / "weak.pdf", {"meta": {"id": "x"}}, "g", [1, 3])
    expect("宣言の無い group の「…について」 頁は止めずに行に出す", any("説明書き?" in x for x in wl), wl)
