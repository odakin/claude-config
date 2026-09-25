"""Excel の操作は全部ここ (staging 経由・前面に出さない・1 回に 1 job)。

- ``export_pdfs([(xlsx, pdf[, ops]), ...])``: 1 回の Excel session で workbook を順に開き、 ops (体裁の変更の
  AppleScript 行) を当てて PDF にして閉じる (保存しない)。 ``export_pdf(xlsx, pdf)`` はその 1 件版。
- ``write_cells(xlsx, edits)``: 値を Excel osascript で書く (openpyxl で save しない = 数式 cache・drawing・form control を守る)。
- ``ops_lines(sheet, ops)``: 体裁の変更 (行高・折り返し・揃え・字の大きさ・結合・表示書式・印刷範囲・1 枚に収める・白黒) を
  AppleScript に。 openpyxl で save すると図形 (標題・checkbox 等の drawing) が消える様式は、 体裁を openpyxl の
  読み込みの上で計算し (layout.py)、 Excel に当てる。
Excel には staging root の copy だけを触らせる (層1 office-automation.md#office-inplace-guard)。

osascript の失敗: 標準エラーを逐語で build の出力と ``ExcelError`` の文に出し、
``~/.local/state/formcase/excel-errors.log`` に 1 行ずつ残す (= 長い出力の途中が切れても、 最後の行と log に残る)。
一時的な AppleEvent エラー (-1712 = 応答待ちの timeout / -609 = Excel との接続が切れた / -50 = 起動直後に
引数を受け付けない) は、 自分の staged copy を閉じ・Excel が落ちていれば起こし直してから ``RETRIES`` 回まで
再試行する。 それ以外のエラーと、 再試行を使い切ったものは止める。
"""
from __future__ import annotations

import datetime as _dt
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "lib"))
from office_staging import Stage, office_app  # noqa: E402

TRANSIENT = (-1712, -609, -50)
RETRIES = 2
RETRY_WAIT = (3.0, 8.0)
ERROR_LOG = Path(os.environ.get("XDG_STATE_HOME") or (Path.home() / ".local/state")) / "formcase" / "excel-errors.log"


class ExcelError(RuntimeError):
    """Excel (osascript) の失敗。 文に osascript の標準エラーを逐語で含む。"""


class _ExcelSession:
    """この job が起動した Excel だけを後で終わらせ、 前面を元に戻す (層1 office-automation.md#office-app-reset-guard)。"""

    def __enter__(self):
        self.front = office_app("front-remember").stdout.strip()
        self.launched = "launched=1" in office_app("launch", "excel").stdout
        return self

    def relaunch(self):
        """再試行の前: Excel が落ちていたら (-609) 起こし直す。 ここで起動したら出る時に終わらせる。"""
        if "launched=1" in office_app("launch", "excel").stdout:
            self.launched = True

    def __exit__(self, *exc):
        office_app("release", "excel", "1" if self.launched else "0")
        if self.front:
            office_app("front-restore", self.front, "Microsoft Excel.app")
        return False


def error_code(text: str):
    """osascript の標準エラー末尾の ``(-1712)`` から AppleEvent のエラー番号を取る (無ければ None)。"""
    m = re.search(r"\((-?\d+)\)\s*$", (text or "").strip())
    return int(m.group(1)) if m else None


def _log_error(label, attempt, rc, text) -> None:
    try:
        ERROR_LOG.parent.mkdir(parents=True, exist_ok=True)
        one = " ⏎ ".join((text or "").strip().splitlines())
        with ERROR_LOG.open("a", encoding="utf-8") as fh:
            fh.write(f"{_dt.datetime.now().isoformat(timespec='seconds')}\t{label}\ttry={attempt}\trc={rc}\t"
                     f"code={error_code(text)}\t{one}\n")
    except OSError:
        pass


PROBE_SCRIPT = '''
with timeout of 15 seconds
  tell application id "com.microsoft.Excel" to get count of workbooks
end timeout
'''


def excel_responds(runner=None) -> bool:
    """Excel が 15 秒以内に AppleEvent に答えるか (起動していなければ True = 再試行で起こせる)。"""
    runner = runner or (lambda argv: subprocess.run(argv, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True))
    if "not-running" in office_app("state", "excel").stdout:
        return True
    return runner(["osascript", "-e", PROBE_SCRIPT]).returncode == 0


def run_osascript(script: str, args, label: str, session=None, staged=(), runner=None, sleep=time.sleep,
                  responds=None) -> str:
    """osascript を回す。 失敗の標準エラーは逐語で出し、 一時的なエラー (TRANSIENT) は staged copy を閉じてから
    RETRIES 回まで再試行する。 ただし -1712 の後に Excel が 15 秒の問い合わせにも答えない (= dialog が出ている /
    固まっている) なら再試行しない (= 200 秒の timeout を重ねない。 Excel を終わらせることもしない = 人が見る)。
    runner / sleep / responds は selftest 用の差し替え口。"""
    args = [str(a) for a in args]
    runner = runner or (lambda argv: subprocess.run(argv, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True))
    responds = responds or excel_responds
    for attempt in range(1, RETRIES + 2):
        r = runner(["osascript", "-e", script, *args])
        if r.returncode == 0:
            return r.stdout
        err = (r.stderr or "").strip() or f"(標準エラーなし。 標準出力: {(r.stdout or '').strip()[:300]!r})"
        code = error_code(err)
        _log_error(label, attempt, r.returncode, err)
        print(f"   ⚠️  Excel ({label}) 試行 {attempt}: osascript exit {r.returncode}: {err}")
        if code == -1712 and not responds():
            _log_error(label, attempt, r.returncode, "Excel が 15 秒の問い合わせに答えない → 再試行しない")
            raise ExcelError(f"Excel ({label}) が応答しない (-1712 の後の 15 秒の問い合わせにも答えない = dialog が出ているか"
                             f"固まっている)。 再試行しない・Excel を終わらせない → Excel の画面を見て dialog を閉じる / "
                             f"保存してから自分で終了して再実行 (log = {ERROR_LOG}): {err}")
        if code in TRANSIENT and attempt <= RETRIES:
            if staged:
                office_app("close-ours", "excel", *[str(p) for p in staged])
            sleep(RETRY_WAIT[min(attempt, len(RETRY_WAIT)) - 1])
            if session is not None:
                session.relaunch()
            print(f"   ↻ 一時的なエラー ({code}) → 自分の copy を閉じて再試行 ({attempt}/{RETRIES})")
            continue
        why = "一時的なエラーの再試行を使い切った" if code in TRANSIENT else "一時的なエラーでない = 再試行しない"
        raise ExcelError(f"Excel ({label}) が失敗 ({why}、 試行 {attempt} 回、 log = {ERROR_LOG}): {err}")
    raise AssertionError("unreachable")


# display alerts = false: 体裁の変更 (merge 等) が確認 dialog を出すと AppleEvent が返らず -1712 で止まる
# (probe で実測: merge で 200 秒 timeout → 再試行の open が -609)。 終わったら true に戻す。
EXPORT_SCRIPT = '''
on run argv
  set xlsxPath to item 1 of argv
  set pdfPath to item 2 of argv
  with timeout of 200 seconds
    tell application "Microsoft Excel"
      set display alerts to false
      set formcaseStep to "open"
      try
        delay 2
        set wbk to open workbook workbook file name (POSIX file xlsxPath)
        delay 3
        set formcaseStep to "ops"
        %OPS%
        set formcaseStep to "save as PDF"
        save as active sheet of wbk filename (POSIX file pdfPath) file format PDF file format with overwrite
        delay 2
        set formcaseStep to "close"
        close wbk saving no
      on error errMsg number errNum
        set display alerts to true
        error ("[" & formcaseStep & "] " & errMsg) number errNum
      end try
      set display alerts to true
    end tell
  end timeout
end run
'''


def export_pdfs(jobs) -> None:
    """jobs = [(xlsx, pdf)] か [(xlsx, pdf, ops_lines)]。 1 回の Excel session で順に PDF にする (= job ごとに Excel を
    起動・終了しない)。 ops_lines = ``ops_lines()`` の行 (staged copy に当て、 保存しない)。 workbook 全体の PDF になる。"""
    jobs = [(str(j[0]), str(j[1]), list(j[2]) if len(j) > 2 and j[2] else []) for j in jobs]
    for _x, pdf, _o in jobs:
        if os.path.exists(pdf):
            os.remove(pdf)
    names = [os.path.basename(x) for x, _p, _o in jobs]
    if len(set(names)) != len(names):
        raise ExcelError(f"同じ名前の workbook は 1 つの staging dir に置けない: {names}")
    with _ExcelSession() as ses, Stage(*[x for x, _p, _o in jobs]) as st:
        for (x, pdf, ops), sx in zip(jobs, st.paths):
            spdf = os.path.join(st.dir, os.path.basename(pdf)) if st.active else pdf
            script = EXPORT_SCRIPT.replace("%OPS%", "\n        ".join(ops))
            run_osascript(script, [sx, spdf], f"PDF 化 {os.path.basename(x)}", session=ses, staged=[sx])
            if st.active and os.path.exists(spdf):
                shutil.copy2(spdf, pdf)
    for _x, pdf, _o in jobs:
        if not os.path.exists(pdf):
            raise ExcelError(f"Excel の PDF 出力に失敗 (osascript は成功したが file が無い): {pdf}")


def export_pdf(xlsx, pdf, ops=()) -> None:
    export_pdfs([(xlsx, pdf, list(ops))])


def _q(s: str) -> str:
    parts = str(s).replace("\\", "\\\\").replace('"', '\\"').split("\n")
    return " & linefeed & ".join(f'"{p}"' for p in parts)


_HALIGN = {"left": "horizontal align left", "center": "horizontal align center", "right": "horizontal align right",
           "general": "horizontal align general", "justify": "horizontal align justify",
           "distributed": "horizontal align distributed", "centerContinuous": "horizontal align center across selection"}
_VALIGN = {"top": "vertical alignment top", "center": "vertical alignment center", "bottom": "vertical alignment bottom",
           "justify": "vertical alignment justify", "distributed": "vertical alignment distributed"}


def _dollar(area: str) -> str:
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.utils import get_column_letter

    c0, r0, c1, r1 = range_boundaries(area.replace("$", ""))
    return f"${get_column_letter(c0)}${r0}:${get_column_letter(c1)}${r1}"


def ops_lines(sheet: str, ops) -> list:
    """体裁の変更 (layout.excel_ops の tuple) を AppleScript の行に。 未対応の変更は ExcelError (黙って落とさない)。

    ops の形: ("unmerge", range) / ("merge", range) / ("row_height", row, pt) / ("wrap", cell, bool) /
    ("halign", cell, name) / ("valign", cell, name) / ("font_size", cell, pt) / ("number_format", cell, fmt) /
    ("print_area", range) / ("one_page",) / ("black_and_white", bool) / ("hide_sheet", name)。"""
    ws = f"worksheet {_q(sheet)} of wbk"
    out = []
    for op in ops:
        k = op[0]
        if k == "unmerge":
            out.append(f'unmerge range "{op[1]}" of {ws}')
        elif k == "merge":
            out.append(f'merge range "{op[1]}" of {ws}')
        elif k == "row_height":
            out.append(f'set row height of range "{op[1]}:{op[1]}" of {ws} to {float(op[2])}')
        elif k == "wrap":
            out.append(f'set wrap text of range "{op[1]}" of {ws} to {"true" if op[2] else "false"}')
        elif k == "halign" and op[2] in _HALIGN:
            out.append(f'set horizontal alignment of range "{op[1]}" of {ws} to {_HALIGN[op[2]]}')
        elif k == "valign" and op[2] in _VALIGN:
            out.append(f'set vertical alignment of range "{op[1]}" of {ws} to {_VALIGN[op[2]]}')
        elif k == "font_size":
            out.append(f'set font size of font object of range "{op[1]}" of {ws} to {float(op[2])}')
        elif k == "number_format":
            out.append(f'set number format of range "{op[1]}" of {ws} to {_q(op[2])}')
        elif k == "print_area":
            out.append(f'set print area of page setup object of {ws} to "{_dollar(op[1])}"')
        elif k == "one_page":
            out += [f"set zoom of page setup object of {ws} to false",
                    f"set fit to pages wide of page setup object of {ws} to 1",
                    f"set fit to pages tall of page setup object of {ws} to 1"]
        elif k == "black_and_white":
            out.append(f'set black and white of page setup object of {ws} to {"true" if op[1] else "false"}')
        elif k == "hide_sheet":            # 素刷り (fidelity.blank_pdf): 他の sheet を刷らない = 非表示 (削除しない = 参照の数式を壊さない)
            out.append(f"set visible of worksheet {_q(op[1])} of wbk to sheet hidden")
        else:
            raise ExcelError(f"Excel に当てる体裁の変更に未対応: {op!r} ({sheet})")
    return out


def _serial(d) -> float:
    if isinstance(d, _dt.datetime):
        base = _dt.datetime(1899, 12, 30)
        return (d - base).total_seconds() / 86400.0
    return float((d - _dt.date(1899, 12, 30)).days)


def script_for(edits) -> str:
    """edits = [(sheet, cell, kind, value)]。 kind = text | number | date | textfmt | clear | formula | general | fontsize。

    formula = 数式を書く (出張報告書が出張願を参照する欄)、 general = 表示書式を General に (value 不要。
    書式 @ のままだと数式が文字列として入る = 実測)。 clear は merge の全域 (範囲で渡す)。"""
    lines = []
    for sheet, cell, kind, val in edits:
        r = f'range "{cell}" of worksheet "{sheet}" of wbk'
        if kind == "formula":
            lines.append(f"set formula of {r} to {_q(val)}")
            continue
        if kind == "general":
            lines.append(f'set number format of {r} to "General"')
            continue
        if kind == "fontsize":            # 長い値が結合セルで切れる欄 (spec の font_size)
            lines.append(f"set font size of font object of {r} to {float(val)}")
            continue
        if kind == "text":
            lines.append(f"set value of {r} to {_q(val)}")
        elif kind == "number":
            lines.append(f"set value of {r} to {val}")
        elif kind == "date":
            lines.append(f"set value of {r} to {_serial(val)}")
        elif kind == "textfmt":
            lines.append(f'set number format of {r} to "@"')
            lines.append(f"set value of {r} to {_q(val)}")
        elif kind == "clear":
            lines.append(f"clear contents of {r}")
        else:
            raise ValueError(f"未知の kind {kind!r} ({sheet}!{cell})")
    body = "\n      ".join(lines)
    return f'''
on run argv
  set xlsxPath to item 1 of argv
  with timeout of 300 seconds
    tell application "Microsoft Excel"
      delay 2
      set wbk to open workbook workbook file name (POSIX file xlsxPath)
      delay 3
      {body}
      calculate
      save wbk
      delay 2
      close wbk saving no
    end tell
  end timeout
end run
'''


def single_sheet_values(src, dst, sheet: str, pre_ops=()) -> None:
    """src の copy から sheet だけを残し、 その sheet の数式を計算済みの値に固定した workbook を dst に作る。

    相手に送る入力用 xlsx 用: 他 sheet を消すと参照の数式が #REF! になるので先に値へ。 Excel で作る
    (openpyxl で save すると標題の drawing が消える = guideline #openpyxl 派生 xlsx の事故)。
    pre_ops = 刷った PDF と同じ体裁 (``ops_lines`` の行) を先に当てる (= 講師が刷る xlsx と当方の PDF の体裁を揃える)。"""
    import openpyxl
    import warnings

    src, dst = str(src), str(dst)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(src)
    if sheet not in wb.sheetnames:
        raise RuntimeError(f"sheet {sheet!r} が {src} に無い")
    formulas = [c.coordinate for row in wb[sheet].iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    others = [n for n in wb.sheetnames if n != sheet]
    lines = list(pre_ops)
    lines += [f'set value of range "{c}" of ws to (get value of range "{c}" of ws)' for c in formulas]
    lines += [f"delete worksheet {_q(n)} of wbk" for n in others]
    body = "\n      ".join(lines)
    script = f'''
on run argv
  set xlsxPath to item 1 of argv
  with timeout of 300 seconds
    tell application "Microsoft Excel"
      delay 2
      set display alerts to false
      set wbk to open workbook workbook file name (POSIX file xlsxPath)
      delay 3
      set ws to worksheet {_q(sheet)} of wbk
      calculate
      {body}
      save wbk
      delay 2
      close wbk saving no
      set display alerts to true
    end tell
  end timeout
end run
'''
    shutil.copy2(src, dst)
    with _ExcelSession() as ses, Stage(dst) as st:
        sx = st.paths[0]
        run_osascript(script, [sx], f"入力用 xlsx {os.path.basename(dst)}", session=ses, staged=[sx])
        if st.active:
            st.copy_back(sx, dst)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        got = openpyxl.load_workbook(dst)
    left = [c.coordinate for row in got[sheet].iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")] if sheet in got.sheetnames else ["<sheet 無し>"]
    if got.sheetnames != [sheet] or left:
        raise RuntimeError(f"入力用 xlsx の作成が不完全 (sheet {got.sheetnames}、 数式の残り {left[:5]})")


def merge_full_ranges(xlsx, edits) -> list:
    """clear の対象が merge の中の cell なら merge の全域に置き換える。

    Excel の ``clear contents of range "R31"`` は R31 が merge (R31:AA32 等) の anchor でも**黙って効かない**
    (実測)。 全域を渡すと効く。"""
    import openpyxl
    import warnings

    if not any(k == "clear" for _s, _c, k, _v in edits):
        return list(edits)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(str(xlsx))
    out = []
    for sheet, cell, kind, val in edits:
        if kind == "clear" and ":" not in cell and sheet in wb.sheetnames:
            for mr in wb[sheet].merged_cells.ranges:
                if cell in mr:
                    cell = str(mr)
                    break
        out.append((sheet, cell, kind, val))
    return out


def write_cells(xlsx, edits) -> None:
    xlsx = str(xlsx)
    script = script_for(merge_full_ranges(xlsx, edits))
    with _ExcelSession() as ses, Stage(xlsx) as st:
        sx = st.paths[0]
        run_osascript(script, [sx], f"記入 {os.path.basename(xlsx)}", session=ses, staged=[sx])
        if st.active:
            st.copy_back(sx, xlsx)
