"""Word (docx) 様式の記入・gate・fingerprint・PDF の重ね書き (spec の ``meta.kind: docx``)。 使い方の正本 = form-case-pipeline.md #docx。

Excel 様式と違い、 docx 様式の値は 2 か所に分かれる:

1. **docx に打つ値** (spec の ``docx.fields``) = 記入 stub の ``FIELDS``。 fill は**毎回 雛形から** docx を作り直す
   (段落を足す欄 = mode ``append`` があるので、 記入済みの docx に書き足すと二重になる = 冪等にするため)。
2. **PDF に重ねる値** = 記入 stub の ``CHOICES`` (spec の ``docx.choices`` = ○で囲む選択肢) と ``TEXTS``
   (``docx.texts`` = 罫線の右に書く短い値)。 docx に打つと autofit の表の列幅が動いて頁がはみ出す (実測) ので
   PDF の上に重ねる。 fill が ``<docx の stem>.overlay.yaml`` に書き、 build が読む。

``docx.render`` = 値に依らない雛形の直し (1 字溢れる行の字の大きさ等、 各 entry に理由)。 fill が毎回当てる。
``docx.seal`` = 認印の位置。 build が print 版にだけ ``overlay-seal-pdf.py`` で重ねる (印影は設定の
``seal_image_cmd`` が毎回引く)。

欄の場所 (``at``) の書き方:
  {para: N}                         本文の N 番目の段落 (python-docx の document.paragraphs)
  {table: T, row: R, col: C}        表 T の R 行 C 列の cell (python-docx の tables[T].rows[R].cells[C])
  {table: T, row: R, col: C, paragraph: K}   その cell の K 番目の段落だけ
format = 値の前後に雛形の字を残す ("{value}　　印" = 氏名の後ろの「印」 を押印の位置の目印に)。
選択肢の kind = word (語を囲む) / paren (「（ ）語」 の括弧、 値は True) / near (見出しの行の語) / between (前の選択肢の右〜before の左)。
choice_groups = どれか min 個以上を選ぶ選択肢の組。 texts の kind = right_of_label (既定) / after_paren。
mode:
  text     段落 / cell の文字を値で置き換える (最初の run の書式を使う)
  append   cell の最後の段落の後ろに、 その段落の書式で段落を 1 つ足す (雛形の注記を残して値を下に書く欄)
"""
from __future__ import annotations

import copy
import hashlib
import subprocess
from pathlib import Path

import yaml

def _yaml_safe_load(stream):  # yaml.safe_load と同じ結果を C 版 (libyaml) で返す = 約 10 倍速 (2026-09-23)
    import yaml
    return yaml.load(stream, Loader=getattr(yaml, "CSafeLoader", yaml.SafeLoader))


KIND = "docx"
OVERLAY_SUFFIX = ".overlay.yaml"


class DocxFormError(Exception):
    pass


def is_docx(spec) -> bool:
    return str(((spec or {}).get("meta") or {}).get("kind") or "") == KIND


def conf(spec) -> dict:
    return (spec or {}).get("docx") or {}


def fields(spec) -> list:
    return list(conf(spec).get("fields") or [])


def choices(spec) -> list:
    return list(conf(spec).get("choices") or [])


def texts(spec) -> list:
    return list(conf(spec).get("texts") or [])


def overlay_path(docx_path) -> Path:
    p = Path(docx_path)
    return p.with_name(p.stem + OVERLAY_SUFFIX)


def pages(spec) -> int:
    return int(((spec or {}).get("meta") or {}).get("pages") or 1)


# ---------------------------------------------------------------------------
# docx の欄
# ---------------------------------------------------------------------------
def _doc(path):
    import docx

    return docx.Document(str(path))


def _cell(d, at):
    try:
        return d.tables[int(at["table"])].rows[int(at["row"])].cells[int(at["col"])]
    except (IndexError, KeyError) as e:
        raise DocxFormError(f"欄の場所 {at} が docx に無い ({type(e).__name__}) = 雛形が spec と違う") from e


def _paragraphs(d, at) -> list:
    if "table" in at:
        return list(_cell(d, at).paragraphs)
    try:
        return [d.paragraphs[int(at["para"])]]
    except (IndexError, KeyError) as e:
        raise DocxFormError(f"段落の場所 {at} が docx に無い = 雛形が spec と違う") from e


def _set_runs(par, text, size=None, bold=None):
    from docx.shared import Pt

    runs = par.runs or [par.add_run("")]
    runs[0].text = text
    if size:
        runs[0].font.size = Pt(float(size))
    if bold is not None:
        runs[0].font.bold = bold
    for r in runs[1:]:
        r.text = ""


def _add_par(cell, text, size=None):
    import docx

    src = cell.paragraphs[-1]
    new = copy.deepcopy(src._p)
    src._p.addnext(new)
    par = docx.text.paragraph.Paragraph(new, src._parent)
    _set_runs(par, text, size=size, bold=False)


def field_text(d, f, template_count=None) -> str:
    """欄の今の文字。 text = 段落 / cell の全文 (at に paragraph があればその段落)、 append = 雛形より後ろに足された段落
    (template_count = 雛形の段落数)。"""
    ps = _paragraphs(d, f["at"])
    if f.get("mode", "text") == "text" and "paragraph" in f["at"]:
        return ps[int(f["at"]["paragraph"])].text
    if f.get("mode", "text") == "append":
        return "\n".join(p.text for p in ps[int(template_count or 0):])
    return "\n".join(p.text for p in ps)


def apply_render(d, spec) -> None:
    for r in conf(spec).get("render") or []:
        ps = _paragraphs(d, r["at"])
        idx = r["at"].get("paragraph")
        targets = [ps[int(idx)]] if idx is not None else ps
        if r.get("size"):
            from docx.shared import Pt

            for p in targets:
                for run in p.runs:
                    run.font.size = Pt(float(r["size"]))


def write(spec, template: Path, values: dict, out: Path) -> None:
    """雛形から docx を作り直す (values = {field id: 文字}、 None の欄は雛形のまま)。"""
    d = _doc(template)
    apply_render(d, spec)
    for f in fields(spec):
        v = values.get(f["id"])
        if v is None:
            continue
        v = rendered(f, v)
        mode = f.get("mode", "text")
        if mode == "text":
            if "paragraph" in f["at"]:
                _set_runs(_paragraphs(d, f["at"])[int(f["at"]["paragraph"])], v, size=f.get("size"))
            elif "table" in f["at"]:
                cell = _cell(d, f["at"])
                _set_runs(cell.paragraphs[0], str(v), size=f.get("size"))
                for p in cell.paragraphs[1:]:
                    _set_runs(p, "")
            else:
                _set_runs(_paragraphs(d, f["at"])[0], str(v), size=f.get("size"))
        elif mode == "append":
            if "table" not in f["at"]:
                raise DocxFormError(f"{f['id']}: mode append は表の cell だけ")
            _add_par(_cell(d, f["at"]), str(v), size=f.get("size"))
        else:
            raise DocxFormError(f"{f['id']}: mode {mode!r} は text / append のどれでもない")
    d.save(str(out))


def rendered(f, v) -> str:
    """欄に打つ文字 = spec の format (例: 氏名の後ろに雛形の「印」 を残す) に値を入れたもの。"""
    return str(f["format"]).format(value=v) if f.get("format") else str(v)


def template_counts(spec, template: Path) -> dict:
    d = _doc(template)
    return {f["id"]: len(_paragraphs(d, f["at"])) for f in fields(spec) if f.get("mode") == "append"}


def readback(spec, template: Path, docx_path: Path, values: dict) -> list:
    """[(field id, 期待, 実際)] = 書いたはずの値と docx の文字が違う欄。"""
    d, counts = _doc(docx_path), template_counts(spec, template)
    bad = []
    for f in fields(spec):
        v = values.get(f["id"])
        if v is None:
            continue
        got = field_text(d, f, counts.get(f["id"]))
        if got != rendered(f, v):
            bad.append((f["id"], rendered(f, v), got))
    return bad


# ---------------------------------------------------------------------------
# gate (記入内容)
# ---------------------------------------------------------------------------
def load_overlay(docx_path) -> dict:
    p = overlay_path(docx_path)
    if not p.exists():
        return {}
    return _yaml_safe_load(p.read_text(encoding="utf-8")) or {}


def gate(spec, template: Path, docx_path: Path) -> tuple:
    """(ok, lines)。 docx の欄と overlay.yaml の選択肢・短い値を spec と照合する。

    - 必須の欄 (optional でない) が雛形のまま / 空 = 未記入
    - fixed の欄が fixed の値でない
    - 選択肢の値が options に無い / 必須の選択肢が未選択
    - 必須の短い値が空"""
    d, t = _doc(docx_path), _doc(template)
    counts = template_counts(spec, template)
    lines, ok = [], True
    for f in fields(spec):
        fid = f["id"]
        got = field_text(d, f, counts.get(fid))
        tpl = field_text(t, f, counts.get(fid))
        label = f.get("label", fid)
        if "fixed" in f and got != str(f["fixed"]):
            lines.append(f"🔴 {fid} ({label}): {got!r} ≠ 固定値 {f['fixed']!r}")
            ok = False
        elif not f.get("optional") and (got == tpl or not got.strip()):
            lines.append(f"🔴 {fid} ({label}): 未記入 (雛形のまま)")
            ok = False
        elif f.get("optional") and got == tpl:
            lines.append(f"🧊 {fid} ({label}): 空欄 (任意)")
        else:
            lines.append(f"✅ {fid} ({label}): {got[:40]!r}" + ("…" if len(got) > 40 else ""))
    ov = load_overlay(docx_path)
    for c in choices(spec):
        cid, val = c["id"], (ov.get("choices") or {}).get(c["id"])
        label = c.get("label", cid)
        opts = c.get("options")
        if val in (None, "", False) and c.get("kind") == "paren":
            val = None
        if val is None:
            if c.get("optional"):
                lines.append(f"🧊 {cid} ({label}): 選ばない (任意)")
            else:
                lines.append(f"🔴 {cid} ({label}): 未選択")
                ok = False
        elif opts is not None and val not in opts:
            lines.append(f"🔴 {cid} ({label}): {val!r} は選択肢 {opts} に無い")
            ok = False
        else:
            lines.append(f"✅ {cid} ({label}): ○ {val!r}")
    for x in texts(spec):
        xid, val = x["id"], (ov.get("texts") or {}).get(x["id"])
        if val in (None, ""):
            if x.get("optional"):
                lines.append(f"🧊 {xid} ({x.get('label', xid)}): 空欄 (任意)")
            else:
                lines.append(f"🔴 {xid} ({x.get('label', xid)}): 未記入")
                ok = False
        else:
            lines.append(f"✅ {xid} ({x.get('label', xid)}): 重ねる値あり")
    for grp in conf(spec).get("choice_groups") or []:
        picked = [i for i in grp["ids"] if (ov.get("choices") or {}).get(i) not in (None, "", False)]
        if len(picked) < int(grp.get("min", 1)):
            lines.append(f"🔴 {grp.get('label', grp['ids'])}: どれも選ばれていない ({grp['ids']} から {grp.get('min', 1)} 個以上)")
            ok = False
    for kind, ids in (("choices", {c["id"] for c in choices(spec)}), ("texts", {x["id"] for x in texts(spec)})):
        extra = set((ov.get(kind) or {}).keys()) - ids
        if extra:
            lines.append(f"🔴 overlay.yaml の {kind} に spec に無い id {sorted(extra)}")
            ok = False
    return ok, lines


# ---------------------------------------------------------------------------
# fingerprint (凍結の記録)
# ---------------------------------------------------------------------------
def _w(tag):
    return "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag


def docx_digest(path) -> dict:
    """{"values": 本文の文字の digest, "format": 紙の見た目を変える書式の digest}。

    values = 本文の段落と表の cell の文字を文書の順に (結合 cell は 1 回)。
    format = 段落ごとの揃え・style と run ごとの字の大きさ・太字 + 表の列幅 (gridCol) + 用紙・余白 (sectPr)。
    Word が再保存で書き換える属性 (rsid 等) は含めない。"""
    import zipfile

    from lxml import etree

    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(_w("body"))
    vals, fmt = hashlib.sha256(), hashlib.sha256()

    def ptext(p):
        return "".join(t.text or "" for t in p.iter(_w("t")))

    def pformat(p):
        ppr = p.find(_w("pPr"))
        jc = ppr.find(_w("jc")) if ppr is not None else None
        st = ppr.find(_w("pStyle")) if ppr is not None else None
        runs = []
        for r in p.iter(_w("r")):
            rpr = r.find(_w("rPr"))
            sz = rpr.find(_w("sz")) if rpr is not None else None
            b = rpr.find(_w("b")) if rpr is not None else None
            runs.append(f"{sz.get(_w('val')) if sz is not None else ''}/{'b' if b is not None else ''}")
        return (f"p\t{jc.get(_w('val')) if jc is not None else ''}\t{st.get(_w('val')) if st is not None else ''}\t"
                + ",".join(runs))

    for el in body:
        if el.tag == _w("p"):
            vals.update(f"p\t{ptext(el)}\n".encode("utf-8"))
            fmt.update((pformat(el) + "\n").encode("utf-8"))
        elif el.tag == _w("tbl"):
            grid = el.find(_w("tblGrid"))
            if grid is not None:
                fmt.update(("grid\t" + ",".join(g.get(_w("w")) or "" for g in grid.findall(_w("gridCol")))
                            + "\n").encode("utf-8"))
            for ri, tr in enumerate(el.iter(_w("tr"))):
                for ci, tc in enumerate(tr.findall(_w("tc"))):
                    txt = "\n".join(ptext(p) for p in tc.findall(_w("p")))
                    vals.update(f"tc\t{ri}\t{ci}\t{txt}\n".encode("utf-8"))
                    for p in tc.findall(_w("p")):
                        fmt.update((f"tc\t{ri}\t{ci}\t" + pformat(p) + "\n").encode("utf-8"))
        elif el.tag == _w("sectPr"):
            for tag in ("pgSz", "pgMar"):
                x = el.find(_w(tag))
                if x is not None:
                    fmt.update((tag + "\t" + "\t".join(f"{k.split('}')[-1]}={v}" for k, v in sorted(x.attrib.items()))
                                + "\n").encode("utf-8"))
    return {"values": vals.hexdigest()[:24], "format": fmt.hexdigest()[:24]}


def overlay_digest(docx_path) -> str | None:
    p = overlay_path(docx_path)
    return hashlib.sha256(p.read_bytes()).hexdigest()[:24] if p.exists() else None


# ---------------------------------------------------------------------------
# PDF: Word で PDF にし、 選択肢の ○・短い値を重ね、 押印の位置を出す
# ---------------------------------------------------------------------------
CC = Path(__file__).resolve().parent.parent          # 同じ repo の scripts/ (overlay-seal-pdf.py 等)
DOCX2PDF = CC / "docx-to-pdf.sh"


def to_pdf(docx_path: Path, pdf: Path) -> Path:
    r = subprocess.run(["bash", str(DOCX2PDF), str(docx_path), str(pdf)], stdout=subprocess.PIPE,
                       stderr=subprocess.STDOUT, text=True)
    if r.returncode != 0 or not pdf.exists():
        raise DocxFormError("Word で PDF にできない (docx-to-pdf.sh): " + r.stdout[-800:])
    return pdf


def _font():
    r = subprocess.run(["kpsewhich", "HaranoAjiGothic-Regular.otf"], capture_output=True, text=True)
    return r.stdout.strip() or None


def _first(page, key, below=None, above=None):
    rs = page.search_for(key)
    if below is not None:
        rs = [r for r in rs if r.y0 > below]
    if above is not None:
        rs = [r for r in rs if r.y0 < above]
    rs = sorted(rs, key=lambda r: (r.y0, r.x0))
    if not rs:
        raise DocxFormError(f"PDF に「{key}」 が見つからない = 雛形か Word の変換が spec と違う")
    return rs[0]


def overlay(pdf_in: Path, pdf_out: Path, spec, ov: dict) -> None:
    """選択肢の ○ と短い値を重ねる (押印はしない = 確認用はこの PDF そのもの)。"""
    import sys

    import fitz

    sys.path.insert(0, str(CC))
    from pdf_form_fill import circle_paren_gap, circle_word, put_value_right_of_label

    d = fitz.open(str(pdf_in))
    chosen = {}
    for c in choices(spec):
        val = (ov.get("choices") or {}).get(c["id"])
        if val in (None, "", False):
            continue
        p = d[int(c.get("page", 1)) - 1]
        kind = c.get("kind", "word")
        if kind == "word":
            chosen[c["id"]] = circle_word(p, str(val))
        elif kind == "paren":
            chosen[c["id"]] = circle_paren_gap(p, str(c["anchor"]))
        elif kind == "near":
            lab = _first(p, c["label_text"])
            chosen[c["id"]] = circle_word(p, str(val), near=lab, y_tol=float(c.get("y_tol", 8.0)))
        elif kind == "between":
            lab = _first(p, c["label_text"])
            left = chosen.get(c["after"])
            if left is None:
                raise DocxFormError(f"{c['id']}: 先に {c['after']} を選ぶ (その語の右から探す)")
            stops = [r.x0 for r in p.search_for(c["before"]) if abs(r.y0 - lab.y0) < 8 and r.x0 > left.x1 + 20]
            if not stops:
                raise DocxFormError(f"{c['id']}: 行に「{c['before']}」 が見つからない")
            chosen[c["id"]] = circle_word(p, str(val), near=lab, x_min=left.x1, x_max=min(stops), pick="unique")
        else:
            raise DocxFormError(f"{c['id']}: kind {kind!r} は word / paren / near / between のどれでもない")
    tv = ov.get("texts") or {}
    if any(tv.get(x["id"]) for x in texts(spec)):
        font = _font()
        for x in texts(spec):
            val = tv.get(x["id"])
            if not val:
                continue
            p = d[int(x.get("page", 1)) - 1]
            if font and "hag" not in [f[4] for f in p.get_fonts()]:
                p.insert_font(fontname="hag", fontfile=font)
            fname = "hag" if font else "japan"
            size = float(x.get("size", 10))
            lab = _first(p, x["label_text"])
            if x.get("kind", "right_of_label") == "after_paren":
                paren = [r for r in p.search_for("（") if abs(r.y0 - lab.y0) < 3 and r.x0 > lab.x1]
                if not paren:
                    raise DocxFormError(f"{x['id']}: 「{x['label_text']}」 の行に（ が無い")
                p.insert_text((paren[0].x1 + float(x.get("dx", 9)), lab.y1 - 1.8), str(val), fontsize=size, fontname=fname)
                continue
            # label の右の縦罫線を anchor、 数字は行の縦中心 (層1 helper = office-automation.md#pdf-overlay-anchoring)
            put_value_right_of_label(p, x["label_text"], str(val), fontname=fname, size=size, cjk=bool(x.get("cjk")))
    d.save(str(pdf_out), garbage=3, deflate=True)


def seal_places(pdf: Path, spec) -> dict:
    """{page: place spec (page= を除く)} = 層1 overlay-seal-pdf.py の --place。 spec の docx.seal:
    anchor (印の字) / pick (rightmost) / above_label (この語より上の候補だけ) / dx (印字の中心からの横のずれ pt) / size (インク径 pt)。"""
    import fitz

    s = conf(spec).get("seal")
    if not s:
        return {}
    page_no = int(s.get("page", 1))
    d = fitz.open(str(pdf))
    p = d[page_no - 1]
    hits = p.search_for(s["anchor"])
    cand = list(hits)
    if s.get("above_label"):
        lim = _first(p, s["above_label"]).y0
        cand = [r for r in cand if r.y0 < lim]
    if not cand:
        raise DocxFormError(f"押印の「{s['anchor']}」 が見つからない")
    target = sorted(cand, key=lambda r: -r.x0)[0] if s.get("pick", "rightmost") == "rightmost" else cand[0]
    occurrence = next(i for i, r in enumerate(hits, 1) if r == target)
    size = float(s.get("size", 27))
    # engine: 画像の左端 = anchor の右端 + dx、 インク径 = size → インクの中心 = anchor.x1 + dx + size/2
    dx = (target.x0 + target.x1) / 2 + float(s.get("dx", 0)) - target.x1 - size / 2
    return {page_no: f"anchor={s['anchor']},occurrence={occurrence},size={size:g},dx={dx:.2f},dy={float(s.get('dy', 0)):g},avoid=none"}
