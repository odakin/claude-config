"""engine の内蔵 fixture test (Office 不要)。 ``python3 formcase.py --selftest``。

fixture は全部合成 (架空の様式 ``fx`` / sheet A…E / 架空の人名) で、 呼び元の spec も案件も読まない
(= engine の test が instance に依存しない)。 instance 固有の test (実 spec・実雛形に対する照合) は
設定の ``instance_selftest`` が指す file の ``run(expect, tmp)`` が担う。
"""
from __future__ import annotations

import contextlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

from . import check as C
from . import config as CF
from . import guard as G
from . import lifecycle as L
from . import manifest as M
from . import specs as S

MAIN, PLAN, REQ, REPORT, EXTRA = "sheet A", "sheet P", "sheet R", "sheet B", "sheet X"

# ---------------------------------------------------------------------------
# 合成の instance (spec / 雛形 / 設定) — engine の test はこれだけを見る
# ---------------------------------------------------------------------------
SPEC = {
    "meta": {"id": "fx", "sheet": MAIN, "template": "templates/fx.xlsx", "title": "例の様式",
             "form": "fx-template", "history_home": "../notes/history.md"},
    "groups": {"g1": {"sheets": [MAIN, PLAN, REQ], "pages": [1, 2, 3]},
               "g2": {"sheets": [REPORT], "depends": [MAIN], "pages": [4]}},
    "render": [
        {"sheet": MAIN, "ref": "I16", "number_format": "h:mm", "why": "雛形の表示書式が時刻でない"},
        {"sheet": MAIN, "row": 37, "row_height": 15, "why": "雛形の label が行高に収まらない"},
        {"sheet": MAIN, "print_area": "A1:AH60", "why": "雛形に印刷範囲が無い"},
        {"sheet": REPORT, "ref": "L35:X35", "merge": "L35:X35", "why": "雛形の結合が抜けている"},
        {"sheet": REPORT, "ref": "J54:V60", "merge": "J54:V60", "when": {"empty": "J57"},
         "why": "下段が空なら 1 つの欄にする"},
        {"sheet": REPORT, "ref": "J54:V56", "merge": "J54:V56", "when": {"filled": "J57"},
         "why": "下段に値があれば 2 段"},
        {"sheet": REPORT, "ref": "M29", "vertical": "center", "why": "label が上に寄る"},
        {"sheet": REPORT, "black_and_white": True, "why": "色刷りにしない"},
    ],
    "cells": [
        {"ref": "AA11", "label": "氏名", "state": "filled", "rule": "name", "group": "g1",
         "summary": "氏名は登録の表記で書く",
         "claims": [{"about": ["氏名"], "forbid": ["略称"], "says": "氏名は略称で書かない"}]},
        {"ref": "I37", "label": "職名", "state": "filled", "rule": "shokumei", "group": "g1",
         "summary": "職名は登録の区分で書く",
         "claims": [{"about": ["職名"], "forbid": ["甲種"], "says": "職名に旧区分を書かない"}]},
        {"ref": f"{REQ}!Z8", "label": "立替", "state": "fixed", "value": "○甲案立替", "rule": "tatekae",
         "group": "g1", "summary": "立替は ○甲案立替 を選ぶ",
         "superseded": ["立替は ○乙案立替 を選ぶ"],
         "claims": [{"about": ["立替"], "forbid": ["乙案"], "says": "立替の選択は甲案"}]},
        {"ref": f"{REQ}!AG38", "label": "請求", "state": "fixed", "value": "☑", "rule": "seikyuu",
         "group": "g1", "summary": "請求の欄は ☑ にする",
         "claims": [{"about": ["請求"], "negation": True, "says": "請求の欄は空で出さない"}]},
        {"ref": f"{REQ}!M38", "label": "日付欄", "state": "empty", "rule": "hizuke-empty", "group": "g1",
         "summary": "日付欄は空で出す (窓口が書く)",
         "claims": [{"about": ["日付欄"], "affirm": True, "says": "日付欄に当方が書かない"}]},
        {"ref": f"{REQ}!D11", "label": "区分", "state": "filled", "rule": "kubun", "group": "g1",
         "font_size": 7, "summary": "区分は正式な語で書く",
         "claims": [{"about": ["区分"], "forbid_re": ["(?<!研究)協力者"], "says": "区分を略さない"}]},
        {"ref": f"{REQ}!V11", "label": "番号", "state": "fixed", "value": "12345678", "rule": "bangou",
         "group": "g1", "summary": "番号は登録値を文字列で書く",
         "claims": [{"about": ["番号"], "forbid": ["99999999"], "says": "番号は登録値"}]},
        {"ref": f"{REPORT}!X10", "label": "用務", "state": "filled", "rule": "youmu", "group": "g2",
         "summary": "用務は目的が分かる語で書く",
         "claims": [{"about": ["用務"], "negation": True, "says": "用務は空で出さない"}]},
        {"ref": f"{REPORT}!AC37", "label": "合計", "state": "fixed", "value": '=IF(A1="","",A1)',
         "rule": "goukei", "group": "g2", "summary": "合計は数式で入れる",
         "claims": [{"about": ["合計"], "negation": True, "says": "合計を空にしない"}]},
        {"ref": f"{REPORT}!L25", "label": "内訳", "state": "changed", "rule": "uchiwake", "group": "g2",
         "summary": "雛形の値のままにしない", "why": "雛形の例が残る",
         "claims": [{"about": ["内訳"], "forbid": ["雛形のまま"], "says": "雛形の値を残さない"}]},
        {"ref": f"{REQ}!U47", "label": "同意", "state": "checkbox", "rule": "doui", "group": "g1",
         "summary": "同意の欄は ☑選ばない を消す",
         "markers": ["同意の欄は ☑選ばない を消す"],
         "claims": [{"about": ["同意"], "forbid": ["☑選ばない"], "says": "同意は片方だけ"}]},
        {"ref": "U10", "label": "経路", "state": "empty", "rule": "keiro", "group": "g1",
         "summary": "経路は空で出す",
         "claims": [{"about": ["経路"], "affirm": True, "says": "経路に当方が書かない"}]},
        {"ref": "AB5", "label": "種別", "state": "fixed", "value": "甲", "rule": "shubetsu", "group": "g1",
         "summary": "種別は 甲 に固定",
         "claims": [{"about": ["種別"], "forbid": ["乙"], "says": "種別は甲"}]},
    ],
    "nittei": {"sheet": PLAN, "anchors": [10, 14, 18],
               "roles": {"date": {"col": "C"}, "duty": {"col": "I"}, "place": {"col": "U"},
                         "days": {"col": "AF"}},
               "layout": {"collapsed_row_max_pt": 5.0},
               "rules": [{"id": "days", "label": "滞在日数",
                          "summary": "滞在日数は 0.5 刻みで書く",
                          "superseded": ["滞在日数は 0.25 刻み"],
                          "markers": ["滞在日数は 0.5 刻みで書く"],
                          "claims": [{"about": ["滞在日数"], "value": "0.25", "says": "刻みは 0.5"}]},
                         {"id": "tanka", "label": "単価", "summary": "単価は 5,000 円",
                          "claims": [{"about": ["単価"], "other_numbers": ["5000"], "unit": "円",
                                      "says": "単価は 5,000 円"}]}]},
    "cross_checks": [{"id": "days-sum", "expr": "sum(days) == nights", "same_as": "fx/days"},
                     {"id": "order", "summary": "先の書類を出してから次の書類を出す",
                      "claims_na": "値でなく手順の順番"}],
}


def _instance(tmp: Path) -> Path:
    """合成の instance dir (spec / 雛形 / 設定 / ack) を作り、 engine をそこに configure する。"""
    import openpyxl
    import yaml

    inst = tmp / "inst"
    (inst / "reference").mkdir(parents=True)
    (inst / "templates").mkdir()
    (inst / "notes").mkdir()
    (inst / "notes" / "history.md").write_text("# 経緯の home\n", encoding="utf-8")
    (inst / "reference" / "fx.yaml").write_text(yaml.safe_dump(SPEC, allow_unicode=True, sort_keys=False),
                                                encoding="utf-8")
    (inst / "lint-ack.yaml").write_text("acks: []\n", encoding="utf-8")
    # 配布雛形 (= 合成): 印刷範囲・手動改ページ・結合セル・長い label
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MAIN
    ws["I16"] = "9:00"
    ws["A37"] = "とても長い label が 1 行に入らない欄"
    ws.merge_cells("G17:H18")            # 長い値で行が伸びる欄 (狭い)
    ws.merge_cells("G20:AH21")           # 短い値の欄 (広い = 伸びない)
    ws.column_dimensions["G"].width = 8.0
    ws.column_dimensions["H"].width = 8.0
    for r in (17, 18, 20, 21):
        ws.row_dimensions[r].height = 13.5
    from openpyxl.worksheet.pagebreak import Break

    ws.row_breaks.append(Break(id=30))
    for name in (PLAN, REQ, REPORT, EXTRA):
        s = wb.create_sheet(name)
        s.print_area = "A1:AH60"
        if name == PLAN:
            for a in (10, 14, 18):
                s[f"C{a}"] = "：\n～"
                s.merge_cells(f"I{a}:T{a + 1}")
                for r in range(a, a + 4):
                    s.row_dimensions[r].height = 13.5
        if name == REPORT:
            s["J54"] = "上段"
            s["L35"] = "内訳"
            s["M29"] = "label"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb.save(inst / "templates" / "fx.xlsx")
    cfg = {
        "workspace_root": str(tmp), "workspace_label": "<ws>/",
        "spec_dir": "reference", "template_root": ".",
        "case_roots": ["repo-a/docs", "repo-b/docs"],
        "case_label_with_parent": ["sub"],
        "repos": ["repo-a", "repo-b"],
        "usage_doc": "<usage doc>", "cli": "formcase.py", "stub_sys_path": str(inst),
        "precedent_doc": "layer1/conventions/office-automation.md#template-base-not-precedent-base",
        "gates": [], "default_gates": [], "gates_by_form": {}, "seal_image_cmd": [],
        "recipes": [], "instance_selftest": "",
        "lint": {"ack": "lint-ack.yaml",
                 "targets": [{"glob": "repo-a/docs/*/README.md"},
                             {"glob": "repo-a/docs/*/notes*.md", "require_sibling": "submission.yaml"},
                             {"glob": "repo-a/docs/reference/*.md"},
                             {"glob": "repo-b/docs/*/sub/README.md"}],
                 "case_readme": [{"glob": "repo-a/docs/*/README.md", "name_re": r"^\d{4}-",
                                  "or_sibling": "submission.yaml"},
                                 {"glob": "repo-b/docs/*/sub/README.md"}],
                 "forms_by_path": [{"glob": "repo-b/**", "forms": ["fx"]}],
                 "context_stopwords": ["共通"]},
        "views": {"roots": ["repo-a/docs"], "files": []},
        "scaffold": {"spec_hint": "reference/{spec}", "process_hint": "手順 = <手順 doc>",
                     "derived_workbook_suffix": {"g2": "_second"}},
    }
    import json

    (inst / CF.CONFIG_NAME).write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    CF.configure(inst / CF.CONFIG_NAME)
    return inst


def _fixture(root: Path) -> Path:
    import openpyxl

    case = root / "repo-a" / "docs" / "2026-01-01-case-a"
    case.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MAIN
    ws["AA11"] = "甲野 太郎"
    for name, cell, val in ((PLAN, "U10", "出発地（架空市）→ 目的地"), (REQ, "Z8", 46200),
                            (REQ, "U47", "☑選ばない"), (REPORT, "L23", "架空大学")):
        wb.create_sheet(name)[cell] = val
    from openpyxl.styles import Border, Side

    plan = wb[PLAN]
    plan.print_area = "A1:AJ60"
    plan.sheet_properties.pageSetUpPr.fitToPage = True
    plan["V20"].border = Border(bottom=Side(style="thin"))     # 値の無い cell の罫線 (印刷範囲の中)
    plan.row_dimensions[12].height = 14.25
    plan.column_dimensions["C"].width = 3.0
    plan.column_dimensions["D"].width = 1.25
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb.save(case / "book.xlsx")
    (case / "g1.pdf").write_bytes(b"%PDF-g1")
    (case / "g2.pdf").write_bytes(b"%PDF-g2")
    (case / "gen.py").write_text("from formcase.guard import legacy_guard; legacy_guard(__file__)\n")
    (case / "other.py").write_text("print(1)\n")
    data = {"schema": M.SCHEMA, "case": case.name,
            "documents": {"d1": {"form": "fx", "workbook": "book.xlsx",
                                 "drivers": {"gen.py": ["g1", "g2"]},
                                 "groups": {"g1": {"current": {"state": "draft", "outputs": {"print": "g1.pdf"}}},
                                            "g2": {"current": {"state": "draft", "outputs": {"print": "g2.pdf"}}}}}}}
    (case / M.MANIFEST_NAME).write_text(M.dump_text(data), encoding="utf-8")
    return case


def _mark_mac(path: Path) -> None:
    """fixture の docProps/app.xml を「最後の保存 = Mac Excel」 にする (freeze の前提。 openpyxl は自分の名前を書く)。"""
    import io as _io
    import re
    import zipfile

    zin = zipfile.ZipFile(path)
    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == "docProps/app.xml":
                d = re.sub(rb"<Application>[^<]*</Application>",
                           b"<Application>Microsoft Macintosh Excel</Application>", d)
            zout.writestr(it, d)
    zin.close()
    Path(path).write_bytes(buf.getvalue())


def _anchor_xml(name: str, text: str, cx: int = 100, run: str = "<a:r><a:t>{t}</a:t></a:r>") -> str:
    return ('<xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row>'
            f'<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="{cx}" cy="100"/><xdr:sp><xdr:nvSpPr>'
            f'<xdr:cNvPr id="1" name="{name}"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/>'
            f'<a:p>{run.format(t=text)}</a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:oneCellAnchor>')


def _inject_drawing(path: Path, sheet_part: str, text: str, anchors: str | None = None) -> None:
    """sheet に図形 (textbox) を足す (openpyxl が落とすものの代表)。 anchors = wsDr の中身 (既定 = text の 1 つ)。"""
    import io as _io
    import zipfile

    draw = ('<?xml version="1.0"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"'
            ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            + (anchors if anchors is not None else _anchor_xml("t", text)) + '</xdr:wsDr>')
    rels = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdD1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"'
            ' Target="../drawings/drawing9.xml"/></Relationships>')
    zin = zipfile.ZipFile(path)
    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == f"xl/worksheets/{sheet_part}":
                d = d.replace(b"</worksheet>", b'<drawing r:id="rIdD1"/></worksheet>')
                if b'xmlns:r=' not in d:
                    d = d.replace(b"<worksheet ", b'<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                                                 b'officeDocument/2006/relationships" ', 1)
            zout.writestr(it, d)
        zout.writestr("xl/drawings/drawing9.xml", draw)
        zout.writestr(f"xl/worksheets/_rels/{sheet_part}.rels", rels)
    zin.close()
    Path(path).write_bytes(buf.getvalue())


def _sheet_part(path: Path, name: str) -> str:
    import zipfile

    from . import fingerprint as FP

    with zipfile.ZipFile(path) as z:
        return FP._sheet_parts(z)[name].rsplit("/", 1)[-1]


def _set(case: Path, sheet: str, cell: str, val):
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(case / "book.xlsx")
        wb[sheet][cell] = val
        wb.save(case / "book.xlsx")


# ---------------------------------------------------------------------------
def _synth_font(dir_: Path, name: str, ascent: int, descent: int, line_gap: int) -> None:
    """hhea の行の高さだけを持つ最小の TTF (1000 upem、 glyph は .notdef と 'a')。 fontTools が要る。"""
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen

    fb = FontBuilder(1000, isTTF=True)
    fb.setupGlyphOrder([".notdef", "a"])
    fb.setupCharacterMap({0x61: "a"})
    pen = TTGlyphPen(None)
    pen.moveTo((50, 0)); pen.lineTo((550, 0)); pen.lineTo((550, 700)); pen.lineTo((50, 700)); pen.closePath()
    fb.setupGlyf({".notdef": TTGlyphPen(None).glyph(), "a": pen.glyph()})
    fb.setupHorizontalMetrics({".notdef": (500, 0), "a": (600, 50)})
    fb.setupHorizontalHeader(ascent=ascent, descent=descent, lineGap=line_gap)
    fb.setupNameTable({"familyName": name, "styleName": "Regular"})
    fb.setupOS2(sTypoAscender=ascent, sTypoDescender=descent, usWinAscent=ascent, usWinDescent=-descent)
    fb.setupPost()
    fb.save(str(dir_ / (name.replace(" ", "-") + ".ttf")))


def _synth_font_checks(expect) -> None:
    """line_pitch が font の行の高さから出ることを、 合成 font で (機械の font に依らず) 確かめる。"""
    from . import layout as LY

    try:
        import fontTools.fontBuilder  # noqa: F401
    except ImportError:
        print("  SKIP 行の送り (fontTools が無い = font の行の高さは読めず、 既定 LINE_PITCH に倒れる経路だけ見る)")
        expect("行の送り: font が読めなければ既定 LINE_PITCH", LY.line_pitch("ＭＳ Ｐゴシック") == LY.LINE_PITCH)
        return
    tall, flat = "Synth Tall Gothic", "Synth Flat Gothic"
    with tempfile.TemporaryDirectory() as td:
        fdir = Path(td)
        _synth_font(fdir, tall, 1160, -442, 0)   # 1.602 em × LINE_GAP 1.04 = 1.666 (> LINE_PITCH 1.3)
        _synth_font(fdir, flat, 880, -120, 0)    # 1.0 em × 1.04 = 1.04 (< LINE_PITCH → 既定が勝つ)
        saved = (LY.DFONTS, LY.FONT_FILES)
        LY.DFONTS, LY.FONT_FILES = fdir, tuple(p.name for p in sorted(fdir.iterdir()))
        for fn in (LY._fonts, LY._advance, LY.line_pitch):
            fn.cache_clear()
        try:
            expect("行の送り: 行の高い font (合成、 hhea 1.602 em = 游ゴシック相当) は MS 系より高い (font の行の高さから)",
                   LY.line_pitch(tall) > 1.5, LY.line_pitch(tall))
            expect("行の送り: 行の高さ 1 em の font (合成、 MS 系相当) は既定 LINE_PITCH に留まる",
                   LY.line_pitch(flat) == LY.LINE_PITCH, LY.line_pitch(flat))
            expect("行の送り: 無い font は既定 LINE_PITCH", LY.line_pitch("ＭＳ Ｐゴシック") == LY.LINE_PITCH)
            expect("行の送り: 字の幅も合成 font の hmtx から (a = 0.6 em)", LY._advance(tall, "a") == 0.6, LY._advance(tall, "a"))
        finally:
            LY.DFONTS, LY.FONT_FILES = saved
            for fn in (LY._fonts, LY._advance, LY.line_pitch):
                fn.cache_clear()
    if (LY.DFONTS / "YuGothR.ttc").exists():
        expect("行の送り: 実機の 游ゴシック (Excel 同梱) は MS 系より高い",
               LY.line_pitch("游ゴシック") > 1.5 and LY.line_pitch("ＭＳ Ｐゴシック") == LY.LINE_PITCH)
    else:
        print("  SKIP 行の送り: Excel 同梱の font が無い機械 (実機 font の項は合成 font で代替済)")


def run() -> int:
    ok = True

    def expect(label, cond, detail=""):
        nonlocal ok
        print(f"  {'✅' if cond else '❌'} {label}" + (f"  ({detail})" if detail and not cond else ""))
        ok &= bool(cond)

    def fails(case):
        return [x for x in C.check_case(M.load(case)) if x[0] == M.FAIL]

    tmp = Path(tempfile.mkdtemp(prefix="formcase-selftest-"))
    saved = dict(CF._STATE)
    try:
        inst = _instance(tmp)
        expect("合成 spec が読める (規則・claims の網羅)", S.get("fx") is not None)
        case = _fixture(tmp)
        expect("fixture manifest は FAIL なし", not fails(case), fails(case))
        _freeze_tests(tmp, case, expect, fails)
        _format_tests(tmp, case, expect, fails)
        _guard_tests(tmp, case, expect, fails)
        _annotate_tests(tmp, case, expect, fails)
        _lifecycle_tests(tmp, case, expect, fails)
        _precommit_tests(tmp, case, expect)
        _exit_code_tests(expect)
        _recipe_tests(tmp, inst, expect)
        _excel_tests(tmp, expect)
        _layout_tests(tmp, expect)
        _view_lint_tests(tmp, expect)
        _case_readme_tests(tmp, expect)
        _value_from_tests(tmp, expect)
        _copy_page_tests(tmp, expect)
        _graft_tests(tmp, expect)
        from .selftest_docx import run_docx_tests, run_page_role_tests

        run_docx_tests(tmp, expect)
        run_page_role_tests(tmp, expect)
        if saved.get("cfg") is not None:        # 呼び元の設定に戻してから instance の selftest を探す
            CF._STATE.update(saved)
            CF._invalidate()
            inst_st = CF.instance_selftest()
        else:
            inst_st = None
        if inst_st:
            print(f"── instance の selftest ({inst_st})")
            import importlib.util

            sp = importlib.util.spec_from_file_location("formcase_instance_selftest", inst_st)
            mod = importlib.util.module_from_spec(sp)
            sys.modules[sp.name] = mod
            sp.loader.exec_module(mod)
            mod.run(expect, tmp)
    finally:
        CF._STATE.update(saved)
        CF._invalidate()
        shutil.rmtree(tmp, ignore_errors=True)
    print("selftest:", "ALL PASS" if ok else "FAILED")
    return 0 if ok else 1


def _copy_page_tests(tmp, expect) -> None:
    """recipes.with_copy_page = 当事者に送る 1 本 (#recipient-send-version): 1 頁だけ白黒の raster に差し替える。"""
    import fitz

    from .recipes import PP, BuildError, with_copy_page

    plain = tmp / "cp_plain.pdf"
    d = fitz.open()
    for t in ("合成 旅費請求書", "合成 依頼書", "合成 承諾書"):
        d.new_page().insert_text((72, 72), t, fontname="japan")
    PP.write_record(d, [{"role": "submit", "label": t} for t in ("a", "b", "c")], "selftest")
    d.save(plain)
    out = with_copy_page(plain, plain, 1, tmp / "cp_out.pdf", anchor="依頼書")
    o = fitz.open(out)
    expect("with_copy_page: 頁数は変わらない", o.page_count == 3)
    expect("with_copy_page: 差し替えた頁は字を持たない (raster)", o[1].get_text().strip() == "")
    expect("with_copy_page: 他の頁は元のまま", "旅費請求書" in "".join(o[0].get_text().split()))
    expect("with_copy_page: 刷る頁の宣言を引き継ぐ", PP.read_record(o) is not None)
    try:
        with_copy_page(plain, plain, 0, tmp / "cp_bad.pdf", anchor="依頼書")
        stopped = False
    except BuildError:
        stopped = True
    expect("with_copy_page: 頁の目印が無ければ止める", stopped)


def _graft_tests(tmp, expect) -> None:
    """openpyxl の save で落ちる図形を _save が元 workbook から移し直す (drawings.graft_drawings)。"""
    import re
    import zipfile

    import openpyxl

    from . import drawings as DR
    from . import recipes as RC

    src = tmp / "graft-src.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb["S"]["A1"] = "x"
    wb.create_sheet("T")["A1"] = "y"
    wb.save(src)
    ms = '<a:rPr sz="1600"><a:latin typeface="MS Gothic"/><a:ea typeface="MS Gothic"/></a:rPr>'
    # form control の DrawingML 側 = cNvPr の拡張に a14:compatExt (VML 側の shape id) を持つ (Excel が書く形)
    ctrl = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">'
            + _anchor_xml("Check Box 1", "").replace('name="Check Box 1"/>', 'name="Check Box 1"><a:extLst>'
                                                     '<a:ext uri="{63B3BB69-23CF-44E3-9099-C40C66FF867C}">'
                                                     '<a14:compatExt spid="_x0000_s1025"/></a:ext></a:extLst>'
                                                     '</xdr:cNvPr>')
            + '</mc:Choice><mc:Fallback/></mc:AlternateContent>')
    # 16pt × 4 字 = 64pt、 枠 78pt = 余白の既定 7.2pt × 2 を足すと入らない (実雛形の寸法)
    anchors = (_anchor_xml("label", "外部資金", cx=78 * 12700, run="<a:r>" + ms + "<a:t>{t}</a:t></a:r>")
               + _anchor_xml("dup", "二重") + ctrl)
    _inject_drawing(src, _sheet_part(src, "S"), "", anchors=anchors)

    def drawn(path):
        with zipfile.ZipFile(path) as z:
            return "".join(z.read(n).decode() for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n))

    w = RC._load(src)
    w.remove(w["T"])
    out = RC._save(w, tmp / "graft-out.xlsx")
    d = drawn(out)
    expect("図形の移植: openpyxl の save の後も図形の字が残る", "外部資金" in d and "二重" in d, d[:200])
    expect("図形の移植: form control (mc:AlternateContent) は移さない", "AlternateContent" not in d and "Check Box" not in d)
    lab = re.search(r'name="label"/>.*?</xdr:oneCellAnchor>', d, re.S).group(0)
    ins = re.search(r'lIns="(\d+)"', lab)
    expect("図形の移植: 枠に入り切らない 1 行の label は余白を縮めて折り返さない",
           'wrap="none"' in lab and ins is not None and int(ins.group(1)) < 91440, lab[:300])
    w = RC._load(src)
    w.__dict__["_formcase_drop_shapes"] = {"S": {"dup"}}
    d = drawn(RC._save(w, tmp / "graft-drop.xlsx"))
    expect("図形の移植: spec の drop_shape の図形は移さない", "外部資金" in d and "二重" not in d)
    w = RC._load(src)
    w.__dict__["_formcase_drop_shapes"] = {"S": {"無い図形"}}
    try:
        RC._save(w, tmp / "graft-stale.xlsx")
        expect("図形の移植: 雛形に無い drop_shape は止める", False)
    except RC.BuildError as e:
        expect("図形の移植: 雛形に無い drop_shape は止める", "無い図形" in str(e), e)
    expect("図形の移植: data_only の読み込み (値の照合用) は移植しない",
           getattr(RC._load(src, data_only=True), "_formcase_source", None) is None)
    try:
        DR.graft_drawings(src, out)
        expect("図形の移植: drawing を既に持つ sheet には重ねない", False)
    except DR.GraftError as e:
        expect("図形の移植: drawing を既に持つ sheet には重ねない", "既にある" in str(e), e)
    # 2026-09-24 レビュー: form control の印の無い AlternateContent (数式入りの textbox 等 = 紙に出る) を黙って落とさない
    src2 = tmp / "graft-src-ac.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.save(src2)
    eq = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
          '<mc:Choice Requires="a14">' + _anchor_xml("数式", "x=1") + '</mc:Choice><mc:Fallback>'
          + _anchor_xml("数式", "x=1") + '</mc:Fallback></mc:AlternateContent>')
    _inject_drawing(src2, _sheet_part(src2, "S"), "", anchors=_anchor_xml("label", "見出し") + eq)
    try:
        RC._save(RC._load(src2), tmp / "graft-ac.xlsx")
        expect("図形の移植: form control でない mc:AlternateContent は止める (黙って落とさない)", False)
    except RC.BuildError as e:
        expect("図形の移植: form control でない mc:AlternateContent は止める (黙って落とさない)", "form control でない" in str(e), e)
    # 縦書きの 1 行 label は枠の幅と比べない (字は高さ方向に並ぶ) = 余白・折り返しを触らない
    vert = _anchor_xml("縦", "業務内容", cx=20 * 12700, run="<a:r>" + ms + "<a:t>{t}</a:t></a:r>").replace(
        "<a:bodyPr/>", '<a:bodyPr vert="eaVert"/>')
    expect("図形の移植: 縦書きの label は余白・折り返しを変えない", DR._fit_single_line(vert) == vert)


def _value_from_tests(tmp, expect) -> None:
    """fill.value_from = 個人情報を stub に書かずに別の workbook から実行時に読む (#pii-runtime-source)。"""
    import openpyxl

    from .fill import value_from

    src = tmp / "vf_src.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb["S"]["B2"] = "合成住所 1-2-3"
    wb.save(src)
    expect("value_from: cell の値を読む", value_from(src, "S", "B2") == "合成住所 1-2-3")
    for label, args in (("空の cell", (src, "S", "C3")), ("無い sheet", (src, "X", "B2")),
                        ("無い file", (tmp / "nope.xlsx", "S", "B2"))):
        try:
            value_from(*args)
            stopped = False
        except SystemExit:
            stopped = True
        expect(f"value_from: {label}なら止める", stopped)


# ---------------------------------------------------------------------------
# freeze / 凍結の不変条件
# ---------------------------------------------------------------------------
def _freeze_tests(tmp, case, expect, fails) -> None:
    m = M.load(case)
    try:
        L.freeze(m, "d1", "g1", "submitted", "2026-07-31", evidence="fixture")
        expect("paper: を言わない submitted の freeze は拒否", False)
    except M.ManifestError as e:
        expect("paper: を言わない submitted の freeze は拒否", "--paper" in str(e))
    try:
        L.freeze(m, "d1", "g1", "submitted", "2026-07-31", paper="differs")
        expect("paper differs に差の説明が無い freeze は拒否", False)
    except M.ManifestError as e:
        expect("paper differs に差の説明が無い freeze は拒否", "--paper-diff" in str(e))
    expect("拒否された freeze は manifest を変えない",
           M.load(case).group("d1", "g1")["current"]["state"] == "draft"
           and m.group("d1", "g1")["current"]["state"] == "draft")
    try:
        L.freeze(m, "d1", "g1", "submitted", "2026-07-31", paper="same", paper_basis="fixture で同じ run")
        expect("最後の保存が Mac Excel でない workbook の freeze は拒否 (normalize を案内)", False)
    except M.ManifestError as e:
        expect("最後の保存が Mac Excel でない workbook の freeze は拒否 (normalize を案内)", "normalize" in str(e), e)
    _mark_mac(case / "book.xlsx")
    m = M.load(case)
    L.freeze(m, "d1", "g1", "submitted", "2026-07-31", evidence="fixture", paper="same",
             paper_basis="fixture で同じ run")
    m.save()
    cur = M.load(case).group("d1", "g1")["current"]
    expect("freeze が sha256 と group の sheet digest を記録", cur["frozen"]["sha256"].get("print")
           and set(cur["frozen"]["sheet_digest"]) == {MAIN, PLAN, REQ})
    expect("freeze 直後の check は FAIL なし", not fails(case), fails(case))

    _set(case, REPORT, "L25", "架空の内訳")
    expect("draft group の sheet 変更は FAIL にしない", not fails(case), fails(case))
    _set(case, PLAN, "U10", "出発地 → 目的地")
    fs = fails(case)
    expect("凍結 group の sheet 変更は FAIL", any("提出済み" in x[2] for x in fs), fs)
    _set(case, PLAN, "U10", "出発地（架空市）→ 目的地")
    (case / "g1.pdf").write_bytes(b"%PDF-g1-overwritten")
    fs = fails(case)
    expect("凍結出力の bytes 変更は FAIL", any("bytes" in x[2] for x in fs), fs)
    (case / "g1.pdf").write_bytes(b"%PDF-g1")


# ---------------------------------------------------------------------------
# 書式の fingerprint (v3 / v2 / v1 の版ごとの射程)
# ---------------------------------------------------------------------------
def _format_tests(tmp, case, expect, fails) -> None:
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Border, Font, PatternFill, Side

    fr0 = M.load(case).group("d1", "g1")["current"]["frozen"]
    expect("新しい freeze は v1 と v3 を記録 (v2 は書かない)",
           set(fr0.get("sheet_digest_v3") or {}) == {MAIN, PLAN, REQ} and "sheet_digest_v2" not in fr0)

    def _fmt(fn):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wbx = openpyxl.load_workbook(case / "book.xlsx")
            fn(wbx)
            wbx.save(case / "book.xlsx")

    def fmt_fails():
        return any("書式" in x[2] for x in fails(case))

    props = [
        ("行高 (0.5 刻みの丸めでは見えない 14.25 → 14.0)",
         lambda w: setattr(w[PLAN].row_dimensions[12], "height", 14.0),
         lambda w: setattr(w[PLAN].row_dimensions[12], "height", 14.25), True),
        ("行高 30 (大きな変更)", lambda w: setattr(w[PLAN].row_dimensions[10], "height", 30),
         lambda w: setattr(w[PLAN].row_dimensions[10], "height", None), True),
        ("列幅 1 pixel (3.0 = 18px → 2.83 = 17px、 0.5 刻みの丸めでは同じ)",
         lambda w: setattr(w[PLAN].column_dimensions["C"], "width", 2.83203125),
         lambda w: setattr(w[PLAN].column_dimensions["C"], "width", 3.0), True),
        ("列幅を Excel が格子に丸め直しただけ (1.25 → 1.1640625 = 同じ 7px) は FAIL にしない",
         lambda w: setattr(w[PLAN].column_dimensions["D"], "width", 1.1640625),
         lambda w: setattr(w[PLAN].column_dimensions["D"], "width", 1.25), False),
        ("幅の記録が無かった列に既定幅の entry が足されただけは FAIL にしない",
         lambda w: setattr(w[PLAN].column_dimensions["F"], "width", 53 / 6),
         lambda w: w[PLAN].column_dimensions.pop("F", None), False),
        ("値のある cell の罫線", lambda w: setattr(w[PLAN]["U10"], "border", Border(left=Side(style="thin"))),
         lambda w: setattr(w[PLAN]["U10"], "border", Border()), True),
        ("値のある cell の字の色", lambda w: setattr(w[PLAN]["U10"], "font", Font(color="FFFF0000")),
         lambda w: setattr(w[PLAN]["U10"], "font", Font()), True),
        ("値のある cell の塗り", lambda w: setattr(w[PLAN]["U10"], "fill", PatternFill("solid", fgColor="FFFFFF00")),
         lambda w: setattr(w[PLAN]["U10"], "fill", PatternFill()), True),
        ("値の無い cell の罫線 (印刷範囲の中)", lambda w: setattr(w[PLAN]["V20"], "border", Border()),
         lambda w: setattr(w[PLAN]["V20"], "border", Border(bottom=Side(style="thin"))), True),
        ("値の無い cell の塗り", lambda w: setattr(w[PLAN]["V21"], "fill", PatternFill("solid", fgColor="FFCCCCCC")),
         lambda w: setattr(w[PLAN]["V21"], "fill", PatternFill()), True),
        ("値の無い cell の font だけ (紙に出ない = Excel の初回保存が揺らす) は FAIL にしない",
         lambda w: setattr(w[PLAN]["V22"], "font", Font(name="ＭＳ Ｐゴシック")),
         lambda w: setattr(w[PLAN]["V22"], "font", Font()), False),
        ("値のある cell の表示書式", lambda w: setattr(w[PLAN]["U10"], "number_format", "@"),
         lambda w: setattr(w[PLAN]["U10"], "number_format", "General"), True),
        ("結合範囲", lambda w: w[PLAN].merge_cells("U10:W10"), lambda w: w[PLAN].unmerge_cells("U10:W10"), True),
        ("条件付き書式", lambda w: w[PLAN].conditional_formatting.add(
            "U10", CellIsRule(operator="equal", formula=['"x"'], font=Font(color="FFFF0000"))),
         lambda w: setattr(w[PLAN], "conditional_formatting", type(w[PLAN].conditional_formatting)()), True),
        ("footer", lambda w: setattr(w[PLAN].oddFooter.center, "text", "p.2"),
         lambda w: setattr(w[PLAN].oddFooter.center, "text", None), True),
        ("fitToPage の時の scale (Excel が計算して書き直す) は FAIL にしない",
         lambda w: setattr(w[PLAN].page_setup, "scale", 77),
         lambda w: setattr(w[PLAN].page_setup, "scale", None), False),
        ("fitToPage を外す (印刷設定)",
         lambda w: setattr(w[PLAN].sheet_properties.pageSetUpPr, "fitToPage", False),
         lambda w: setattr(w[PLAN].sheet_properties.pageSetUpPr, "fitToPage", True), True),
    ]
    for label, change, undo, want in props:
        blob = (case / "book.xlsx").read_bytes()
        _fmt(change)
        got = fmt_fails()
        expect(f"v3: {label}{' → FAIL (書式)' if want else ''}", got == want, fails(case))
        _fmt(undo)
        if fails(case):                     # 戻し方が既定の style と違う (Font() 等) = bytes で戻す
            (case / "book.xlsx").write_bytes(blob)
        expect(f"   戻せば FAIL なし ({label[:12]}…)", not fails(case), fails(case))

    # v2 / v1 の issue は記録した版で照合する (v3 の射程を後から当てない)
    m1 = M.load(case)
    fr1 = m1.group("d1", "g1")["current"]["frozen"]
    v3 = fr1.pop("sheet_digest_v3")
    fr1["sheet_digest_v2"] = M.sheet_digests_v2(case / "book.xlsx", list(v3))
    m1.save()
    _fmt(lambda w: setattr(w[PLAN]["U10"], "border", Border(left=Side(style="thin"))))
    expect("v2 の issue は罫線の変更を見ない (v2 の射程のまま = 既存の digest は有効)", not fails(case), fails(case))
    _fmt(lambda w: setattr(w[PLAN].row_dimensions[10], "height", 30))
    expect("v2 の issue も行高 30 は FAIL", fmt_fails())
    _fmt(lambda w: setattr(w[PLAN].row_dimensions[10], "height", None))
    _fmt(lambda w: setattr(w[PLAN]["U10"], "border", Border()))
    m1 = M.load(case)
    m1.group("d1", "g1")["current"]["frozen"].pop("sheet_digest_v2")
    m1.save()
    _fmt(lambda w: setattr(w[PLAN].row_dimensions[10], "height", 30))
    expect("v1 だけの凍結 issue は書式の変更を見ない (= 既存の digest は有効なまま)", not fails(case), fails(case))
    _set(case, PLAN, "U10", "出発地 → 目的地")
    expect("v1 だけの凍結 issue も値の変更は FAIL", any("提出済み" in x[2] for x in fails(case)))
    _set(case, PLAN, "U10", "出発地（架空市）→ 目的地")
    _fmt(lambda w: setattr(w[PLAN].row_dimensions[10], "height", None))
    m1 = M.load(case)
    m1.group("d1", "g1")["current"]["frozen"]["sheet_digest_v3"] = v3
    m1.save()
    expect("v3 を戻して FAIL なし", not fails(case), fails(case))

    # 図形・form control (openpyxl の保存は落とす = 紙から checkbox・標題が消える)
    from . import fingerprint as FPX

    dwb = tmp / "draw.xlsx"
    wd = openpyxl.Workbook()
    wd.active.title = "S"
    wd["S"]["A1"] = "x"
    wd.save(dwb)
    _inject_drawing(dwb, _sheet_part(dwb, "S"), "標題")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        d_before = FPX.sheet_digests_v3(dwb, ["S"], M.sheet_digests(dwb, ["S"]))
        wdl = openpyxl.load_workbook(dwb)
        expect("v3 は図形の文字を読む", any("標題" in x for x in FPX.format_items(wdl, wdl["S"], dwb)))
        openpyxl.load_workbook(dwb).save(dwb)
    d_after = FPX.sheet_digests_v3(dwb, ["S"], M.sheet_digests(dwb, ["S"]))
    expect("図形が落ちた (openpyxl の保存) は v3 の書式が変わる",
           d_before["S"]["format"] != d_after["S"]["format"] and d_before["S"]["values"] == d_after["S"]["values"])


# ---------------------------------------------------------------------------
# legacy driver の guard
# ---------------------------------------------------------------------------
def _guard_tests(tmp, case, expect, fails) -> None:
    try:
        G.legacy_guard(case / "gen.py")
        expect("凍結 group を書く legacy driver は止まる", False)
    except SystemExit as e:
        expect("凍結 group を書く legacy driver は止まる", "g1" in str(e.code))
    try:
        G.legacy_guard(case / "gen.py", groups=["g2"])
        expect("draft group だけに絞った legacy driver は通る", True)
    except SystemExit as e:
        expect("draft group だけに絞った legacy driver は通る", False, e.code)
    try:
        G.legacy_guard(case / "other.py")
        expect("manifest に宣言の無い driver は止まる", False)
    except SystemExit as e:
        expect("manifest に宣言の無い driver は止まる", "宣言されていない" in str(e.code))
    nomani = tmp / "repo-a" / "docs" / "legacy-case"
    nomani.mkdir()
    (nomani / "d.py").write_text("x")
    try:
        G.legacy_guard(nomani / "d.py")
        expect("manifest の無い dir の driver は素通り (旧案件互換)", True)
    except SystemExit:
        expect("manifest の無い dir の driver は素通り (旧案件互換)", False)


# ---------------------------------------------------------------------------
# 記録と紙の一致 (paper) / 日付不明 (date_ack) / 出力なし (outputs_ack)
# ---------------------------------------------------------------------------
def _annotate_tests(tmp, case, expect, fails) -> None:
    from . import markers as MK

    m = M.load(case)
    g1 = m.group("d1", "g1")["current"]
    fr_before = repr(g1["frozen"])
    del g1["paper"]
    lv = [x for x in M.validate(m, S.get) if "paper" in x[2]]
    expect("paper: の無い submitted issue は 🟡", lv and lv[0][0] == M.WARN, lv)
    g1["paper"] = "unverified"
    lv = [x for x in M.validate(m, S.get) if x[0] == M.FAIL]
    expect("paper unverified に paper_diff が無いと 🔴", any("paper_diff" in x[2] for x in lv), lv)
    g1["paper"] = "maybe"
    expect("paper の値の誤りは 🔴", any("maybe" in x[2] for x in M.validate(m, S.get) if x[0] == M.FAIL))
    L.annotate(m, "d1", "g1", paper="differs", paper_diff="窓口で手書き訂正 2 点", paper_commit="abc1234")
    m.save()
    m = M.load(case)
    g1 = m.group("d1", "g1")["current"]
    expect("annotate は frozen (sha256 / sheet_digest) を変えない", repr(g1["frozen"]) == fr_before)
    expect("annotate は paper を note / log / frozen の前に置く",
           list(g1).index("paper") < list(g1).index("frozen"))
    fs = C.check_case(m)
    pl = [x for x in fs if x[0] == M.PAPER]
    expect("differs は 📄 の行になり 🔴/🟡 にはならない", pl and "記録は提出した紙と違う" in pl[0][2]
           and "abc1234" in pl[0][2] and not [x for x in fs if x[0] in (M.FAIL, M.WARN) and "paper" in x[2]], fs)
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        bad = C.render("case-a", fs, quiet_ok=True)
    expect("check --quiet でも 📄 は出て exit には効かない", "📄≠" in buf.getvalue() and not bad, buf.getvalue())
    mk = MK.render(m)
    expect("隔離 marker に 📄≠ の行と 紙 列が出る", "📄≠ 記録は提出した紙と違う: 窓口で手書き訂正 2 点" in mk
           and "≠ **違う**" in mk, mk)
    rows = C.status_rows(m)
    expect("status に 📄≠ の行", any("📄≠" in e for r in rows for e in r[7]), rows)
    try:
        L.annotate(m, "d1", "g2", paper="same")
        expect("draft の issue に paper は書けない", False)
    except M.ManifestError:
        expect("draft の issue に paper は書けない", True)
    g1["date"] = "unknown"
    w = [x for x in M.validate(m, S.get) if "日付" in x[2]]
    expect("date unknown は 🟡 (owner に確認)", w and w[0][0] == M.WARN, w)
    L.annotate(m, "d1", "g1", date_ack="owner: 分からないと確認済")
    w = [x for x in M.validate(m, S.get) if "日付" in x[2]]
    expect("date_ack があれば 🔵 (聞き直さない)", w and w[0][0] == M.INFO and "分からない" in w[0][2], w)
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        C.render("case-a", C.check_case(m), quiet_ok=True)
    expect("date_ack の 🔵 は --quiet では出さない", "owner 確認済" not in buf.getvalue(), buf.getvalue())
    try:
        L.annotate(m, "d1", "g2", date_ack="x")
        expect("date が unknown でない issue に date_ack は書けない", False)
    except M.ManifestError:
        expect("date が unknown でない issue に date_ack は書けない", True)
    g1["date"] = "2026-07-31"
    del g1["date_ack"]
    outs_saved, sha_saved = g1.pop("outputs"), g1["frozen"].pop("sha256")
    w = [x for x in M.validate(m, S.get) if "outputs" in x[2] or "出力 file" in x[2]]
    expect("凍結 issue に outputs が無いと 🟡", w and w[0][0] == M.WARN, w)
    L.annotate(m, "d1", "g1", outputs_ack="fixture: folder と git history に PDF なし")
    w = [x for x in M.validate(m, S.get) if "outputs" in x[2] or "出力 file" in x[2]]
    expect("outputs_ack があれば 🔵 (理由つきの事実)", w and w[0][0] == M.INFO and "PDF なし" in w[0][2], w)
    expect("隔離 marker に outputs_ack の理由", "PDF なし" in MK.render(m))
    del g1["outputs_ack"]
    g1["outputs"], g1["frozen"]["sha256"] = outs_saved, sha_saved
    try:
        L.annotate(m, "d1", "g1", outputs_ack="x")
        expect("outputs のある issue に outputs_ack は書けない", False)
    except M.ManifestError:
        expect("outputs のある issue に outputs_ack は書けない", True)
    m.save()


# ---------------------------------------------------------------------------
# printed → STALE / reopen / 出力 path の共有
# ---------------------------------------------------------------------------
def _lifecycle_tests(tmp, case, expect, fails) -> None:
    _mark_mac(case / "book.xlsx")
    m = M.load(case)
    L.freeze(m, "d1", "g2", "printed", "2026-09-18", paper="same", paper_basis="fixture")
    m.save()
    _set(case, REPORT, "L25", "別の内訳")
    fs = fails(case)
    expect("printed の後の sheet 変更は STALE", any("STALE" in x[2] for x in fs), fs)
    m = M.load(case)
    try:
        L.freeze(m, "d1", "g2", "submitted", "2026-09-20", paper="same")
        expect("旧紙の submitted 化は拒否", False)
    except M.ManifestError:
        expect("旧紙の submitted 化は拒否", True)
    new = L.reopen(m, "d1", "g2", "記入誤りの訂正", "2026-09-19")
    m.save()
    expect("reopen は _r2 の新しい出力名で draft を作る",
           new["outputs"]["print"] == "g2_r2.pdf" and new["state"] == "draft")
    fz = M.load(case).frozen_outputs()
    expect("reopen 後も前の printed 出力は凍結のまま", str((case / "g2.pdf").resolve()) in fz)
    expect("reopen 後の check は FAIL なし", not fails(case), fails(case))
    m = M.load(case)
    g = m.group("d1", "g2")
    g["previous"][0]["outputs"] = {"print": "g2_r2.pdf"}
    m.save()
    fs = fails(case)
    expect("前の凍結 issue と current が path を共有 (commit なし) は FAIL", any("commit" in x[2] for x in fs), fs)
    g["previous"][0]["commit"] = "abc1234"
    m.save()
    expect("commit を書けば共有 path を許す", not fails(case), fails(case))
    g["previous"][0]["outputs"] = {"print": "g2.pdf"}
    del g["previous"][0]["commit"]
    m.save()


# ---------------------------------------------------------------------------
# pre-commit guard (一時 git repo)
# ---------------------------------------------------------------------------
def _precommit_tests(tmp, case, expect) -> None:
    if not shutil.which("git"):
        print("  SKIP pre-commit guard test (git 不在)")
        return
    import openpyxl

    repo = tmp / "repo-a"
    env = dict(os.environ, GIT_CONFIG_GLOBAL="/dev/null", GIT_CONFIG_SYSTEM="/dev/null")

    def git(*a):
        return subprocess.run(["git", "-C", str(repo), "-c", "user.email=t@t", "-c", "user.name=t", *a],
                              env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)
    rel = str(case.relative_to(repo))
    git("init", "-q")
    git("add", "-A")
    git("commit", "-q", "--no-verify", "-m", "init")
    (case / "g1.pdf").write_bytes(b"%PDF-changed")
    git("add", f"{rel}/g1.pdf")
    bl = G.staged_findings(repo)
    expect("pre-commit: 凍結出力の staged 変更を BLOCK", any("g1.pdf" in b for b in bl), bl)
    git("reset", "-q")
    (case / "g1.pdf").write_bytes(b"%PDF-g1")
    (case / "g2_r2.pdf").write_bytes(b"%PDF-r2")
    git("add", f"{rel}/g2_r2.pdf")
    expect("pre-commit: draft issue の新しい出力は通る", not G.staged_findings(repo), G.staged_findings(repo))
    git("reset", "-q")
    _set(case, PLAN, "U46", "目的地 → 出発地")
    git("add", f"{rel}/book.xlsx")
    bl = G.staged_findings(repo)
    expect("pre-commit: 凍結 sheet の値の変更を BLOCK", any("sheet" in b for b in bl), bl)
    git("reset", "-q")
    _set(case, PLAN, "U46", None)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wbg = openpyxl.load_workbook(case / "book.xlsx")
        wbg[PLAN].row_dimensions[12].height = 40
        wbg.save(case / "book.xlsx")
    git("add", f"{rel}/book.xlsx")
    bl = G.staged_findings(repo)
    expect("pre-commit: 凍結 sheet の書式だけの変更も BLOCK", any("(書式)" in b for b in bl), bl)
    m = M.load(case)
    L.reopen(m, "d1", "g1", "訂正の指示があった", "2026-09-19")
    m.save()
    git("add", f"{rel}/{M.MANIFEST_NAME}")
    bl = G.staged_findings(repo)
    expect("pre-commit: reopen を同じ commit に入れれば sheet 変更は通る", not bl, bl)
    git("reset", "-q")
    (case / M.MANIFEST_NAME).write_text("schema: formcase/1\ndocuments: {x: {form: nope}}\n")
    git("add", f"{rel}/{M.MANIFEST_NAME}")
    bl = G.staged_findings(repo)
    expect("pre-commit: 壊れた manifest を BLOCK", any("form" in b for b in bl), bl)
    git("checkout", "--", f"{rel}/{M.MANIFEST_NAME}")
    git("reset", "-q")


def _exit_code_tests(expect) -> None:
    """pre-commit の入口 (guard / lint --staged) は、 内部の故障で 1 を返してはいけない。

    1 = 「明示の BLOCK」 は呼び元 (pre-commit chain) との約束なので、 engine の故障が 1 になると
    **その repo のどの commit も止まる** — しかも「凍結を守った」 のと区別がつかない。 2026-09-20 に
    実測 (層1 CLI の読み込みが失敗した状態で、 無関係な新規 file の commit まで止まった)。
    対話の入口はこの受けに入れない = 本当の bug を traceback のまま見せる。"""
    import importlib.util

    cli_path = Path(__file__).resolve().parent.parent / "formcase.py"
    sp = importlib.util.spec_from_file_location("formcase_cli_selftest", cli_path)
    cli = importlib.util.module_from_spec(sp)
    sys.modules[sp.name] = cli
    sp.loader.exec_module(cli)

    def boom(args):
        raise RuntimeError("simulated internal failure")

    saved = {k: getattr(cli, k) for k in ("cmd_guard", "cmd_lint", "cmd_status")}
    try:
        cli.cmd_guard = cli.cmd_lint = cli.cmd_status = boom
        for argv, name in ((["guard", "--staged"], "guard --staged"), (["lint", "--staged"], "lint --staged")):
            with contextlib.redirect_stderr(io.StringIO()):
                rc = cli.main(argv)
            expect(f"pre-commit: {name} の内部エラーは 3 (= commit を止めない)", rc == 3, rc)
        expect("pre-commit でない入口は例外のまま (故障を握りつぶさない)",
               _raises(RuntimeError, cli.main, ["status"]))
    finally:
        for k, v in saved.items():
            setattr(cli, k, v)


# ---------------------------------------------------------------------------
# recipe の登録 / 出力名 / scaffold
# ---------------------------------------------------------------------------
def _recipe_tests(tmp, inst, expect) -> None:
    from . import fill as FI
    from . import recipes as RC
    from . import scaffold as SC

    class R1(RC.Recipe):
        form = "fx"
        outputs = {"g1": {"print": "{stem}_a.pdf", "confirm": "{stem}_a_c.pdf"},
                   "g2": {"print": "{base}_b.pdf"}}
        seals = {1: "anchor=印,occurrence=1,size=27"}

        def name_vars(self, stem):
            return {"stem": stem, "base": stem[:-len("_second")] if stem.endswith("_second") else stem}

    RC.register("fx", R1)
    r = RC.recipe_for("fx")
    o = r.default_outputs("d1_second")
    expect("recipe: name_vars で出力名の変数を足せる ({base} は派生の接尾を除く)",
           o["g1"]["print"] == "d1_second_a.pdf" and o["g2"]["print"] == "d1_b.pdf", o)
    expect("recipe: 既定の name_vars は {stem} だけ", RC.Recipe().name_vars("x") == {"stem": "x"})
    expect("recipe: 登録の無い様式は BuildError", _raises(RC.BuildError, RC.recipe_for, "nope"))
    expect("recipe: seals_for は既定で seals", r.seals_for("g1") == R1.seals)
    expect("押印画像の command が設定に無ければ BuildError",
           _raises(RC.BuildError, RC._seal_image))
    expect("seal_mode: 既定は image、 physical を受け付け、 他の値は ConfigError",
           CF.seal_mode() == "image" and CF.check_seal_mode("physical") == "physical"
           and _raises(CF.ConfigError, CF.check_seal_mode, "stamp"))
    expect("stamp_hint: anchor と occurrence を紙の上で探せる言い方に",
           RC.stamp_hint("anchor=印,occurrence=2,size=27") == "「印」 の 2 個目 の欄"
           and RC.stamp_hint("anchor=㊞,occurrence=1") == "「㊞」 の欄"
           and RC.stamp_hint("x=1").startswith("押印の位置"))

    expect("fill: clear / general は値なしでも書く行、 value=None の text は未記入",
           FI._writes(("S", "A1", "clear", None)) and FI._writes(("S", "A1", "general", None))
           and not FI._writes(("S", "A1", "text", None)))

    c2 = tmp / "repo-a" / "docs" / "2026-02-02-case-b"
    SC.new_case("fx", c2, "d1")
    res = SC.new_case("fx", c2, "d1-second", group="g2", from_doc="d1")
    m2 = M.load(c2)
    d2 = m2.doc("d1-second")
    stub = res["stub"].read_text(encoding="utf-8")
    expect("new --group --from-doc: base の workbook の copy・その group だけの document",
           d2["workbook"] == "d1_second.xlsx" and list(d2["groups"]) == ["g2"]
           and (c2 / "d1_second.xlsx").read_bytes() == (c2 / "d1.xlsx").read_bytes()
           and d2["groups"]["g2"]["current"]["outputs"]["print"] == "d1_b.pdf", d2)
    expect("stub はその group の欄だけ (他 group の欄・日程表の案内を出さない) + 数式の欄は general → formula",
           f"'{REPORT}', 'X10'" in stub and f"'{REQ}'" not in stub and "1 日 1 block の表" not in stub
           and f"('{REPORT}', 'AC37', 'general', None, 'g2')" in stub
           and f"('{REPORT}', 'AC37', 'formula'" in stub, stub[-900:])
    expect("stub は設定の cli / sys.path を使う (engine に私的な絶対 path を焼かない)",
           CF.stub_sys_path() in stub and f"python3 {CF.cli()} build" in stub, stub[:1200])
    expect("base がまだ未提出なら notice を返す", res["notice"] and "提出の記録が無い" in res["notice"])
    expect("--from-doc が manifest に無ければ拒否",
           _raises(M.ManifestError, SC.new_case, "fx", c2, "x", group="g2", from_doc="nope"))

    # add-group: 先の group だけで作った案件に、 後の group を足す (workbook は触らない)
    c3 = tmp / "repo-a" / "docs" / "2026-02-03-case-c"
    SC.new_case("fx", c3, "d3", group="g1")
    wb3 = (c3 / "d3.xlsx").read_bytes()
    r3 = SC.add_group(c3, "d3", "g2")
    d3 = M.load(c3).doc("d3")
    stub3 = r3["stub"].read_text(encoding="utf-8")
    expect("add-group: group が draft + recipe の既定の出力名で足され、 先の group は残り、 workbook は同じ bytes",
           list(d3["groups"]) == ["g1", "g2"] and d3["groups"]["g2"]["current"]["state"] == "draft"
           and d3["groups"]["g2"]["current"]["outputs"] == RC.recipe_for("fx").default_outputs("d3")["g2"]
           and (c3 / "d3.xlsx").read_bytes() == wb3, d3)
    expect("add-group: stub = fill_<doc>_<group>.py、 その group の欄だけ",
           r3["stub"].name == "fill_d3_g2.py" and f"'{REPORT}', 'X10'" in stub3 and f"'{REQ}'" not in stub3,
           stub3[-600:])
    expect("add-group: 既にある group / spec に無い group / 無い document は拒否 (stub も manifest も作らない)",
           _raises(M.ManifestError, SC.add_group, c3, "d3", "g2")
           and _raises(M.ManifestError, SC.add_group, c3, "d3", "nope")
           and _raises(M.ManifestError, SC.add_group, c3, "zz", "g2"))

    # value_from_text: 台帳の text から 1 つ取る。 当たり 0 / 2 件以上 / 空は止める
    led = tmp / "ledger.md"
    led.write_text("# 記録\n## 口座\n- 番号 1234567 / 種別 普通\n- 旧番号 7654321\n", encoding="utf-8")
    expect("value_from_text: 1 件だけ当たる pattern の group を返す",
           FI.value_from_text(led, r"^- 番号 (\d+)") == "1234567"
           and FI.value_from_text(led, r"種別 (?P<k>\S+)", "k") == "普通")
    expect("value_from_text: 当たり 0 件・2 件以上・file 無しは止める",
           _raises(SystemExit, FI.value_from_text, led, r"^- 口座名 (\S+)")
           and _raises(SystemExit, FI.value_from_text, led, r"番号 (\d+)")
           and _raises(SystemExit, FI.value_from_text, tmp / "none.md", r"x"))
    stub1 = (c2 / "fill_d1.py").read_text(encoding="utf-8")
    expect("stub: fixed の値は spec から入り、 数字だけの文字列は textfmt、 font_size の欄は fontsize 行",
           "'○甲案立替'" in stub1 and "'textfmt', '12345678'" in stub1 and "'fontsize', 7" in stub1,
           [x for x in stub1.split("\n") if "12345678" in x or "fontsize" in x])
    readme = (c2 / "README.md").read_text(encoding="utf-8")
    expect("README 枠は kind=status の生成表 + 設定の spec_hint / process_hint",
           "formcase:view kind=status" in readme and "reference/fx.yaml" in readme and "<手順 doc>" in readme, readme)
    from . import lint as LI
    hits = [(ln, LI.state_hits(ln)) for ln in readme.splitlines() if LI.state_hits(ln)]
    expect("new が書いた README の定型文は状態の lint に掛からない (掛かると new した案件の最初の commit が止まる)",
           not hits, hits)


def _raises(exc, fn, *a, **k) -> bool:
    try:
        fn(*a, **k)
    except exc:
        return True
    except Exception:  # noqa: BLE001
        return False
    return False


# ---------------------------------------------------------------------------
# Excel の script 生成と osascript の失敗の扱い (Excel は起こさない)
# ---------------------------------------------------------------------------
def _excel_tests(tmp, expect) -> None:
    from . import excel as XL

    s = XL.script_for([("S", "AC37", "general", None), ("S", "AC37", "formula", '=IF(A1="","",A1)'),
                       ("S", "J59:V59", "clear", None)])
    expect("Excel script: general = 書式 / formula = 数式 / clear = merge 全域の clear contents",
           'number format of range "AC37"' in s and 'set formula of range "AC37"' in s
           and 'clear contents of range "J59:V59"' in s and "calculate" in s, s)
    expect("Excel script: fontsize", "set font size of font object" in XL.script_for([("S", "D11", "fontsize", 7)]))

    class _R:
        def __init__(self, rc, err=""):
            self.returncode, self.stderr, self.stdout = rc, err, ""

    def _runner(seq):
        it = iter(seq)
        return lambda argv: next(it)

    old_log = XL.ERROR_LOG
    XL.ERROR_LOG = tmp / "excel-errors.log"
    try:
        t609 = "1:2: execution error: Microsoft Excel でエラーが起きました: 接続が無効です。 (-609)"
        XL.run_osascript("x", [], "probe", runner=_runner([_R(1, t609), _R(0)]), sleep=lambda s: None,
                         responds=lambda: True)
        expect("osascript: -609 は再試行して通る + log に逐語で残る",
               "(-609)" in XL.ERROR_LOG.read_text(encoding="utf-8"))
        try:
            XL.run_osascript("x", [], "probe", runner=_runner([_R(1, "1:2: execution error: 架空の構文エラー (-2741)")]),
                             sleep=lambda s: None, responds=lambda: True)
            expect("osascript: 一時的でないエラーは再試行せず、 標準エラーを逐語で ExcelError に", False)
        except XL.ExcelError as e:
            expect("osascript: 一時的でないエラーは再試行せず、 標準エラーを逐語で ExcelError に",
                   "架空の構文エラー (-2741)" in str(e) and "試行 1 回" in str(e), str(e))
        t1712 = "9:9: execution error: AppleEvent がタイムアウトしました。 (-1712)"
        try:
            XL.run_osascript("x", [], "probe", runner=_runner([_R(1, t1712)] * (XL.RETRIES + 1)),
                             sleep=lambda s: None, responds=lambda: True)
            expect("osascript: -1712 の再試行は RETRIES 回まで", False)
        except XL.ExcelError as e:
            expect("osascript: -1712 の再試行は RETRIES 回まで", f"試行 {XL.RETRIES + 1} 回" in str(e), str(e))
        try:
            XL.run_osascript("x", [], "probe", runner=_runner([_R(1, t1712)] * 3), sleep=lambda s: None,
                             responds=lambda: False)
            expect("osascript: -1712 の後に Excel が答えないなら再試行せずに止める", False)
        except XL.ExcelError as e:
            expect("osascript: -1712 の後に Excel が答えないなら再試行せずに止める", "応答しない" in str(e), str(e))
    finally:
        XL.ERROR_LOG = old_log


# ---------------------------------------------------------------------------
# 体裁 (layout) — 合成の雛形に対して
# ---------------------------------------------------------------------------
def _layout_tests(tmp, expect) -> None:
    import openpyxl

    from . import excel as XL
    from . import layout as LY

    spec = S.get("fx")
    tpl = S.template_path(spec)
    expect("雛形が合成されている", tpl.exists(), tpl)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(tpl)
        wbv = openpyxl.load_workbook(tpl, data_only=True)

    def shape(w):
        return {n: ({r: d.height for r, d in w[n].row_dimensions.items() if d.height is not None},
                    {str(x) for x in w[n].merged_cells.ranges},
                    {c.coordinate: (c.alignment.wrap_text, c.number_format,
                                    c.border.bottom.style if c.border.bottom else None)
                     for row in w[n].iter_rows() for c in row if c.has_style})
                for n in w.sheetnames}

    first = LY.prepare(wb, wbv, spec)
    s1 = shape(wb)
    LY.prepare(wb, wbv, spec)
    expect("体裁 (雛形の行高・render:・日程表の正規化・値の行高) は冪等", bool(first) and shape(wb) == s1)
    expect("render: の表示書式・行高の下限・揃え・白黒が当たる",
           wb[MAIN]["I16"].number_format == "h:mm" and LY._row_h(wb[MAIN], 37) >= 15
           and wb[REPORT]["M29"].alignment.vertical == "center" and wb[REPORT].page_setup.blackAndWhite, first[:6])
    expect("render: の merge が当たる", "L35:X35" in {str(x) for x in wb[REPORT].merged_cells.ranges})
    expect("page = 印刷範囲 (render: の print_area) × 雛形の手動改ページ",
           LY.pages_for(spec, MAIN) == ["A1:AH30", "A31:AH60"], LY.pages_for(spec, MAIN))
    expect("page = 雛形の印刷範囲 (render: に print_area が無い sheet)",
           LY.pages_for(spec, REPORT) == ["A1:AH60"], LY.pages_for(spec, REPORT))
    # when: 下段が空なら 1 つの欄 / 値があれば 2 段
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        w2 = openpyxl.load_workbook(tpl)
    v2 = openpyxl.Workbook()
    v2.active.title = REPORT
    v2[REPORT]["J54"] = "上段"
    LY.apply_render_fixes(w2, spec, wbv=v2, sheets=[REPORT])
    mg = {str(x) for x in w2[REPORT].merged_cells.ranges}
    expect("render: の when (empty) = 下段が空なら 1 つの欄", "J54:V60" in mg and "J54:V56" not in mg, sorted(mg)[:6])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        w3 = openpyxl.load_workbook(tpl)
    v2[REPORT]["J57"] = "下段"
    LY.apply_render_fixes(w3, spec, wbv=v2, sheets=[REPORT])
    mg = {str(x) for x in w3[REPORT].merged_cells.ranges}
    expect("render: の when (filled) = 下段に値があれば 2 段", "J54:V56" in mg and "J54:V60" not in mg, sorted(mg)[:6])
    # 値の長さ → 折り返し + 行を伸ばす / 短い値は触らない / 潰した行は伸ばさない
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        w4 = openpyxl.load_workbook(tpl)
        w4v = openpyxl.load_workbook(tpl, data_only=True)
    h17 = LY._row_h(w4[MAIN], 17) + LY._row_h(w4[MAIN], 18)
    h20 = LY._row_h(w4[MAIN], 20) + LY._row_h(w4[MAIN], 21)
    w4v[MAIN]["G17"] = "架空の用務の説明を長く書いた欄。" * 6
    w4v[MAIN]["G20"] = "短い値"
    ch = LY.prepare(w4, w4v, spec)
    expect("長い値の欄 = 折り返し + 行を伸ばす", w4[MAIN]["G17"].alignment.wrap_text
           and LY._row_h(w4[MAIN], 17) + LY._row_h(w4[MAIN], 18) > h17, [c for c in ch if MAIN in c])
    expect("短い値の欄は触らない (行を伸ばさない)",
           LY._row_h(w4[MAIN], 20) + LY._row_h(w4[MAIN], 21) == h20 and not any("G20" in c for c in ch), ch)
    collapsed = [r for r in range(9, 25) if LY._row_h(w4[PLAN], r) == LY.COLLAPSED_PT]
    expect("日程表の使わない下段の行は潰れている (normalize_nittei)", len(collapsed) >= 3, collapsed)
    w4v[PLAN]["I10"] = "架空の用務" * 30
    ch2 = LY.fit_text(w4, w4v, [PLAN], frozen_max=5.0)
    expect("日程表の長い値は block の上段の行だけ伸ばし、 潰した行 (3pt) は伸ばさない",
           any("I10" in c for c in ch2)
           and all(LY._row_h(w4[PLAN], r) == LY.COLLAPSED_PT for r in collapsed), (ch2, collapsed))
    before = LY._row_h(w4[MAIN], 17)
    LY.fit_text(w4, w4v, [MAIN], extra={(MAIN, "G17"): {"pt": 10}})
    expect("PDF で足りなかった欄 (extra) は今の高さに足す", LY._row_h(w4[MAIN], 17) > before)
    # 行の送りは font の hhea (ascent − descent + lineGap) から。 Excel 同梱の font は macOS にしか無いので、 合成した
    # 2 本の TTF (行の高さ 1.602 em = 游ゴシック相当 / 1.0 em = MS 系相当) を一時 dir に置いて読ませ、 実機の font が
    # 在ればそれも見る (Linux の CI は合成だけ。 2026-09-22: 実機 font 前提の 1 項が CI で常に赤だった)。
    _synth_font_checks(expect)
    # snapshot → 体裁 → excel_ops (openpyxl で save すると図形が消える雛形の経路)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        w5 = openpyxl.load_workbook(tpl)
        w5v = openpyxl.load_workbook(tpl, data_only=True)
    ws5 = w5[MAIN]
    snap = LY.snapshot(ws5, "A1:AH60")
    w5v[MAIN]["G17"] = "架空の用務の説明を長く書いた欄。" * 6      # 狭い結合 = 折り返し + 行を伸ばす
    LY.prepare(w5, w5v, spec, sheets=[MAIN])
    ops = LY.excel_ops(snap, ws5, "A1:AH60")
    expect("excel_ops: 折り返しと行高の差分を Excel の操作として出す",
           ("wrap", "G17", True) in ops and any(o[0] == "row_height" and o[1] == 17 for o in ops), ops[:4])
    lines = XL.ops_lines(MAIN, ops + [("print_area", "A1:AH60"), ("one_page",), ("black_and_white", True)])
    expect("ops_lines: 行高・折り返し・印刷範囲・1 枚に収める・白黒を AppleScript に",
           any('row height of range "17:17"' in x for x in lines) and any("wrap text" in x for x in lines)
           and any('"$A$1:$AH$60"' in x for x in lines) and any("fit to pages tall" in x for x in lines)
           and any("black and white" in x for x in lines), lines[:3])
    expect("ops_lines: 未対応の体裁の変更は止める (黙って落とさない)",
           _raises(XL.ExcelError, XL.ops_lines, MAIN, [("border_bottom", "A1", "thin")]))


# ---------------------------------------------------------------------------
# generated view / lint
# ---------------------------------------------------------------------------
def _view_lint_tests(tmp, expect) -> None:
    from . import lint as LI
    from . import rules as RU
    from . import views as VI

    allr = RU.all_rules()
    expect("spec の規則が読める", "fx/days" in allr and "0.5" in allr["fx/days"]["summary"])
    expect("same_as が本文を借りる", allr.get("fx/days-sum", {}).get("summary") == allr["fx/days"]["summary"])
    expect("claims の網羅: 合成 spec は全部の規則が claims か claims_na を持つ",
           not LI.claims_coverage(), LI.claims_coverage())
    md = tmp / "repo-a" / "docs" / "reference" / "doc.md"
    md.parent.mkdir(parents=True, exist_ok=True)
    md.write_text("# t\n<!-- formcase:view kind=table rules=fx/days -->\n<!-- /formcase:view -->\n"
                  "説明に `<!-- formcase:view kind=table rules=nope/x -->` と書いても marker ではない\n",
                  encoding="utf-8")
    expect("空の view は stale", VI.check_files([md])[0][1] == "stale")
    expect("views --write で描ける (backtick 内の例は無視)", VI.check_files([md], write=True)[0][1] == "written")
    expect("描いた後は ok", VI.check_files([md])[0][1] == "ok")
    md.write_text(md.read_text(encoding="utf-8").replace("0.5", "0.25"), encoding="utf-8")
    expect("view を手で書き換えると stale", VI.check_files([md])[0][1] == "stale")
    md.write_text("<!-- formcase:view kind=table rules=fx/no-such-rule -->\n<!-- /formcase:view -->\n",
                  encoding="utf-8")
    expect("存在しない規則 id の view は error", VI.check_files([md])[0][1] == "error")
    md.write_text("<!-- formcase:view kind=cells form=fx refs=AB5,U10 -->\n<!-- /formcase:view -->\n",
                  encoding="utf-8")
    VI.check_files([md], write=True)
    body = md.read_text(encoding="utf-8")
    expect("kind=cells は fixed の値と (state) を表にする", "**甲**" in body and "(empty)" in body, body)
    old = allr["fx/days"]["superseded"][0]
    md.write_text(f"手順: {old} で書く\n<!-- formcase:history -->経緯: {old}<!-- /formcase:history -->\n"
                  f"<!-- formcase:view kind=checklist rules=fx/tatekae -->\n- ○甲案立替\n<!-- /formcase:view -->\n",
                  encoding="utf-8")
    f = LI.scan(paths=[md], root=tmp)
    expect("lint: 廃止版の言い回しを検出 (history 区間と view の中は除外)",
           [(x[1], x[4]) for x in f] == [(1, "superseded")], f)
    md.write_text("立替は ○甲案立替 と書く\n", encoding="utf-8")
    expect("lint: 記号つき固定値の書き写しを検出",
           any(x[4] == "value" for x in LI.scan(paths=[md], root=tmp)), LI.scan(paths=[md], root=tmp))

    def lines(ls, where=md, statement=False):
        where.parent.mkdir(parents=True, exist_ok=True)
        where.write_text("\n".join(ls) + "\n", encoding="utf-8")
        return [x for x in LI.scan(paths=[where], root=tmp) if statement or x[4] != "statement"]

    cases = [
        ("廃止版そのまま", "滞在日数は 0.25 刻み", True),
        ("全角・空白・助詞の違いを吸う", "滞在日数は　０．２５　刻み", True),
        ("規則の定義文の書き写し (marker)", "同意の欄は ☑選ばない を消す", True),
        ("今の規則そのものは検出しない (= それが本文)", "滞在日数は 0.5 刻みで書く", True),
        ("規則に関係ない行", "この doc は手順の入口である", False),
    ]
    for label, line, want in cases:
        expect(f"lint (正規化): {label} → {'検出' if want else '検出しない'}", bool(lines([line])) == want, lines([line]))
    kinds = [
        ("claims value: 廃止した刻みの数", "滞在日数は 0.25 で計算する", "fx/days"),
        ("claims forbid", "例の様式の氏名は略称で出す", "fx/name"),
        ("claims forbid の直後が否定なら拾わない", "例の様式の氏名は略称にしない", None),
        ("claims forbid_re", "例の様式の区分は「協力者」", "fx/kubun"),
        ("claims forbid_re: 除外の語は拾わない", "例の様式の区分 = 研究協力者", None),
        ("claims affirm: 空が正の欄に書く", "例の様式の経路は本人が書いて出す", "fx/keiro"),
        ("claims affirm の直後が否定なら拾わない", "例の様式の経路は書かない", None),
        ("claims negation", "例の様式の請求の欄は不要", "fx/seikyuu"),
        ("claims other_numbers + unit", "単価は 10,000 円", "fx/tanka"),
        ("claims other_numbers: 日付の数は拾わない", "単価は 5,000 円 (2026-09-02 確定)", None),
    ]
    for label, line, rid in kinds:
        f = lines([line])
        got = {x[3] for x in f}
        expect(f"lint (claims): {label} → {rid or '検出しない'}", (rid in got) if rid else not f, f)
    f = lines(["- AG38 は空欄で出す", "- M38 は「日付選択」 のまま印刷"])
    expect("lint (cell の値の主張): spec と違う値の主張を検出", {x[3] for x in f} >= {"fx/AG38", "fx/M38"}, f)
    expect("lint (cell の値の主張): spec と同じ値は検出しない", not lines(["- AG38 = ☑"]), lines(["- AG38 = ☑"]))
    stmt = [
        ("cell 番地 × 値 (矛盾していなくても書き写し)", "- 例の様式の I37 は空欄で出す", True),
        ("規則の語 × 引用の値 (様式の話の印つき)", "- 例の様式の職名は「乙種」", True),
        ("規則 id の参照だけは書いてよい", "- 例の様式の職名 = 規則 `fx/shokumei`", False),
        ("history 区間の中は経緯",
         "<!-- formcase:history -->職名は「甲種」 と昔は書いた<!-- /formcase:history -->", False),
        ("様式の話の印が無い一般語は拾わない", "- 別のシステムの経路は「申請」 から", False),
    ]
    for label, line, want in stmt:
        f = [x for x in lines([line], statement=True) if x[4] == "statement"]
        expect(f"lint (statement): {label} → {'検出' if want else '検出しない'}", bool(f) == want, f)
    cov = LI.claims_coverage([("t/a", {"rule": "a", "state": "fixed", "value": "x", "summary": "A = x"}),
                              ("t/b", {"rule": "b", "state": "fixed", "summary": "B = y", "claims_na": "理由"}),
                              ("t/c", {"id": "c", "summary": "手順の順番", "claims_na": "値でなく手順の順番"}),
                              ("t/d", {"id": "d", "summary": "手順の順番"}),
                              ("t/e", {"rule": "e", "state": "empty", "summary": "空",
                                       "claims": [{"about": ["E"], "forbid": ["x"], "says": "s"}]}),
                              ("t/f", {"rule": "f", "same_as": "t/e"})])
    expect("claims の網羅: 値の規則に claims 無し / 値の規則に claims_na / 理由なしは 🔴、"
           " claims あり・same_as 先・値でない規則の claims_na は通る",
           {r for r, _w in cov} == {"t/a", "t/b", "t/d"}, cov)
    expect("lint の承認一覧は設定が指す file", CF.lint_ack() == CF.config_dir() / "lint-ack.yaml")


# ---------------------------------------------------------------------------
# 案件 README に状態を書かない仕組み
# ---------------------------------------------------------------------------
def _case_readme_tests(tmp, expect) -> None:
    import importlib.util

    from . import lint as LI
    from . import views as VI

    case = tmp / "repo-a" / "docs" / "2026-03-03-readme-case"
    case.mkdir(parents=True)
    data = {"schema": M.SCHEMA, "case": case.name, "todo": ["repo-a:2026-03-03-x"],
            "documents": {"d1": {"form": "fx", "workbook": "book.xlsx",
                                 "groups": {"g1": {"current": {"state": "draft", "outputs": {"print": "j.pdf"}}}}}}}
    (case / M.MANIFEST_NAME).write_text(M.dump_text(data), encoding="utf-8")
    readme = case / "README.md"
    readme.write_text("# t\n\n<!-- formcase:view kind=status -->\n<!-- /formcase:view -->\n\n## 値の出典\n",
                      encoding="utf-8")
    st = VI.check_files([readme], write=True)[0]
    body = readme.read_text(encoding="utf-8")
    expect("status view: manifest から document / group / 状態と案件の TODO を描く",
           st[1] == "written" and "| d1 | fx | g1 | ✏️ draft |" in body and "repo-a:2026-03-03-x" in body, (st, body))
    m = M.load(case)
    L.annotate(m, "d1", "g1", note="状態の説明はここ")
    m.save()
    expect("annotate --note は note を置き換え、 状態表は変わらない",
           M.load(case).group("d1", "g1")["current"].get("note") == "状態の説明はここ"
           and VI.check_files([readme])[0][1] == "ok")
    m = M.load(case)
    m.group("d1", "g1")["current"]["state"] = "sent"
    m.save()
    expect("manifest の状態が変わると status view は stale", VI.check_files([readme])[0][1] == "stale")
    line = VI.refresh_case(case)
    expect("refresh_case (freeze / annotate / reopen / new が呼ぶ) が描き直す",
           line and "描き直した" in line and "🧊 sent" in readme.read_text(encoding="utf-8")
           and VI.check_files([readme])[0][1] == "ok", line)
    from . import markers as MK2
    expect("refresh_case は隔離 marker も描き直す (凍結 issue ができたら marker が現れ、 audit が stale を出さない)",
           "隔離 marker" in (line or "") and MK2.check_case(M.load(case))[0] == "ok", line)
    other = tmp / "repo-a" / "docs" / "reference" / "x.md"
    other.parent.mkdir(parents=True, exist_ok=True)
    other.write_text("<!-- formcase:view kind=status -->\n<!-- /formcase:view -->\n", encoding="utf-8")
    expect("manifest の無い dir の kind=status は error", VI.check_files([other])[0][1] == "error")

    def state_findings(text, path=readme):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return [x for x in LI.scan(paths=[path], root=tmp) if x[4] == "state"]

    expect("state lint: 案件 README の「提出予定」「未決」「印刷版」 を拾う",
           len(state_findings("## 甲野分 (9/18 提出予定)\n- 未決 (確認): 報告日\n- 印刷版: なし\n")) == 3)
    expect("state lint: history 区間・生成表・backtick の file 名・値の出典表は見ない", not state_findings(
        "<!-- formcase:history -->\n- 8/21 に提出済、 返事待ちだった\n<!-- /formcase:history -->\n"
        "- `houkoku_未提出.pdf` = 確認用\n"
        "| 欄 | 値 | 出典 |\n|---|---|---|\n| 報告日 | 9/18 (提出予定日) | 本人 |\n"))
    expect("state lint: 表の途中の行に「出典」 があっても header に無ければ出典表ではない",
           len(state_findings("| file | 何か |\n|---|---|\n| a.py | 値と出典は docstring |\n| b.txt | 送信済・返事待ち |\n")) == 1)
    sub = tmp / "repo-b" / "docs" / "2026-04-04-x" / "sub" / "README.md"
    expect("state lint: 案件 dir の下の書類 dir の README も案件 README", len(state_findings("## 残\n", sub)) == 1)
    runbook = tmp / "repo-a" / "docs" / "reference" / "runbook.md"
    expect("state lint: 手順 doc (案件 README でない) は見ない (点呼行の規約を説明してよい)",
           not state_findings("TODO notes に 印刷版: の点呼行を書く (未提出のうちに)\n", runbook))
    nested = ("<!-- formcase:history -->\n- 経緯\n<!-- formcase:history -->**強調**<!-- /formcase:history -->\n"
              "- 9/18 提出予定\n<!-- /formcase:history -->\n")
    readme.write_text(nested, encoding="utf-8")
    got_r = [x for x in LI.scan(paths=[readme], root=tmp) if x[4] == "region"]
    expect("region lint: history 区間の入れ子 = 開きと宙に浮いた閉じの 2 件、 その間の状態の手書きも拾う",
           [x[1] for x in got_r] == [3, 5] and state_findings(nested) != [], got_r)
    expect("region lint: 閉じ忘れと、 backtick / code fence の中の marker (書き方の説明) は区別する",
           [x[0] for x in VI.region_problems("a\n<!-- formcase:history -->\nb\n")] == [2]
           and not VI.region_problems("`<!-- formcase:history -->` の書き方\n```\n<!-- /formcase:history -->\n```\n"))
    sp = importlib.util.spec_from_file_location("formcase_cli",
                                                Path(__file__).resolve().parents[1] / "formcase.py")
    cli = importlib.util.module_from_spec(sp)
    sp.loader.exec_module(cli)
    got = cli.staged_case_docs(tmp / "repo-a", ["docs/2026-03-03-readme-case/submission.yaml",
                                                "docs/2026-03-03-readme-case/README.md",
                                                "docs/reference/runbook.md", "TODO.yaml"])
    expect("lint --staged: stage した README と submission.yaml の案件 README だけを選ぶ (重複なし)",
           [p.resolve() for p in got] == [readme.resolve()], got)
    expect("案件の見出しは設定の case_label_with_parent で親も出す",
           cli.case_label(sub.parent) == "2026-04-04-x/sub" and cli.case_label(case) == case.name)
    found = [p.resolve() for p in cli.discover()]
    expect("discover は設定の case_roots を見て、 case_root_skip の下は見ない",
           case.resolve() in found and not any("reference" in str(p) for p in found), found)
