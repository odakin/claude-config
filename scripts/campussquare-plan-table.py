#!/usr/bin/env python3
"""campussquare-plan-table.py — CampusSquare の「授業計画表」 (xlsx 出力) を読み、 教員ごとの登録 (開講期・曜時・科目) を並べて予定と突き合わせる。

学科の取りまとめ役が CampusSquare から出力して配る確認用の xlsx (科目グループごとに 1 file) を、 画面を開かずに照合する。
形式 (実測) と読み方の罠 = conventions/campussquare.md#plan-table-xlsx。 要点:
  - 数行の表題 (「授業計画表」 / 年度 / 科目グループ) の後に見出し行
    `NO | 開講期 | 曜時 | 科目名 | 担当者 | 選必 | 年次 | 曜時2 | 科目名2 | …` (曜日ごとに 5 列の組が横に並ぶ)
  - 1 行に複数の曜日の組が入る。 担当者名は姓と名の間が全角空白
  - 末尾に曜時の無い行 (曜時未設定・担当者空欄) が並ぶことがある

  python3 campussquare-plan-table.py A.xlsx [B.xlsx …] --teacher 山田
      → その教員の行を開講期・曜時の順に並べる (科目グループつき)
  python3 campussquare-plan-table.py A.xlsx B.xlsx --teacher 山田 --expect "前期 月1 科目A" --expect "通年 - 科目B"
      → 予定と照合: ✅ 一致 / ❌ 予定にあるのに登録が無い / ➕ 登録にあるのに予定に無い
  --selftest : 合成の xlsx で回す (名前・科目・年度はすべて架空)

照合の単位 = (開講期, 曜時, 科目名)。 曜時が無い行は予定で `-` と書く。 科目名と教員名は空白 (全角含む) を無視して比べる。
予定は全 file の登録の和と照合する = 新課程と旧課程の合同開講で同じ授業が 2 つの科目グループに別名で出るときは、
旧課程側の名前も予定に書く (書かなければ ➕ で出る)。

終了コード: 0 = 照合が全部一致 (照合しないときも 0) / 1 = ❌ か ➕ がある / 2 = 入力の形が想定外 (見出し行が無い 等)。
依存: openpyxl。
"""
from __future__ import annotations

import argparse
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

TERM_ORDER = {"前期": 0, "後期": 1, "通年": 2, "集中": 3}
DAY_ORDER = {d: i for i, d in enumerate("月火水木金土日他")}


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def norm(s) -> str:
    return re.sub(r"\s+", "", _s(s))


def key(term, slot, course):
    return (term, slot, norm(course))


class FormatError(Exception):
    pass


def parse(path: Path) -> list[dict]:
    """授業計画表 1 file → 行の list (group, year, term, slot, course, teacher, years)。"""
    import openpyxl
    rows_out = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        group = year = ""
        blocks, term_col = None, None
        for row in ws.iter_rows(values_only=True):
            row = list(row)
            head = _s(row[0]) if row else ""
            if head == "科目グループ" and len(row) > 2:
                group = _s(row[2])
                continue
            if head == "年度" and len(row) > 2:
                year = _s(row[2])
                continue
            if head == "NO":
                blocks = [i for i, v in enumerate(row) if _s(v).startswith("曜時")]
                names = [_s(v) for v in row]
                if "開講期" not in names or not blocks:
                    raise FormatError(f"{path.name}: 見出し行に 開講期 / 曜時 が無い")
                term_col = names.index("開講期")
                for b in blocks:
                    if not (names[b + 1:b + 2] and names[b + 1].startswith("科目名")
                            and names[b + 2].startswith("担当者")):
                        raise FormatError(f"{path.name}: 曜時の列の後ろが 科目名 / 担当者 でない (列 {b + 1})")
                continue
            if blocks is None:
                continue
            term = _s(row[term_col]) if term_col < len(row) else ""
            for b in blocks:
                cells = (row + [None] * 5)[b:b + 5]
                slot, course, teacher, _sel, years = (_s(c) for c in cells)
                if course:
                    rows_out.append(dict(group=group, year=year, term=term, slot=slot,
                                         course=course, teacher=teacher, years=years))
    if not any(True for _ in rows_out) and not _has_header(path):
        raise FormatError(f"{path.name}: 見出し行 (NO | 開講期 | 曜時 …) が無い")
    return rows_out


def _has_header(path: Path) -> bool:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return any(row and _s(row[0]) == "NO" for ws in wb.worksheets for row in ws.iter_rows(values_only=True))


def sort_key(r):
    slot = r["slot"]
    day = DAY_ORDER.get(slot[:1], 9) if slot else 10
    period = int(re.sub(r"\D", "", slot) or 0)
    return (TERM_ORDER.get(r["term"], 9), day, period, r["course"], r["group"])


def compare(rows: list[dict], expects: list[str]):
    """→ (一致, 予定にあるのに無い, 登録にあるのに予定に無い) を (開講期, 曜時, 科目名) で。"""
    reg = Counter(key(r["term"], r["slot"] or "-", r["course"]) for r in rows)
    want = []
    for e in expects:
        parts = e.split(maxsplit=2)
        if len(parts) != 3:
            raise FormatError(f'--expect は "開講期 曜時 科目名" の 3 つ (曜時が無ければ -): {e!r}')
        want.append(key(*parts))
    matched = [k for k in want if k in reg]
    missing = [k for k in want if k not in reg]
    extra = [k for k in reg if k not in set(want)]
    return matched, missing, extra


def run(paths, teacher, expects, out=sys.stdout) -> int:
    rows = []
    for p in paths:
        rows += parse(Path(p))
    if teacher:
        t = norm(teacher)
        rows = [r for r in rows if t in norm(r["teacher"])]
    groups = []
    for r in sorted(rows, key=sort_key):
        g = f'{r["group"]} ({r["year"]})' if r["year"] else r["group"]
        if g not in groups:
            groups.append(g)
    for g in groups:
        print(g, file=out)
        for r in sorted((r for r in rows if (f'{r["group"]} ({r["year"]})' if r["year"] else r["group"]) == g),
                        key=sort_key):
            print(f'  {r["term"] or "-":<3} {r["slot"] or "-":<4} {r["course"]}  [{r["teacher"] or "担当者なし"}]'
                  f'  {r["years"]}', file=out)
    if not rows:
        print("(該当する行なし)", file=out)
    if not expects:
        return 0
    matched, missing, extra = compare(rows, expects)
    print("\n照合 (開講期 曜時 科目名):", file=out)
    for k in matched:
        print(f"  ✅ {' '.join(k)}", file=out)
    for k in missing:
        print(f"  ❌ 予定にあるのに登録が無い: {' '.join(k)}", file=out)
    for k in extra:
        print(f"  ➕ 登録にあるのに予定に無い: {' '.join(k)}", file=out)
    return 1 if (missing or extra) else 0


# ---- selftest (合成データ。 名前・科目・年度はすべて架空) ---------------------------------------

def _make_xlsx(path: Path, group: str, rows: list[list], header=True):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None])
    ws.append(["授業計画表"])
    ws.append(["年度", None, "2031年度"])
    ws.append(["科目グループ", None, group])
    if header:
        h = ["NO", "開講期"]
        for i in range(1, 6):
            sfx = "" if i == 1 else str(i)
            h += [f"曜時{sfx}", f"科目名{sfx}", f"担当者{sfx}", "選必", "年次"]
        ws.append(h)
    for r in rows:
        ws.append(r)
    wb.save(path)


def _block(slot, course, teacher, years="2・3・4"):
    return [slot, course, teacher, None, years]


EMPTY = [None] * 5


def selftest() -> int:
    import io
    failures = []

    def check(name, cond, detail=""):
        print(("PASS " if cond else "FAIL ") + name)
        if not cond:
            failures.append(name)
            if detail:
                print("     " + detail.replace("\n", "\n     ")[:800])

    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        a, b, bad = d / "new.xlsx", d / "old.xlsx", d / "bad.xlsx"
        _make_xlsx(a, "【新】合成学科", [
            ["1", "前期"] + _block("月1", "科目甲", "佐藤　花子") + EMPTY + EMPTY
            + _block("木1", "科目乙", "山田　太郎", "3・4") + _block("金1", "科目丙", "佐藤　花子"),
            ["1", "後期"] + EMPTY + EMPTY + EMPTY + EMPTY + _block("金2", "科目丁", "山田　太郎", "1・2・3・4"),
            ["2", "後期"] + EMPTY + EMPTY + EMPTY + _block("木3", "講究2", "山田　太郎", "4") + EMPTY,
            ["1", "", "", "講究9", None, None, "4"],
        ])
        _make_xlsx(b, "【旧】合成学科", [
            ["1", "前期"] + EMPTY + EMPTY + EMPTY + _block("木1", "科目乙", "山田　太郎") + EMPTY,
            ["1", "通年"] + EMPTY + EMPTY + EMPTY + _block("木3", "旧講究", "山田　太郎", "4") + EMPTY,
        ])
        _make_xlsx(bad, "【新】合成学科", [["1", "前期"] + _block("月1", "科目甲", "佐藤　花子")], header=False)

        rows = parse(a)
        check("5 つ目の曜日の列も読む", any(r["course"] == "科目丁" and r["slot"] == "金2" for r in rows),
              repr(rows))
        check("曜時の無い行を落とさない", any(r["course"] == "講究9" and r["slot"] == "" for r in rows), repr(rows))
        check("科目グループと年度を行に付ける", all(r["group"] == "【新】合成学科" and r["year"] == "2031年度"
                                               for r in rows), repr(rows[:2]))

        buf = io.StringIO()
        rc = run([a, b], "山田 太郎", [], out=buf)
        txt = buf.getvalue()
        check("全角空白の名前で当たる", "科目乙" in txt and "科目丁" in txt and rc == 0, txt)
        check("他の教員の行を混ぜない", "科目甲" not in txt and "科目丙" not in txt, txt)
        check("2 file の科目グループを区別して出す", "【新】合成学科" in txt and "【旧】合成学科" in txt, txt)
        check("開講期・曜時の順に並べる",
              "科目乙" in txt and "科目丁" in txt and txt.index("科目乙") < txt.index("科目丁"), txt)

        full = ["前期 木1 科目乙", "後期 金2 科目丁", "後期 木3 講究2", "通年 木3 旧講究"]
        buf = io.StringIO()
        rc = run([a, b], "山田太郎", full, out=buf)
        check("予定どおりなら一致で 0", rc == 0 and "❌" not in buf.getvalue() and "➕" not in buf.getvalue(),
              buf.getvalue())

        buf = io.StringIO()
        rc = run([a, b], "山田太郎", ["後期 木1 科目乙", "後期 金2 科目丁", "後期 木3 講究2", "通年 木3 旧講究"],
                 out=buf)
        txt = buf.getvalue()
        check("学期違いを不一致にする", rc == 1 and "❌ 予定にあるのに登録が無い: 後期 木1 科目乙" in txt
              and "➕ 登録にあるのに予定に無い: 前期 木1 科目乙" in txt, txt)

        buf = io.StringIO()
        rc = run([a, b], "山田太郎", full[:3], out=buf)
        check("旧課程側の登録も予定に無ければ ➕ で出す", rc == 1 and "➕ 登録にあるのに予定に無い: 通年 木3 旧講究"
              in buf.getvalue(), buf.getvalue())

        try:
            parse(bad)
            code = 0
        except FormatError:
            code = 2
        check("見出しが無い file は形式エラー", code == 2)

    print(f"\n{'FAIL' if failures else 'OK'}: {len(failures)} failure(s)")
    return 1 if failures else 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="CampusSquare の授業計画表 (xlsx) を教員ごとに並べ、 予定と照合する")
    ap.add_argument("files", nargs="*", help="授業計画表の xlsx (科目グループごと、 複数可)")
    ap.add_argument("--teacher", default="", help="教員名 (空白は無視、 部分一致)")
    ap.add_argument("--expect", action="append", default=[],
                    help='予定 1 件 = "開講期 曜時 科目名" (曜時が無ければ -)。 複数回')
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args(argv)
    if args.selftest:
        return selftest()
    if not args.files:
        ap.error("xlsx を 1 つ以上")
    try:
        return run(args.files, args.teacher, args.expect)
    except FormatError as e:
        print(f"形式エラー: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
