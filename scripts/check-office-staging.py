#!/usr/bin/env python3
"""check-office-staging.py — Office 事前 grant 済み staging dir が使えない (= wrapper が in-place に落ちる / 落ちた) のを surface する検出器 + 実機 e2e (`--e2e`)。 office-automation.md#office-pregranted-staging-dir

WHY
  層 1 wrapper (docx/xlsx/pptx-to-pdf.sh + affix-image-xlsx.py + office-stage-run.sh) は staging root が
  作れないと旧 in-place に fallback する。 壊れはしないが Office の「ファイル アクセスを許可」 dialog が
  folder ごとに戻り、 remote 操作では押せず -1712 で止まる。 lib (office-staging.sh / office_staging.py) は
  予期せぬ fallback を stderr ⚠️ + fallback log に出す。 本 script はその消費側 (= 呼んだ session 以外、
  別 session / 無人 routine が踏んだ fallback も拾う) と、 新しいマシン・OS / Office 更新の後に
  「このマシンで staging 経路が通るか」 を 1 コマンドで確かめる実機 e2e を持つ。

述語 (= code-as-SoT、 変更時はここが正)
  (A) live probe: macOS ∧ Office group container あり ∧ staging が env で無効化されていない のに
      lib の fallback_reason() が "ok" 以外 (mkdir-failed / not-writable) → 🔴
      (= 次に wrapper を叩いたら fallback する、 を叩く前に知る)。 lib 不在 / 非 macOS / Office 未 install /
      disabled は silent (fail-open)。
  (B) fallback log: lib が書く log (既定 ${XDG_STATE_HOME:-~/.local/state}/claude-office-staging/fallback.log、
      env CLAUDE_OFFICE_STAGING_LOG) のうち、 ack cursor (<log>.acked の行数) より後ろの未 ack entry があれば
      🔴 件数 + 直近 3 行。 `--ack` で cursor を現在行数に進める (log 自体は残す = 履歴)。 log が truncate /
      rotate されて cursor が行数を超えたら全行を未 ack 扱い (= 安全側)。
  - 該当 0 件は沈黙 (SessionStart / dashboard から呼ぶ前提の契約)、 全 error path fail-open。
  - `--fleet-note TEXT`: 対処行に「fleet 実測 = TEXT」 を足す (マシン別の実測記録の所在は呼び出し側が注入)。

e2e (`--e2e`、 macOS の前面で実行 = 初回は osascript の Automation 許可 dialog に応答が要る)
  $HOME 直下の新しい隠し dir に 1 頁の docx / xlsx / pptx を作り (python-docx / openpyxl / python-pptx、
  無い種類は SKIP)、 3 wrapper を順に叩いて、 app ごとに「stderr に staging: 行が出た / PDF が出た /
  頁数 1 / 本文の MARKER 一致 (PyMuPDF があれば)」 を 1 行で出す。 folder-grant dialog が出れば wrapper は
  timeout / -1712 で失敗するので、 成功 = dialog 無しの証拠として扱う (UI を読むには Accessibility が
  要るので dialog そのものは見ない)。 staging root が使えない状態では走らない (= 新しい dir で dialog を
  踏ませない、 exit 2)。 全成功なら probe dir を消す (`--keep` で残す)、 失敗時は残して path を出す。
  app の起動・終了は wrapper の app guard に任せる (user の文書を閉じない・前面を返す)。
  **後始末も見る** (after=): wrapper の後に app guard へ状態を聞き、 app が AppleEvent に答えない
  (= dialog が出ている可能性) か、 自分の staged copy が開いたまま (= staging は掃除済みなので次の起動で
  「開けない」 が出る元) なら、 export が成功していても FAIL にする。 app は起こさず、 quit も kill もしない
  (実測: 長く起動していた PowerPoint で export 成功の後にこの状態になった例がある、 原因は未分離)。

⚠️ 限界: (A) は「root が作れるか」 だけを見る。 root が作れても Office 側が grant を失う形の失敗は (B) か
  e2e でしか拾えず、 (B) は wrapper が実際に走った後にしか出ない。 `--no-stage` / CLAUDE_OFFICE_STAGING=0 の
  意図的 in-place は両方とも射程外 (設計通り)。

Usage
  check-office-staging.py [--fleet-note TEXT]      surface (finding 0 件は沈黙)
  check-office-staging.py --ack                    fallback log を受領 (cursor を進める)
  check-office-staging.py --e2e [--apps word,excel,powerpoint] [--timeout SEC] [--keep] [--parent DIR]
  check-office-staging.py --selftest
"""
from __future__ import annotations

import argparse
import os
import platform
import shutil
import subprocess
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
LIB_DIR = HERE / "lib"
GUARD = LIB_DIR / "office-app-guard.sh"
GROUP_CONTAINER = "Library/Group Containers/UBF8T346G9.Office"
RECENT_N = 3
DOC = "claude-config/conventions/office-automation.md#office-pregranted-staging-dir"

# e2e: app → (wrapper, 入力拡張子, 追加引数 〔入力と出力の間〕, 必要な python module, app bundle)
E2E_APPS = {
    "word": ("docx-to-pdf.sh", "docx", [], "docx", "Microsoft Word.app"),
    "excel": ("xlsx-to-pdf.sh", "xlsx", ["Sheet1"], "openpyxl", "Microsoft Excel.app"),
    "powerpoint": ("pptx-to-pdf.sh", "pptx", [], "pptx", "Microsoft PowerPoint.app"),
}


def _lib():
    """層 1 lib を import (不在なら None = fail-open)."""
    try:
        if str(LIB_DIR) not in sys.path:
            sys.path.insert(0, str(LIB_DIR))
        import office_staging  # type: ignore
        return office_staging
    except Exception:
        return None


def log_path(lib=None) -> Path:
    if lib is not None:
        try:
            return Path(lib.log_path())
        except Exception:
            pass
    return Path(os.environ.get("CLAUDE_OFFICE_STAGING_LOG") or os.path.join(
        os.environ.get("XDG_STATE_HOME") or str(Path.home() / ".local" / "state"),
        "claude-office-staging", "fallback.log"))


def read_log(p: Path) -> list[str]:
    try:
        return [l for l in p.read_text(encoding="utf-8", errors="replace").splitlines() if l.strip()]
    except OSError:
        return []


def ack_cursor(p: Path) -> int:
    try:
        return int((p.with_suffix(p.suffix + ".acked")).read_text().strip() or 0)
    except (OSError, ValueError):
        return 0


def write_ack(p: Path, n: int) -> None:
    p.with_suffix(p.suffix + ".acked").write_text(f"{n}\n")


def unacked(lines: list[str], cursor: int) -> list[str]:
    """cursor より後ろの entry。 log が truncate/rotate されて cursor > len なら全行を未 ack 扱い (= 安全側)."""
    if cursor > len(lines):
        return lines
    return lines[cursor:]


def probe_findings(reason: str | None, *, is_darwin: bool, has_office: bool, disabled: bool) -> list[str]:
    """(A) live probe の判定 (= 引数注入で selftest 可能)."""
    if not is_darwin or not has_office or disabled or reason is None or reason == "ok":
        return []
    if reason in ("disabled", "not-darwin", "no-office"):
        return []  # 矛盾した入力は silent (fail-open)
    return [f"  🔴 staging root が作れない ({reason}) = 次の docx/xlsx/pptx 変換は in-place に落ち Office の「ファイル アクセスを許可」 dialog が出る"]


def log_findings(lines: list[str]) -> list[str]:
    if not lines:
        return []
    out = [f"  🔴 staging fallback が {len(lines)} 件 記録済 (未 ack、 = wrapper が in-place で走った):"]
    for l in lines[-RECENT_N:]:
        parts = l.split("\t")
        ts = parts[0] if parts else "?"
        reason = parts[1] if len(parts) > 1 else "?"
        src = parts[3] if len(parts) > 3 else ""
        out.append(f"     - {ts} {reason} {src}".rstrip())
    if len(lines) > RECENT_N:
        out.append(f"     … 他 {len(lines) - RECENT_N} 件")
    return out


def surface_lines(rows: list[str], fleet_note: str = "") -> list[str]:
    """finding 行に見出しと対処行を付ける (rows が空なら空 = 沈黙)."""
    if not rows:
        return []
    out = ["", "🗂 Office staging dir の fallback (= 事前 grant 済み dir が使えず in-place に落ちた / 落ちる)"]
    out += rows
    out.append(f"  → 対処 = {DOC} 注意 1 (TCC 許可 or CLAUDE_OFFICE_STAGING_DIR=<dir> + 1 回 grant)、")
    note = f"fleet 実測 = {fleet_note}。 " if fleet_note else ""
    out.append(f"    {note}対処後 `{Path(__file__).name} --ack` で log を受領")
    return out


def surface(fleet_note: str = "") -> int:
    lib = _lib()
    lp = log_path(lib)
    rows: list[str] = []
    is_darwin = platform.system() == "Darwin"
    has_office = (Path.home() / GROUP_CONTAINER).is_dir()
    disabled = os.environ.get("CLAUDE_OFFICE_STAGING", "1").strip().lower() in ("0", "no", "off", "false")
    reason = None
    if lib is not None:
        try:
            reason = lib.fallback_reason()
        except Exception:
            reason = None
    rows += probe_findings(reason, is_darwin=is_darwin, has_office=has_office, disabled=disabled)
    rows += log_findings(unacked(read_log(lp), ack_cursor(lp)))
    for line in surface_lines(rows, fleet_note):
        print(line)
    return 0


def ack() -> int:
    lp = log_path(_lib())
    n = len(read_log(lp))
    try:
        lp.parent.mkdir(parents=True, exist_ok=True)
        write_ack(lp, n)
        print(f"acked {n} entries ({lp})")
    except OSError as e:
        print(f"ack failed: {e}", file=sys.stderr)
        return 1
    return 0


# ---------------------------------------------------------------- e2e

def staging_line_state(stderr: str) -> str:
    """wrapper の stderr から staging の状態を 1 語で: staged / unavailable / none."""
    state = "none"
    for line in stderr.splitlines():
        s = line.strip()
        if "staging: UNAVAILABLE" in s:
            return "unavailable"
        if s.startswith("staging: "):
            state = "staged"
    return state


def judge_run(*, rc: int | None, stderr: str, pdf_exists: bool, pages: int | None,
              marker_found: bool | None) -> tuple[bool, str]:
    """1 app の e2e 結果を (ok, 理由) に。 rc=None は timeout。 pages / marker_found=None は未照合 (PyMuPDF 無し)."""
    if rc is None:
        return False, "timeout (dialog 待ちで止まった可能性)"
    st = staging_line_state(stderr)
    if st == "unavailable":
        return False, "staging UNAVAILABLE → in-place に落ちた"
    if rc != 0:
        return False, f"wrapper exit {rc}"
    if not pdf_exists:
        return False, "PDF が出ていない"
    if st != "staged":
        return False, "staging: 行が無い (staging を経由していない)"
    if pages is not None and pages != 1:
        return False, f"頁数 {pages} (期待 1)"
    if marker_found is False:
        return False, "本文 MARKER が PDF に無い (stale / 別文書)"
    return True, "ok" if marker_found else "ok (本文未照合 = PyMuPDF 無し)"


def staged_dir_from(stderr: str) -> str:
    """wrapper の `staging: <dir> (pre-granted…)` 行から staging subdir を取る (無ければ空)."""
    for line in stderr.splitlines():
        s = line.strip()
        if s.startswith("staging: ") and "UNAVAILABLE" not in s:
            return s[len("staging: "):].split(" (", 1)[0]
    return ""


def judge_after(state: str | None, copy_open: bool) -> tuple[bool, str]:
    """wrapper の後の app の状態を (ok, 1 語) に。 state = guard の `state` 出力 (None = 問い合わせ自体が timeout)。

    export が成功しても、 app が AppleEvent に答えない (= dialog が出ている可能性) か、 自分の copy が
    開いたまま (= 次の起動で「開けない」 が出る元) なら、 e2e としては失敗。
    """
    if state is None or state == "unknown":
        return False, "unresponsive"
    if copy_open:
        return False, "copy-open"
    return True, "clean"


def _guard(*args: str, timeout: int = 90) -> tuple[int | None, str]:
    try:
        p = subprocess.run(["bash", str(GUARD), *args], capture_output=True, text=True, timeout=timeout)
        return p.returncode, p.stdout.strip()
    except (subprocess.TimeoutExpired, OSError):
        return None, ""


def after_state(app: str, staged: str) -> tuple[bool, str]:
    """app guard に app の状態を聞き、 自分の staged copy が開いたままでないかも見る (app を起動しない)."""
    if not GUARD.is_file():
        return True, "clean (guard 無し = 未確認)"
    rc, out = _guard("state", app)
    state = out.splitlines()[-1].strip() if (rc is not None and out) else None
    copy_open = False
    if state not in (None, "unknown", "not-running") and staged:
        rc_h, _ = _guard("has-path", app, staged)
        copy_open = rc_h == 0
    return judge_after(state, copy_open)


def _make_fixture(kind: str, path: Path, marker: str) -> None:
    if kind == "docx":
        import docx  # type: ignore
        d = docx.Document()
        d.add_heading("Office staging e2e", 1)
        d.add_paragraph(marker)
        d.save(str(path))
    elif kind == "xlsx":
        import openpyxl  # type: ignore
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Office staging e2e"
        ws["A2"] = marker
        wb.save(str(path))
    elif kind == "pptx":
        from pptx import Presentation  # type: ignore
        pr = Presentation()
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = "Office staging e2e"
        s.placeholders[1].text = marker
        pr.save(str(path))
    else:
        raise ValueError(kind)


def _pdf_check(pdf: Path, marker: str) -> tuple[int | None, bool | None]:
    try:
        import fitz  # type: ignore
    except ImportError:
        return None, None
    try:
        doc = fitz.open(str(pdf))
        text = "".join(p.get_text() for p in doc)
        return doc.page_count, marker in text
    except Exception:
        return None, False


def e2e(apps: list[str], timeout: int, keep: bool, parent: Path | None) -> int:
    if platform.system() != "Darwin":
        print("e2e: macOS 専用 (Office の sandbox 経路の確認)。")
        return 2
    lib = _lib()
    reason = lib.fallback_reason() if lib is not None else "no-lib"
    root = lib.staging_root() if (lib is not None and reason == "ok") else None
    print("Office staging e2e (実機。 初回は osascript の Automation 許可 dialog に前面で応答):")
    print(f"  staging root: {root or '-'} ({reason})")
    if reason != "ok" or not root:
        print(f"  ⛔ staging が使えない状態では走らない (新しい dir で dialog を踏ませない)。 対処 = {DOC} 注意 1")
        return 2
    ts = time.strftime("%Y%m%dT%H%M%S")
    base = parent if parent is not None else Path.home()
    probe = base / f".office-staging-e2e-{ts}-{os.getpid()}"
    probe.mkdir(parents=True)
    results = []
    for app in apps:
        wrapper, ext, extra, module, bundle = E2E_APPS[app]
        if not Path("/Applications", bundle).is_dir():
            results.append((app, "SKIP", f"/Applications/{bundle} が無い", None))
            continue
        marker = f"MARKER-{app.upper()}-{ts}"
        src = probe / f"probe-{app}.{ext}"
        pdf = probe / f"probe-{app}.pdf"
        try:
            _make_fixture(ext, src, marker)
        except ImportError:
            results.append((app, "SKIP", f"python module `{module}` が無い (pip install python-{module})"
                            if module != "openpyxl" else "python module `openpyxl` が無い", None))
            continue
        cmd = ["bash", str(HERE / wrapper), str(src)] + extra + [str(pdf)]
        t0 = time.time()
        try:
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
            rc, err = p.returncode, p.stderr
        except subprocess.TimeoutExpired as ex:
            rc, err = None, (ex.stderr or b"").decode("utf-8", "replace") if isinstance(ex.stderr, bytes) else (ex.stderr or "")
        dt = time.time() - t0
        pages, found = _pdf_check(pdf, marker) if pdf.exists() else (None, None)
        ok, why = judge_run(rc=rc, stderr=err, pdf_exists=pdf.exists(), pages=pages, marker_found=found)
        sdir = staged_dir_from(err)
        ok_after, after = after_state(app, str(Path(sdir) / src.name) if sdir else "")
        if ok and not ok_after:
            why = "export は成功、 後始末で失敗"
        tail = [] if ok else err.strip().splitlines()[-3:]
        if after == "unresponsive":
            tail.append(f"⚠️ {bundle[:-4]} が AppleEvent に答えない = dialog が出ている可能性。 その app の窓を人が見る"
                        " (quit / kill はしない = user の文書を持ちうる)")
        elif after == "copy-open":
            tail.append(f"⚠️ 閉じたはずの copy が {bundle[:-4]} に開いたまま (staging は掃除済み = 次の起動で「開けない」 が出うる)")
        results.append((app, "OK" if (ok and ok_after) else "FAIL",
                        f"{dt:5.1f}s  staging={staging_line_state(err)}  after={after}  {why}", tail))
    for app, status, detail, tail in results:
        print(f"  {app:<11} {status:<5} {detail}")
        for t in tail or []:
            print(f"      | {t}")
    ran = [r for r in results if r[1] != "SKIP"]
    failed = [r for r in ran if r[1] == "FAIL"]
    if failed or keep or not ran:
        print(f"  probe dir を残した: {probe}")
    else:
        shutil.rmtree(probe, ignore_errors=True)
    print(f"結果: {len(ran) - len(failed)}/{len(ran)} OK" + (f" (SKIP {len(results) - len(ran)})" if len(results) != len(ran) else ""))
    if not ran:
        return 2
    return 1 if failed else 0


# ---------------------------------------------------------------- selftest

def selftest() -> int:
    import tempfile
    checks: list[tuple[str, bool]] = []

    def ck(name, cond):
        checks.append((name, bool(cond)))

    # (A) probe 判定
    ck("probe: ok は silent", probe_findings("ok", is_darwin=True, has_office=True, disabled=False) == [])
    ck("probe: not-writable を flag", any("not-writable" in r for r in probe_findings("not-writable", is_darwin=True, has_office=True, disabled=False)))
    ck("probe: mkdir-failed を flag", any("mkdir-failed" in r for r in probe_findings("mkdir-failed", is_darwin=True, has_office=True, disabled=False)))
    ck("probe: disabled は silent", probe_findings("not-writable", is_darwin=True, has_office=True, disabled=True) == [])
    ck("probe: 非 macOS は silent", probe_findings("not-writable", is_darwin=False, has_office=True, disabled=False) == [])
    ck("probe: Office 未 install は silent", probe_findings("not-writable", is_darwin=True, has_office=False, disabled=False) == [])
    ck("probe: lib 不在 (None) は silent", probe_findings(None, is_darwin=True, has_office=True, disabled=False) == [])
    # (B) log + ack
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "fallback.log"
        ck("log: 不在は空", read_log(p) == [] and log_findings(unacked(read_log(p), ack_cursor(p))) == [])
        p.write_text("2026-01-01T00:00:00Z\tnot-writable\t/r\t/a.docx\n2026-01-01T01:00:00Z\tmkdir-failed\t/r\t/b.xlsx\n")
        rows = log_findings(unacked(read_log(p), ack_cursor(p)))
        ck("log: 未 ack 2 件を flag", rows and "2 件" in rows[0])
        ck("log: 直近行に src が出る", any("/b.xlsx" in r for r in rows))
        write_ack(p, 2)
        ck("log: ack 後は silent", log_findings(unacked(read_log(p), ack_cursor(p))) == [])
        p.write_text(p.read_text() + "2026-01-01T02:00:00Z\tnot-writable\t/r\t/c.pptx\n")
        rows = log_findings(unacked(read_log(p), ack_cursor(p)))
        ck("log: ack 後の新着 1 件だけ flag", rows and "1 件" in rows[0] and not any("/a.docx" in r for r in rows))
        p.write_text("2026-01-02T00:00:00Z\tnot-writable\t/r\t/d.docx\n")
        ck("log: rotate で cursor 超過 → 全行未 ack", len(unacked(read_log(p), ack_cursor(p))) == 1)
        p.write_text("".join(f"2026-01-0{i + 1}T00:00:00Z\tnot-writable\t/r\t/{i}.docx\n" for i in range(5)))
        write_ack(p, 0)
        rows = log_findings(unacked(read_log(p), ack_cursor(p)))
        ck("log: 5 件は直近 3 + 省略行", any("他 2 件" in r for r in rows))
    # surface の組み立て
    ck("surface: finding 0 件は沈黙", surface_lines([]) == [])
    s = surface_lines(["  🔴 x"], "記録の在処")
    ck("surface: fleet-note が対処行に入る", any("fleet 実測 = 記録の在処。 対処後" in l for l in s))
    ck("surface: note 無しは fleet 実測 行を出さない", not any("fleet 実測" in l for l in surface_lines(["  🔴 x"])))
    # e2e の判定 (純関数)
    good = "staging: /r/x (pre-granted, no sandbox dialog)\nPDF: /p.pdf\n"
    ck("e2e: staging 行を読む", staging_line_state(good) == "staged")
    ck("e2e: UNAVAILABLE を読む", staging_line_state("⚠️  staging: UNAVAILABLE (not-writable: /r) → in-place\n") == "unavailable")
    ck("e2e: 行が無ければ none", staging_line_state("PDF: /p.pdf\n") == "none")
    ck("e2e: 全条件で ok", judge_run(rc=0, stderr=good, pdf_exists=True, pages=1, marker_found=True)[0])
    ck("e2e: timeout は失敗", not judge_run(rc=None, stderr="", pdf_exists=False, pages=None, marker_found=None)[0])
    ck("e2e: in-place fallback は成功でも失敗扱い",
       not judge_run(rc=0, stderr="⚠️  staging: UNAVAILABLE (x: /r)\n", pdf_exists=True, pages=1, marker_found=True)[0])
    ck("e2e: staging 行なしは失敗", not judge_run(rc=0, stderr="PDF: /p\n", pdf_exists=True, pages=1, marker_found=True)[0])
    ck("e2e: 頁数違いは失敗", not judge_run(rc=0, stderr=good, pdf_exists=True, pages=2, marker_found=True)[0])
    ck("e2e: MARKER 不一致は失敗 (stale)", not judge_run(rc=0, stderr=good, pdf_exists=True, pages=1, marker_found=False)[0])
    ck("e2e: PyMuPDF 無しは未照合 ok", judge_run(rc=0, stderr=good, pdf_exists=True, pages=None, marker_found=None)
       == (True, "ok (本文未照合 = PyMuPDF 無し)"))
    ck("e2e: wrapper の非 0 exit は失敗", not judge_run(rc=1, stderr=good, pdf_exists=True, pages=1, marker_found=True)[0])
    ck("e2e: 3 app の wrapper が隣に在る", all((HERE / v[0]).is_file() for v in E2E_APPS.values()))
    ck("e2e: staging 行から subdir を取る", staged_dir_from("x\nstaging: /r/sub-1 (pre-granted, no sandbox dialog)\n") == "/r/sub-1")
    ck("e2e: UNAVAILABLE 行は subdir にしない", staged_dir_from("⚠️  staging: UNAVAILABLE (x: /r)\n") == "")
    ck("後始末: 起動していない / 空なら clean", judge_after("not-running", False) == (True, "clean")
       and judge_after("clear", False) == (True, "clean"))
    ck("後始末: 応答しない app は失敗", judge_after("unknown", False) == (False, "unresponsive")
       and judge_after(None, False) == (False, "unresponsive"))
    ck("後始末: 自分の copy が開いたままは失敗", judge_after("clear", True) == (False, "copy-open"))
    # lib との契約 (lib が在れば log_path が一致)
    lib = _lib()
    if lib is not None:
        old = os.environ.get("CLAUDE_OFFICE_STAGING_LOG")
        os.environ["CLAUDE_OFFICE_STAGING_LOG"] = "/nonexistent/x/fallback.log"
        ck("lib: log_path が env に従う", str(log_path(lib)) == "/nonexistent/x/fallback.log")
        if old is None:
            del os.environ["CLAUDE_OFFICE_STAGING_LOG"]
        else:
            os.environ["CLAUDE_OFFICE_STAGING_LOG"] = old
        ck("lib: fallback_reason が callable", callable(getattr(lib, "fallback_reason", None)))
    n_pass = sum(1 for _, c in checks if c)
    for name, c in checks:
        print(f"  [{'PASS' if c else 'FAIL'}] {name}")
    print(f"selftest: {n_pass}/{len(checks)} passed")
    return 0 if n_pass == len(checks) else 1


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0], allow_abbrev=False)
    ap.add_argument("--ack", action="store_true")
    ap.add_argument("--fleet-note", default="")
    ap.add_argument("--e2e", action="store_true")
    ap.add_argument("--apps", default="word,excel,powerpoint")
    ap.add_argument("--timeout", type=int, default=300)
    ap.add_argument("--keep", action="store_true")
    ap.add_argument("--parent", default="")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    if a.ack:
        return ack()
    if a.e2e:
        apps = [x.strip() for x in a.apps.split(",") if x.strip()]
        bad = [x for x in apps if x not in E2E_APPS]
        if bad:
            print(f"--apps: 不明 {bad} (word / excel / powerpoint)", file=sys.stderr)
            return 2
        return e2e(apps, a.timeout, a.keep, Path(a.parent).expanduser() if a.parent else None)
    return surface(a.fleet_note)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
